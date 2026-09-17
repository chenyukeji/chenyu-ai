from __future__ import annotations

import math
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from statistics import mean, median
from typing import Any, Iterable

from .io_utils import clamp, parse_number, read_delimited
from .models import Candidate
from .normalizer import infer_product_type, normalize_title


FIELD_ALIASES = {
    "asin": ("asin", "asin码", "商品asin"),
    "title": ("title", "标题", "商品标题", "产品标题"),
    "product_type": ("producttype", "产品类型", "商品类型", "细分类目", "小类目", "subcategory"),
    "sales": ("sales", "月销量", "销量", "monthlysales", "units"),
    "growth": (
        "growth",
        "增长率",
        "销量增长率",
        "销量环比增长率",
        "月增长率",
        "growthrate",
        "salesgrowth",
    ),
    "yoy_growth": ("销量同比增长率", "同比增长率", "yearoveryeargrowth", "yoygrowth"),
    "reviews": ("reviews", "reviewcount", "评论数", "评价数", "评分数", "ratings"),
    "price": ("price", "价格", "售价"),
    "month": ("month", "月份", "年月", "date", "统计月份"),
    "listing_age_days": ("上架天数", "listingagedays", "dayslisted"),
    "weight": (
        "包装重量（单位换算）",
        "包装重量",
        "商品重量（单位换算）",
        "商品重量",
        "packageweight",
        "itemweight",
        "weight",
    ),
    "dimensions": (
        "包装尺寸（单位换算）",
        "包装尺寸",
        "商品尺寸（单位换算）",
        "商品尺寸",
        "packagedimensions",
        "itemdimensions",
        "dimensions",
    ),
}

DEFAULT_MAX_PRICE = 30.0
DEFAULT_MAX_PACKAGE_WEIGHT_G = 1000.0
DEFAULT_MAX_PACKAGE_DIMENSIONS_CM = (45.0, 35.0, 25.0)
DEFAULT_RECENT_LISTING_DAYS = 365


def _key(value: Any) -> str:
    return re.sub(r"[^a-z0-9\u4e00-\u9fff]", "", str(value).casefold())


def _field_map(headers: Iterable[Any]) -> dict[str, str]:
    normalized = {_key(header): str(header) for header in headers if header is not None}
    mapping: dict[str, str] = {}
    for field, aliases in FIELD_ALIASES.items():
        for alias in aliases:
            if _key(alias) in normalized:
                mapping[field] = normalized[_key(alias)]
                break
    missing = [field for field in ("asin", "title") if field not in mapping]
    if missing:
        raise ValueError(f"SellerSprite file is missing required columns: {', '.join(missing)}")
    return mapping


def _read_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Reading .xlsx requires openpyxl. Run: .\\env\\setup.ps1") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        headers = [str(value or "").strip() for value in next(rows)]
    except StopIteration:
        return []
    result = [dict(zip(headers, row)) for row in rows if any(value not in (None, "") for value in row)]
    workbook.close()
    return result


def read_sellersprite(path: str | Path) -> list[dict[str, Any]]:
    source = Path(path)
    suffix = source.suffix.casefold()
    if suffix in {".csv", ".tsv"}:
        return read_delimited(source)
    if suffix == ".xlsx":
        return _read_xlsx(source)
    raise ValueError("SellerSprite input must be .xlsx, .csv, or .tsv")


def _month(value: Any) -> str:
    if isinstance(value, (date, datetime)):
        return value.strftime("%Y-%m")
    text = str(value or "").strip()
    match = re.search(r"(20\d{2})\D*([01]?\d)", text)
    if match:
        return f"{match.group(1)}-{int(match.group(2)):02d}"
    return text


def _growth_percent(value: Any) -> float:
    """Normalize spreadsheet percentage cells and textual percentages to percentage points."""
    parsed = parse_number(value)
    if isinstance(value, (int, float)) and -2 <= parsed <= 2:
        return parsed * 100
    text = str(value or "").strip()
    if text and "%" not in text and -2 <= parsed <= 2:
        return parsed * 100
    return parsed


def _present(value: Any) -> bool:
    return value is not None and str(value).strip() != ""


def _weight_grams(value: Any) -> float | None:
    if not _present(value):
        return None
    text = str(value).strip().casefold().replace(",", ".")
    match = re.search(r"-?\d+(?:\.\d+)?", text)
    if not match:
        return None
    amount = float(match.group())
    if amount < 0:
        return None
    if "kg" in text:
        return amount * 1000
    if "mg" in text:
        return amount / 1000
    if "lb" in text:
        return amount * 453.59237
    if "oz" in text:
        return amount * 28.349523125
    return amount


def _dimensions_cm(value: Any) -> tuple[float, float, float] | None:
    if not _present(value):
        return None
    text = str(value).strip().casefold().replace(",", ".").replace("×", "x").replace("*", "x")
    values = [float(item) for item in re.findall(r"\d+(?:\.\d+)?", text)[:3]]
    if len(values) != 3:
        return None
    if "mm" in text or "毫米" in text:
        values = [item / 10 for item in values]
    elif re.search(r"(?:inch|inches|\bin\b|\")", text):
        values = [item * 2.54 for item in values]
    elif re.search(r"(?:\bm\b|米)", text) and "cm" not in text and "厘米" not in text:
        values = [item * 100 for item in values]
    return tuple(sorted(values, reverse=True))


def _percentile(values: list[float], percentile: float) -> float:
    if not values:
        return 0.0
    values = sorted(values)
    position = (len(values) - 1) * percentile
    lower = math.floor(position)
    upper = math.ceil(position)
    if lower == upper:
        return values[lower]
    return values[lower] + (values[upper] - values[lower]) * (position - lower)


def _positive_percentile_rank(value: float, population: list[float]) -> float:
    positive = [item for item in population if item > 0]
    if value <= 0 or not positive:
        return 0.0
    return sum(1 for item in positive if item <= value) / len(positive)


def _latest_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    latest: dict[str, dict[str, Any]] = {}
    for row in rows:
        current = latest.get(row["asin"])
        if current is None or (row["month"], row["row_number"]) >= (current["month"], current["row_number"]):
            latest[row["asin"]] = row
    return list(latest.values())


def _absolute_sales_increase(row: dict[str, Any]) -> float:
    if not row["growth_known"] or row["sales"] <= 0 or row["growth"] <= -100:
        return 0.0
    previous_sales = row["sales"] / (1 + row["growth"] / 100)
    return max(0.0, row["sales"] - previous_sales)


def _eligibility_reason(
    row: dict[str, Any],
    *,
    max_price: float | None,
    max_package_weight_g: float | None,
    max_package_dimensions_cm: tuple[float, float, float] | None,
    require_complete_size: bool,
) -> str:
    if not row["price_known"]:
        return "missing_price"
    if max_price is not None and row["price"] > max_price:
        return "price_above_limit"
    dimensions = row["dimensions_cm"]
    weight_g = row["weight_g"]
    if require_complete_size and (dimensions is None or weight_g is None):
        return "missing_size_or_weight"
    if max_package_weight_g is not None and weight_g is not None and weight_g > max_package_weight_g:
        return "weight_above_limit"
    if max_package_dimensions_cm is not None and dimensions is not None:
        limits = tuple(sorted(max_package_dimensions_cm, reverse=True))
        if any(actual > limit for actual, limit in zip(dimensions, limits)):
            return "dimensions_above_limit"
    return "eligible"


def analyze_sellersprite(
    path: str | Path,
    *,
    marketplace: str = "",
    spike_growth_threshold: float = 30.0,
    low_review_threshold: int = 100,
    minimum_asins: int = 2,
    max_price: float | None = DEFAULT_MAX_PRICE,
    max_package_weight_g: float | None = DEFAULT_MAX_PACKAGE_WEIGHT_G,
    max_package_dimensions_cm: tuple[float, float, float] | None = DEFAULT_MAX_PACKAGE_DIMENSIONS_CM,
    recent_listing_days: int = DEFAULT_RECENT_LISTING_DAYS,
    require_complete_size: bool = True,
) -> list[Candidate]:
    raw_rows = read_sellersprite(path)
    if not raw_rows:
        return []
    mapping = _field_map(raw_rows[0].keys())
    canonical: list[dict[str, Any]] = []
    seen: set[tuple[str, str]] = set()
    for row_number, row in enumerate(raw_rows, start=2):
        asin = str(row.get(mapping["asin"]) or "").strip().upper()
        title = str(row.get(mapping["title"]) or "").strip()
        month = _month(row.get(mapping.get("month", "")))
        if not asin or not title or (asin, month) in seen:
            continue
        seen.add((asin, month))
        supplied_type = str(row.get(mapping.get("product_type", "")) or "").strip()
        sales_value = row.get(mapping.get("sales", ""))
        growth_value = row.get(mapping.get("growth", ""))
        reviews_value = row.get(mapping.get("reviews", ""))
        price_value = row.get(mapping.get("price", ""))
        canonical.append(
            {
                "asin": asin,
                "title": title,
                "normalized_title": normalize_title(title),
                "product_type": supplied_type or infer_product_type(title),
                "sales": parse_number(sales_value),
                "sales_known": _present(sales_value),
                "growth": _growth_percent(growth_value),
                "growth_known": _present(growth_value),
                "yoy_growth": _growth_percent(row.get(mapping.get("yoy_growth", ""))),
                "reviews": int(parse_number(reviews_value)),
                "reviews_known": _present(reviews_value),
                "price": parse_number(price_value),
                "price_known": _present(price_value) and parse_number(price_value) > 0,
                "month": month,
                "listing_age_days": parse_number(row.get(mapping.get("listing_age_days", "")), -1),
                "weight_g": _weight_grams(row.get(mapping.get("weight", ""))),
                "dimensions_cm": _dimensions_cm(row.get(mapping.get("dimensions", ""))),
                "row_number": row_number,
            }
        )

    rows_by_asin: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in canonical:
        rows_by_asin[row["asin"]].append(row)
    for asin_rows in rows_by_asin.values():
        asin_rows.sort(key=lambda row: (row["month"], row["row_number"]))
        for previous, current in zip(asin_rows, asin_rows[1:]):
            if not current["growth_known"] and previous["sales"] > 0:
                current["growth"] = (current["sales"] / previous["sales"] - 1) * 100
                current["growth_known"] = True

    by_type: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in canonical:
        by_type[row["product_type"]].append(row)

    prepared: list[dict[str, Any]] = []
    source = str(Path(path).resolve())
    for product_type, rows in by_type.items():
        current_rows = _latest_rows(rows)
        evaluated = [
            (
                row,
                _eligibility_reason(
                    row,
                    max_price=max_price,
                    max_package_weight_g=max_package_weight_g,
                    max_package_dimensions_cm=max_package_dimensions_cm,
                    require_complete_size=require_complete_size,
                ),
            )
            for row in current_rows
        ]
        filter_counts = Counter(reason for _, reason in evaluated)
        eligible = [row for row, reason in evaluated if reason == "eligible"]
        unique_asins = sorted({row["asin"] for row in eligible})
        if len(unique_asins) < minimum_asins:
            continue
        spike_rows = [row for row in eligible if row["growth_known"] and row["growth"] >= spike_growth_threshold]
        spike_asins = sorted({row["asin"] for row in spike_rows})
        low_review_success = sorted(
            {
                row["asin"]
                for row in spike_rows
                if row["reviews_known"] and 0 < row["reviews"] <= low_review_threshold
            }
        )
        recent_listing_success = sorted(
            {
                row["asin"]
                for row in spike_rows
                if 0 <= row["listing_age_days"] <= recent_listing_days
            }
        )
        prices = [row["price"] for row in eligible]
        sales_values = [row["sales"] for row in eligible if row["sales"] > 0]
        growth_values = [row["growth"] for row in eligible if row["growth_known"]]
        absolute_increases = [_absolute_sales_increase(row) for row in eligible]
        completeness = mean(
            (
                float(row["sales_known"] and row["sales"] > 0)
                + float(row["growth_known"])
                + float(row["reviews_known"])
            )
            / 3
            for row in eligible
        )
        weights = [row["weight_g"] for row in eligible if row["weight_g"] is not None]
        volumes = [
            row["dimensions_cm"][0] * row["dimensions_cm"][1] * row["dimensions_cm"][2]
            for row in eligible
            if row["dimensions_cm"] is not None
        ]
        prepared.append(
            {
                "product_type": product_type,
                "rows": eligible,
                "all_current_count": len(current_rows),
                "unique_asins": unique_asins,
                "spike_asins": spike_asins,
                "low_review_success": low_review_success,
                "recent_listing_success": recent_listing_success,
                "prices": prices,
                "sales_total": sum(sales_values),
                "sales_median": median(sales_values) if sales_values else 0.0,
                "growth_values": growth_values,
                "median_growth": median(growth_values) if growth_values else 0.0,
                "median_absolute_increase": median(absolute_increases) if absolute_increases else 0.0,
                "growth_breadth": len(spike_asins) / len(eligible),
                "completeness": completeness,
                "median_weight_g": median(weights) if weights else 0.0,
                "median_volume_cm3": median(volumes) if volumes else 0.0,
                "filter_counts": dict(filter_counts),
            }
        )

    sales_medians = [item["sales_median"] for item in prepared]
    sales_totals = [item["sales_total"] for item in prepared]
    absolute_increases = [item["median_absolute_increase"] for item in prepared]
    candidates: list[Candidate] = []
    for item in prepared:
        eligible = item["rows"]
        demand_score = 18 * _positive_percentile_rank(item["sales_median"], sales_medians)
        demand_score += 12 * _positive_percentile_rank(item["sales_total"], sales_totals)
        momentum_score = 18 * _positive_percentile_rank(item["median_absolute_increase"], absolute_increases)
        momentum_score += 12 * clamp(item["median_growth"], 0.0, 100.0) / 100
        breadth_score = 15 * item["growth_breadth"]
        entry_score = 14 * len(item["low_review_success"]) / len(eligible)
        entry_score += 6 * len(item["recent_listing_success"]) / len(eligible)
        confidence_score = 5 * item["completeness"]
        score_components = {
            "demand_30": round(demand_score, 1),
            "momentum_30": round(momentum_score, 1),
            "growth_breadth_15": round(breadth_score, 1),
            "entry_opportunity_20": round(entry_score, 1),
            "data_confidence_5": round(confidence_score, 1),
        }
        score = round(clamp(sum(score_components.values())), 1)
        representatives = sorted(
            eligible,
            key=lambda row: (_absolute_sales_increase(row), row["growth"], row["sales"]),
            reverse=True,
        )
        representative_asins = list(dict.fromkeys(row["asin"] for row in representatives))[:5]
        representative_titles = list(dict.fromkeys(row["title"] for row in representatives))[:3]
        candidates.append(
            Candidate(
                product_type=item["product_type"],
                source="Historical Spike",
                marketplace=marketplace.upper(),
                source_ref=source,
                score=score,
                evidence={
                    "strategy": "short_term_historical_spike",
                    "profit_scored": False,
                    "same_type_asins": len(item["unique_asins"]),
                    "all_current_asins": item["all_current_count"],
                    "spike_asins": len(item["spike_asins"]),
                    "growth_breadth_pct": round(item["growth_breadth"] * 100, 1),
                    "low_review_success": len(item["low_review_success"]),
                    "recent_listing_success": len(item["recent_listing_success"]),
                    "sales_total": round(item["sales_total"], 1),
                    "sales_median": round(item["sales_median"], 1),
                    "average_growth_pct": round(mean(item["growth_values"]), 1) if item["growth_values"] else 0,
                    "median_growth_pct": round(item["median_growth"], 1),
                    "max_growth_pct": round(max(item["growth_values"]), 1) if item["growth_values"] else 0,
                    "median_absolute_sales_increase": round(item["median_absolute_increase"], 1),
                    "median_package_weight_g": round(item["median_weight_g"], 1),
                    "median_package_volume_cm3": round(item["median_volume_cm3"], 1),
                    "data_completeness_pct": round(item["completeness"] * 100, 1),
                    "score_components": score_components,
                    "hard_filter": {
                        "max_price": max_price,
                        "max_package_weight_g": max_package_weight_g,
                        "max_package_dimensions_cm": list(max_package_dimensions_cm)
                        if max_package_dimensions_cm is not None
                        else None,
                        "require_complete_size": require_complete_size,
                        "counts": item["filter_counts"],
                    },
                },
                representative_asins=representative_asins,
                representative_titles=representative_titles,
                price_band={
                    "p25": round(_percentile(item["prices"], 0.25), 2),
                    "median": round(_percentile(item["prices"], 0.5), 2),
                    "p75": round(_percentile(item["prices"], 0.75), 2),
                }
                if item["prices"]
                else {},
            )
        )
    return sorted(
        candidates,
        key=lambda candidate: (
            -(candidate.score or 0),
            candidate.price_band.get("median", float("inf")),
            candidate.evidence.get("median_package_weight_g", float("inf")),
            candidate.evidence.get("median_package_volume_cm3", float("inf")),
        ),
    )
