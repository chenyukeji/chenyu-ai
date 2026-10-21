import importlib.util
import sys
from unittest.mock import patch
import copy
import tempfile
import unittest
import zipfile
from io import BytesIO
from pathlib import Path

BASE = Path(__file__).resolve().parents[1] / 'plugins/chenyu-yunying/skills'
LISTING_ROOT = BASE / 'chenyu-listing/scripts'
YUNYING_ROOT = BASE / 'chenyu-yunying/scripts'


def module(name, root=LISTING_ROOT):
    spec = importlib.util.spec_from_file_location(name, root / (name + '.py'))
    obj = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(obj)
    return obj


extractor = module('extract_development_brief', YUNYING_ROOT)
analyzer = module('analyze_listing')
package_validator = module('validate_listing_package')
description_renderer = module('render_chinese_description')


def validate_legacy(data):
    """Exercise older fixture contracts without using them as current deliveries."""
    return package_validator.validate(data, require_current_search_policy=False)


class ListingTests(unittest.TestCase):
    @staticmethod
    def listing_package():
        return {
            'targets': ['DE'],
            'variants': [{'id': 'V1'}],
            'facts': [{'id': 'F1', 'source_type': 'own_product', 'status': 'confirmed',
                       'field': 'material', 'value': 'Papier', 'variant_ids': ['V1'],
                       'source': {'sheet': 'Product', 'cell': 'B2'}}],
            'keywords': [
                {'marketplace': 'DE', 'phrase': 'Dekoration', 'aliases': ['Deko'], 'is_core': True},
                {'marketplace': 'DE', 'phrase': 'Papierdekoration', 'aliases': ['Dekoration aus Papier'], 'is_core': True},
                {'marketplace': 'DE', 'phrase': 'Festschmuck', 'aliases': ['Feierdeko'], 'is_core': True},
            ],
            'mappings': [{'marketplace': 'DE', 'variant_id': 'V1', 'fact_ids': ['F1'],
                          'buying_reasons': ['Leicht'],
                          'keywords': ['Dekoration', 'Papierdekoration', 'Festschmuck'],
                          'listing_fields': ['title', 'bullet_1', 'bullet_2', 'bullet_3']}],
            'listings': [{'marketplace': 'DE', 'language': 'de-DE', 'variant_id': 'V1',
                          'title_keywords': ['Dekoration', 'Papierdekoration', 'Festschmuck'],
                          'title_scene': 'für Feiern', 'title_quantity': 1,
                          'title_quantity_term': '', 'color_mode': 'not_applicable',
                          'title_color_terms': [], 'critical_differentiators': [],
                          'title': ('Dekoration, Papierdekoration und Festschmuck für Feiern '
                                    'und Feiertage'),
                          'item_highlights': ('Papiermaterial mit klarer Form für Tisch, Regal und '
                                              'Innenraum, einzeln platzierbar oder mit '
                                              'Dekoration kombinierbar'),
                          'bullets': [
                              '📦【Dekoration für kleine Flächen】Die Dekorationselemente setzen auf Tisch, Regal oder Fensterbank einen klaren saisonalen Akzent. Jedes Element kann einzeln platziert oder zusammen mit vorhandener Festdekoration verwendet werden. So lässt sich der verfügbare Platz flexibel nutzen, ohne eine feste Anordnung vorauszusetzen.',
                              '🧩【Leichtes Papiermaterial】Die Dekoration besteht aus Papier und lässt sich unkompliziert an einem trockenen Platz im Innenraum aufstellen. Das leichte Material erleichtert das Umstellen zwischen Tisch, Regal und Fensterbank. Die klaren Formen bleiben dabei gut sichtbar und ergänzen unterschiedliche saisonale Dekorationsstile.',
                              '✨【Flexibel kombinierbar】Die einzelnen Elemente können als ruhiger Akzent verwendet oder mit bereits vorhandener Tisch- und Raumdekoration kombiniert werden. Sie lassen sich nebeneinander verteilen oder an verschiedenen Stellen des Raums platzieren. Dadurch kann die Dekoration an kleine und größere freie Flächen angepasst werden.',
                              '🎉【Für festliche Arrangements】Die Gestaltung eignet sich für Feiern und saisonale Innenraumdekorationen. Auf Esstisch, Kommode oder Regal ergänzt sie bestehende Arrangements, ohne einen bestimmten Aufbau zu verlangen. Nach dem Anlass lassen sich die leichten Elemente abnehmen und für den nächsten Einsatz an einem trockenen Ort aufbewahren.',
                              '💡【Einfach aufstellen und lagern】Die Papierdekoration wird auf einer trockenen, ebenen Fläche platziert und kann bei Bedarf an einen anderen Ort versetzt werden. Für die Aufbewahrung sollte sie vor Feuchtigkeit und starkem Druck geschützt liegen. So bleiben Form und bedruckte Oberfläche zwischen mehreren saisonalen Einsätzen erhalten.',
                          ],
                          'description': ('<p>Eine Dekoration aus Papier für festliche Arrangements. '
                                          'Sie lässt sich einzeln oder zusammen mit vorhandenen Dekorationen einsetzen.</p>'
                                          '<p><b>Eigenschaften:</b><br>'
                                          '1. Platzierung: Die Papierdekoration lässt sich auf geeigneten Flächen gezielt anordnen. Sie ergänzt vorhandene Arrangements, ohne zusätzlich abgebildete Gegenstände als Lieferumfang darzustellen.<br>'
                                          '2. Gestaltung: Die klare Form setzt einen sichtbaren dekorativen Akzent. Sie kann einzeln stehen oder mit abgestimmten Elementen kombiniert werden.<br>'
                                          '3. Anlass: Die Gestaltung ist für bestätigte Feiern und saisonale Innenräume vorgesehen. Der verfügbare Platz bestimmt, ob sie einzeln oder als Teil eines größeren Arrangements verwendet wird.</p>'
                                          '<p><b>Produktdetails:</b><br>Material: Papier<br>Farbe: Weiß<br>Größe: 10 cm</p>'
                                          '<p><b>Lieferumfang:</b><br>1 × Dekoration</p>'),
                          'search_terms': 'dekoartikel schmuckanhänger festbedarf',
                          'front_end_attributes': [],
                          'claim_fact_ids': ['F1'],
                          'field_fact_ids': {
                              'title': ['F1'],
                              'item_highlights': ['F1'],
                              'bullet_1': ['F1'],
                              'bullet_2': ['F1'],
                              'bullet_3': ['F1'],
                              'bullet_4': ['F1'],
                              'bullet_5': ['F1'],
                              'description': ['F1'],
                          },
                          'translations': {
                              'title': '适用于庆典和节日的纸质装饰与派对装饰',
                              'item_highlights': ('纸质材质造型清晰，适合桌面、置物架和室内空间，'
                                                  '可单独摆放或与现有节日装饰搭配'),
                              'bullets': [
                                  '📦【适合小空间装饰】可摆放在桌面、置物架或窗台。',
                                  '🧩【轻巧纸质材质】适合干燥的室内位置。',
                                  '✨【灵活搭配】可单独使用或搭配现有节日装饰。',
                                  '🎉【适合节庆布置】适用于庆典和季节性室内装饰。',
                                  '💡【便于摆放收纳】请置于干燥平面并避免受潮挤压。',
                              ],
                              'description': ('<p>纸质装饰适合节庆布置，可单独使用或搭配现有装饰。</p>'
                                              '<p><b>产品特点：</b><br>'
                                              '1. 摆放：可置于合适的平面。<br>'
                                              '2. 造型：纸质外形清晰，可单独摆放。<br>'
                                              '3. 场合：适合节庆时的室内摆放。</p>'
                                              '<p><b>产品参数：</b><br>材质：纸<br>颜色：白色<br>尺寸：10厘米</p>'
                                              '<p><b>包装内容：</b><br>1 件装饰</p>'),
                              'search_terms': '装饰品 挂饰 派对用品',
                          }}],
        }

    def test_xlsx_evidence(self):
        with tempfile.TemporaryDirectory() as tmp:
            source, out = Path(tmp) / 'sample.xlsx', Path(tmp) / 'out'
            with zipfile.ZipFile(source, 'w') as z:
                z.writestr('xl/workbook.xml', '<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="Product" sheetId="1" r:id="r1"/></sheets></workbook>')
                z.writestr('xl/_rels/workbook.xml.rels', '<Relationships><Relationship Id="r1" Target="worksheets/sheet1.xml"/></Relationships>')
                z.writestr('xl/worksheets/sheet1.xml', '''<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheetData><row r="1">
                <c r="A1" t="inlineStr"><is><t>https://example.com/a https://example.com/b</t></is></c>
                <c r="B1"><f>HYPERLINK("https://example.com/c","ref")</f></c>
                <c r="C1"><f>HYPERLINK(A2,"dynamic")</f></c>
                </row></sheetData><mergeCells><mergeCell ref="A2:B2"/></mergeCells><hyperlinks><hyperlink ref="D1" r:id="h1"/></hyperlinks></worksheet>''')
                z.writestr('xl/worksheets/_rels/sheet1.xml.rels', '<Relationships><Relationship Id="h1" Target="https://example.com/d" TargetMode="External"/></Relationships>')
                z.writestr('xl/embeddings/object.bin', b'not executable')
            result = extractor.extract(source, out)
            self.assertEqual(len(result['links']), 4)
            self.assertEqual(result['links'][-1]['cell'], 'D1')
            self.assertEqual(result['sheets'][0]['merged_ranges'], ['A2:B2'])
            self.assertTrue(any('Dynamic HYPERLINK' in w for w in result['warnings']))
            self.assertEqual(len(result['attachments']), 1)
            with self.assertRaises(ValueError):
                extractor.extract(source, out)

    def test_frequency_union_and_language(self):
        records = [dict(id='a', marketplace='DE', product_group='p1', status='complete', title='Baumschmuck Baumschmuck'),
                   dict(id='b', marketplace='DE', product_group='p1', status='partial', title='Baumdeko'),
                   dict(id='c', marketplace='DE', product_group='p2', status='partial', title='Baumdeko'),
                   dict(id='d', marketplace='FR', product_group='p3', status='complete', title='Baumschmuck'),
                   dict(id='e', marketplace='DE', product_group='p4', status='failed', title='Baumschmuck')]
        report = analyzer.analyze({'competitors': records, 'keywords': [dict(marketplace='DE', phrase='Baumschmuck', aliases=['Baumdeko'])]})
        kw = report['keyword_frequency'][0]
        self.assertEqual((kw['exact_groups'], kw['semantic_groups'], kw['sample_groups']), (1, 2, 2))
        self.assertEqual(kw['field_sample_groups']['description'], 0)

    def test_checks_and_unicode(self):
        phrase = 'one two three four five six seven eight'
        result = analyzer.analyze({'brands': ['ExampleBrand'],
            'competitors': [dict(id='c', marketplace='UK', product_group='p', status='complete', title=phrase)],
            'keywords': [dict(marketplace='UK', phrase='one')],
            'listings': [dict(marketplace='UK', title=phrase + ' ExampleBrand [Brand] B012345678', bullets=[], description='', search_terms='é')],
            'limits': {'UK': {'search_terms_bytes': 1}}})['listing_reviews'][0]
        self.assertEqual(result['lengths']['search_terms_bytes'], 2)
        self.assertEqual(len(result['overlap_review']), 1)
        self.assertIn('title_chars', result['unverified_limits'])
        self.assertTrue(any('exceeds' in i for i in result['issues']))
        self.assertIn('Brand placeholder', result['issues'])
        self.assertFalse(analyzer.contains('stone', 'one'))
        self.assertTrue(analyzer.contains('Étoile', 'étoile'))
        self.assertFalse(analyzer.contains('étoile', 'etoile'))

    def test_item_highlights_requires_a_title_core_keyword(self):
        data = self.listing_package()
        data['listings'][0]['item_highlights'] = (
            data['listings'][0]['item_highlights'].replace(
                'Dekoration kombinierbar', 'Festdeko kombinierbar'
            )
        )
        result = validate_legacy(data)
        self.assertTrue(any(
            'item_highlights must include at least one title core keyword' in error
            for error in result['errors']
        ))

    def test_item_highlights_accepts_verified_title_keyword_alias(self):
        data = self.listing_package()
        data['listings'][0]['item_highlights'] = (
            data['listings'][0]['item_highlights'].replace(
                'Dekoration kombinierbar', 'Deko kombinierbar'
            )
        )
        self.assertTrue(validate_legacy(data)['ready_for_delivery'])

    def test_own_listing_package_ready(self):
        data = self.listing_package()
        result = validate_legacy(data)
        self.assertTrue(result['ready_for_delivery'])
        self.assertEqual(result['coverage']['expected_listings'], 1)
        self.assertTrue(any('editorial target' in warning for warning in result['warnings']))
        review = analyzer.analyze(data)['listing_reviews'][0]
        festschmuck = next(item for item in review['coverage']
                           if item['phrase'] == 'Festschmuck')
        self.assertTrue(festschmuck['selected_for_title'])
        self.assertIn('title', festschmuck['exact_locations'])

    def test_competitor_evidence_requires_references_for_all_listing_fields(self):
        data = self.listing_package()
        data['competitors'] = [{'id': 'C1', 'marketplace': 'DE', 'product_group': 'P1',
                                'status': 'complete', 'asin': 'B012345678',
                                'title': 'Dekoration Papierdekoration Festschmuck Dekoartikel Schmuckanhänger Festbedarf',
                                'bullets': ['Competitor detail 1', 'klarer Form for decoration',
                                            'Competitor detail 3', 'Competitor detail 4',
                                            'Competitor detail 5']}]
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('title_reference is required' in error for error in result['errors']))
        self.assertTrue(any('item_highlights_reference is required' in error
                            for error in result['errors']))
        self.assertTrue(any('description_reference is required' in error for error in result['errors']))
        self.assertTrue(any('search_terms_reference is required' in error for error in result['errors']))
        self.assertTrue(any('bullet_references must contain five' in error for error in result['errors']))

        listing = data['listings'][0]
        listing['title_reference'] = ('Dekoration ← B012345678 标题；Papierdekoration ← B012345678 标题；'
                                      'Festschmuck ← B012345678 标题')
        listing['title_phrase_sources'] = [
            {'phrase': phrase, 'asin': 'B012345678', 'source_field': 'title',
             'source_phrase': phrase}
            for phrase in listing['title_keywords']
        ]
        listing['item_highlights_reference'] = ('Papiermaterial ← 自有事实 F1；'
                                                'klarer Form ← B012345678 第2点')
        listing['item_highlights_phrase_sources'] = [
            {'phrase': 'Papiermaterial', 'fact_id': 'F1'},
            {'phrase': 'klarer Form', 'asin': 'B012345678',
             'source_field': 'bullet_2', 'source_phrase': 'klarer Form'},
        ]
        listing['bullet_references'] = [f'参考 B012345678 第{i}点' for i in range(1, 6)]
        listing['description_reference'] = '参考 B012345678 第1-5点'
        listing['search_terms_reference'] = '参考 B012345678 标题及五点'
        listing['primary_reference_asin'] = 'B012345678'
        data['search_term_audits'] = [
            {
                'marketplace': 'DE', 'variant_id': 'V1',
                'phrase': 'dekoartikel schmuckanhänger festbedarf',
                'source_asin': 'B012345678',
                'source_tool': 'sellersprite_reverse_asin',
                'source_marketplace': 'DE',
                'organic_results_checked': 20, 'relevant_results': 16,
                'relevance_band': 'high', 'decision': 'adopt',
                'local_volume_claimed': False,
            },
            {
                'marketplace': 'DE', 'variant_id': 'V1',
                'phrase': 'dekoartikel schmuckanhänger',
                'source_asin': 'B012345678',
                'source_tool': 'reference_title_terms',
                'source_field': 'title', 'source_marketplace': 'DE',
                'organic_results_checked': 20, 'relevant_results': 15,
                'relevance_band': 'high', 'decision': 'adopt',
                'local_volume_claimed': False,
            },
        ]
        self.assertTrue(validate_legacy(data)['ready_for_delivery'])
        saved_bullet = listing['bullet_references'][0]
        listing['bullet_references'][0] = '自有开发资料 D2/E2；参考竞品标题与图册'
        self.assertTrue(any('must cite ASIN 第N点' in error for error in
                            validate_legacy(data)['errors']))
        listing['bullet_references'][0] = saved_bullet
        data['competitors'][0]['bullets'][0] = ''
        self.assertTrue(any('cites missing competitor point' in error for error in
                            validate_legacy(data)['errors']))
        data['competitors'][0]['bullets'][0] = 'Competitor detail 1'
        listing['title_phrase_sources'][0]['source_phrase'] = 'not in competitor title'
        self.assertTrue(any('source_phrase not found' in error for error in
                            validate_legacy(data)['errors']))
        listing['title_phrase_sources'][0]['source_phrase'] = 'Dekoration'
        listing['item_highlights_reference'] = '参考 B012345678 标题及第2点'
        self.assertTrue(any('phrase missing from visible reference' in error for error in
                            validate_legacy(data)['errors']))
        listing['item_highlights_reference'] = ('Papiermaterial ← 自有事实 F1；'
                                                'klarer Form ← B012345678 第2点')

        original_terms = listing['search_terms']
        data['competitors'].append({
            'id': 'C2', 'marketplace': 'DE', 'product_group': 'P2',
            'status': 'complete', 'asin': 'B012345679',
            'title': 'Tischschmuck aus Papier für Feiern',
            'bullets': [f'Second competitor detail {i}' for i in range(1, 6)],
        })
        extra = {
            'marketplace': 'DE', 'variant_id': 'V1',
            'phrase': 'tischschmuck', 'source_asin': 'B012345679',
            'source_tool': 'same_category_title_terms', 'source_field': 'title',
            'source_marketplace': 'DE', 'organic_results_checked': 20,
            'relevant_results': 16, 'relevance_band': 'high',
            'decision': 'adopt', 'local_volume_claimed': False,
        }
        data['search_term_audits'].append(extra)
        listing['search_terms'] += ' tischschmuck'
        result = validate_legacy(data)
        self.assertTrue(any('title_reference must cite every usable competitor title: B012345679'
                            in error for error in result['errors']))
        # V2 separates the internal candidate audit from visible content sources.
        data['source_mapping_version'] = 2
        v2_errors = validate_legacy(data)['errors']
        self.assertFalse(any('title_reference must cite every usable competitor title'
                             in error for error in v2_errors))
        reference = listing['title_reference']
        listing['title_reference'] = 'Only an unrelated candidate B012345679'
        self.assertTrue(any('title_phrase_sources' in error and 'visible reference' in error
                            for error in validate_legacy(data)['errors']))
        listing['title_reference'] = reference
        del data['source_mapping_version']
        listing['title_reference'] += '；对比 B012345679 标题的产品搜索短语'
        # The primary source remains first; output bullets may merge and reorder sources.
        listing['bullet_references'] = [
            '参考 B012345678 第3点；B012345679 第2点',
            '参考 B012345678 第1点；B012345678 第4点',
            '参考 B012345679 第5点',
            '参考 B012345678 第4点；自有开发资料 Product B2',
            '参考 B012345678 第2点；B012345679 第1点',
        ]
        listing['description_reference'] = '综合 B012345678 第1、3点；B012345679 第2、5点'
        for container in (listing, listing['translations']):
            container['bullets'][0], container['bullets'][1] = (
                container['bullets'][1], container['bullets'][0]
            )
        result = validate_legacy(data)
        self.assertTrue(result['ready_for_delivery'], result['errors'])
        saved_reference = listing['bullet_references'][2]
        listing['bullet_references'][2] = ''
        self.assertTrue(any('bullet_references must contain five' in error
                            for error in validate_legacy(data)['errors']))
        listing['bullet_references'][2] = saved_reference
        extra['source_asin'] = 'B012345699'
        self.assertTrue(any('title-term audit source_asin' in error for error in
                            validate_legacy(data)['errors']))
        extra['source_asin'] = 'B012345679'
        extra['source_field'] = 'bullet_1'
        self.assertTrue(any('must use source_field=title' in error for error in
                            validate_legacy(data)['errors']))
        data['search_term_audits'].pop()
        data['competitors'].pop()
        listing['search_terms'] = original_terms

        data['search_term_audits'] = data['search_term_audits'][:1]
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('requires synonym candidates from reference listing titles' in error
                            for error in result['errors']))

    def test_listing_package_rejects_internal_variant_codes_in_buyer_copy(self):
        data = copy.deepcopy(self.listing_package())
        data['listings'][0]['item_highlights'] = (
            'Design A: Papiermaterial mit klarer Form für Tisch, Regal und Innenraum, '
            'einzeln platzierbar oder mit Dekoration kombinierbar'
        )
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('contains internal variant codes' in error
                            for error in result['errors']))

        data['listings'][0]['buyer_visible_variant_terms'] = ['Design A']
        self.assertTrue(validate_legacy(data)['ready_for_delivery'])

    def test_buyer_copy_rejects_internal_confirmation_and_reference_language(self):
        errors = []
        package_validator.validate_buyer_copy(
            errors,
            'FR',
            'V1',
            [
                ('bullet 1', '【Contenu confirmé】Le kit contient des fils et des aiguilles.'),
                ('bullet 2', ('Les dimensions indiquées concernent l’emballage. '
                              'La pochette du produit de référence ne fait pas partie du kit.')),
            ],
        )
        self.assertEqual(
            sum('contains internal confirmation' in error for error in errors), 2
        )

    def test_listing_package_rejects_internal_process_language_in_chinese_translation(self):
        data = copy.deepcopy(self.listing_package())
        data['listings'][0]['translations']['bullets'][0] = (
            '📦【已确认的套装内容】包装资料记录了尺寸；自有套装不包含竞品图中的收纳包。'
        )
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('Chinese translation bullet 1 contains internal' in error
                            for error in result['errors']))

    def test_listing_package_requires_complete_chinese_translations(self):
        data = copy.deepcopy(self.listing_package())
        del data['listings'][0]['translations']
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('translations must be an object' in error
                            for error in result['errors']))

    def test_chinese_description_preserves_html_format(self):
        data = copy.deepcopy(self.listing_package())
        listing = data['listings'][0]
        self.assertTrue(validate_legacy(data)['ready_for_delivery'])
        translated = listing['translations']['description']
        listing['translations']['description'] = package_validator.visible_html(translated)
        result = validate_legacy(data)
        self.assertTrue(any('Chinese description HTML layout must match' in error
                            for error in result['errors']))

        listing['translations']['description'] = translated.replace('2. 造型', '4. 造型')
        result = validate_legacy(data)
        self.assertTrue(any('Chinese description feature numbers must match' in error
                            for error in result['errors']))

        listing['translations']['description'] = translated.replace('<p>', r'\<p>', 1)
        result = validate_legacy(data)
        self.assertTrue(any('Chinese description must not escape HTML tags' in error
                            for error in result['errors']))

    def test_chinese_description_renders_paragraphs_and_bold_headings(self):
        source = ('<p>两条紫色灯串，适合室内装饰。</p>'
                  '<p><b>产品特点：</b><br>1. 两条独立灯串。<br>'
                  '2. 可弯折细灯线。</p>'
                  '<p><b>包装内容：</b><br>2 条灯串。</p>')
        runs = description_renderer.description_runs(source)
        visible = ''.join(value for value, _ in runs)
        self.assertIn('室内装饰。\n\n产品特点：\n1. 两条独立灯串。', visible)
        self.assertIn(('产品特点：', True), runs)
        self.assertNotIn('<p>', visible)

        try:
            from openpyxl import Workbook, load_workbook
            from openpyxl.cell.rich_text import TextBlock
        except ImportError:
            self.skipTest('openpyxl is available in the backend workbook environment')
        book = Workbook()
        cell = book.active['C8']
        description_renderer.write_chinese_description(cell, source)
        self.assertTrue(cell.alignment.wrap_text)
        self.assertGreater(cell.parent.row_dimensions[8].height, 30)
        stream = BytesIO()
        book.save(stream)
        stream.seek(0)
        loaded = load_workbook(stream, rich_text=True).active['C8'].value
        self.assertEqual(str(loaded), visible)
        self.assertTrue(any(isinstance(part, TextBlock) and part.font.b
                            and part.text == '产品特点：' for part in loaded))

    def test_primary_reference_does_not_require_a_mirrored_outline(self):
        data = self.listing_package()
        listing = data['listings'][0]
        listing['primary_reference_asin'] = 'B012345678'
        result = validate_legacy(data)
        self.assertTrue(result['ready_for_delivery'], result['errors'])

    def test_integrated_copy_still_requires_confirmed_field_facts(self):
        data = self.listing_package()
        listing = data['listings'][0]
        listing['primary_reference_asin'] = 'B012345678'
        listing['field_fact_ids']['bullet_2'] = []
        listing['field_fact_ids']['description'] = []
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('field_fact_ids.bullet_2 must contain' in error
                            for error in result['errors']))
        self.assertTrue(any('field_fact_ids.description must contain' in error
                            for error in result['errors']))
        listing['field_fact_ids']['bullet_2'] = ['F1']
        listing['field_fact_ids']['description'] = ['F1']
        data['facts'][0]['status'] = 'unconfirmed'
        self.assertFalse(validate_legacy(data)['ready_for_delivery'])

    def test_fitment_products_require_confirmed_compatibility_in_front_end_copy(self):
        data = copy.deepcopy(self.listing_package())
        listing = data['listings'][0]
        listing['compatibility_required'] = True
        listing['primary_compatibility_term'] = 'PAULTRA2'
        listing['compatibility_terms'] = [
            'PAULTRA2', 'PureAir Ultra 2', '242047805', '5303918847', 'EAP12364179'
        ]
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('title must contain primary compatibility term' in error
                            for error in result['errors']))
        self.assertTrue(any('omit confirmed compatibility terms' in error
                            for error in result['errors']))

    def test_compatibility_validator_accepts_models_and_part_numbers_in_both_languages(self):
        listing = {
            'compatibility_required': True,
            'primary_compatibility_term': 'PAULTRA2',
            'compatibility_terms': [
                'PAULTRA2', 'PureAir Ultra 2', '242047805', '5303918847', 'EAP12364179'
            ],
            'translations': {
                'title': '6个装PAULTRA2冰箱空气过滤器替换滤芯，95 × 45 × 9毫米',
                'item_highlights': ('兼容PureAir Ultra 2，以及242047805、5303918847、'
                                    'EAP12364179替换件号'),
                'bullets': [
                    '🔧【兼容PAULTRA2】适用于PAULTRA2冰箱空气过滤系统。',
                    '📦【6个装】提供多个替换滤芯。',
                    '🌿【过滤异味】帮助过滤冰箱空气。',
                    '🛠️【便于安装】可装入对应滤芯仓。',
                    '✅【日常替换】适合定期维护。',
                ],
                'description': ('兼容PAULTRA2、PureAir Ultra 2、242047805、'
                                '5303918847和EAP12364179。'),
            },
        }
        errors = []
        package_validator.validate_compatibility_copy(
            errors,
            'FR',
            'V1',
            listing,
            '6 filtres à air PAULTRA2 de rechange pour réfrigérateur, 95 x 45 x 9 mm',
            ('Compatibles Frigidaire PureAir Ultra 2 et Electrolux 242047805, '
             '5303918847, EAP12364179 ; lot de 6'),
            [
                '🔧【Compatibilité PAULTRA2】Compatible avec le système PAULTRA2.',
                '📦【Lot de 6】Six filtres de rechange.',
                '🌿【Filtration】Aide à filtrer l’air du réfrigérateur.',
                '🛠️【Installation】S’insère dans le logement compatible.',
                '✅【Entretien】Convient au remplacement régulier.',
            ],
            ('Compatible PAULTRA2, PureAir Ultra 2, 242047805, 5303918847 '
             'et EAP12364179.'),
        )
        self.assertEqual(errors, [])

        listing['translations']['item_highlights'] = '请核对原滤芯尺寸和冰箱说明书。'
        package_validator.validate_compatibility_copy(
            errors,
            'FR',
            'V1',
            listing,
            '6 filtres à air PAULTRA2 de rechange pour réfrigérateur, 95 x 45 x 9 mm',
            ('Compatibles Frigidaire PureAir Ultra 2 et Electrolux 242047805, '
             '5303918847, EAP12364179 ; lot de 6'),
            [
                '🔧【Compatibilité PAULTRA2】Compatible avec le système PAULTRA2.',
                '📦【Lot de 6】Six filtres de rechange.',
                '🌿【Filtration】Aide à filtrer l’air du réfrigérateur.',
                '🛠️【Installation】S’insère dans le logement compatible.',
                '✅【Entretien】Convient au remplacement régulier.',
            ],
            ('Compatible PAULTRA2, PureAir Ultra 2, 242047805, 5303918847 '
             'et EAP12364179.'),
        )
        self.assertTrue(any('Chinese title and item_highlights omit confirmed' in error
                            for error in errors))

    def test_listing_package_requires_field_level_fact_assignments(self):
        data = copy.deepcopy(self.listing_package())
        listing = data['listings'][0]
        del listing['field_fact_ids']['description']
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('field_fact_ids.description must contain' in error
                            for error in result['errors']))

        listing['field_fact_ids']['description'] = ['F2']
        result = validate_legacy(data)
        self.assertTrue(any('references unknown fact F2' in error
                            for error in result['errors']))

    def test_fitment_copy_requires_primary_term_in_bullet_and_full_list_in_description(self):
        listing = {
            'compatibility_required': True,
            'primary_compatibility_term': 'PAULTRA2',
            'compatibility_terms': ['PAULTRA2', 'PureAir Ultra 2', '242047805'],
            'translations': {
                'title': 'PAULTRA2冰箱空气过滤器',
                'item_highlights': '兼容PureAir Ultra 2和242047805',
                'bullets': ['🔧【尺寸】单个尺寸95 × 45 × 9毫米。'],
                'description': '兼容PAULTRA2和PureAir Ultra 2。',
            },
        }
        errors = []
        package_validator.validate_compatibility_copy(
            errors, 'FR', 'V1', listing,
            'Filtre à air PAULTRA2 pour réfrigérateur',
            'Compatible PureAir Ultra 2 et 242047805',
            ['🔧【Dimensions】Chaque filtre mesure 95 x 45 x 9 mm.'],
            'Compatible PAULTRA2 et PureAir Ultra 2.',
        )
        self.assertTrue(any('bullet 1 must lead with primary compatibility term' in error
                            for error in errors))
        self.assertTrue(any('description omits confirmed compatibility terms: 242047805'
                            in error for error in errors))
        self.assertTrue(any('Chinese bullet 1 must contain primary compatibility term'
                            in error for error in errors))
        self.assertTrue(any('Chinese description omits confirmed compatibility terms'
                            in error for error in errors))

    @classmethod
    def audited_listing_package(cls):
        data = copy.deepcopy(cls.listing_package())
        data['competitors'] = [{
            'id': 'C1', 'marketplace': 'DE', 'product_group': 'P1',
            'status': 'complete', 'asin': 'B012345678',
            'title': 'Dekoration Papierdekoration Festschmuck Dekoartikel Schmuckanhänger Festbedarf Wintermotiv',
            'bullets': ['Competitor detail 1', 'klarer Form for decoration',
                        'Competitor detail 3', 'Competitor detail 4',
                        'Competitor detail 5'],
        }]
        listing = data['listings'][0]
        listing.update({
            'title_reference': ('Dekoration ← B012345678 标题；Papierdekoration ← B012345678 标题；'
                                'Festschmuck ← B012345678 标题'),
            'title_phrase_sources': [
                {'phrase': phrase, 'asin': 'B012345678', 'source_field': 'title',
                 'source_phrase': phrase}
                for phrase in listing['title_keywords']
            ],
            'item_highlights_reference': ('Papiermaterial ← 自有事实 F1；'
                                          'klarer Form ← B012345678 第2点'),
            'item_highlights_phrase_sources': [
                {'phrase': 'Papiermaterial', 'fact_id': 'F1'},
                {'phrase': 'klarer Form', 'asin': 'B012345678',
                 'source_field': 'bullet_2', 'source_phrase': 'klarer Form'},
            ],
            'bullet_references': [f'参考 B012345678 第{i}点' for i in range(1, 6)],
            'description_reference': '参考 B012345678',
            'search_terms_reference': '参考 B012345678 标题及卖家精灵反查',
            'primary_reference_asin': 'B012345678',
        })
        data['search_term_audits'] = [
            {
                'marketplace': 'DE', 'variant_id': 'V1',
                'phrase': 'dekoartikel schmuckanhänger festbedarf',
                'source_asin': 'B012345678',
                'source_tool': 'sellersprite_reverse_asin',
                'source_marketplace': 'DE',
                'organic_results_checked': 20, 'relevant_results': 16,
                'relevance_band': 'high', 'decision': 'adopt',
                'local_volume_claimed': False,
            },
            {
                'marketplace': 'DE', 'variant_id': 'V1',
                'phrase': 'wintermotiv', 'source_asin': 'B012345678',
                'source_tool': 'reference_title_terms', 'source_field': 'title',
                'source_marketplace': 'DE',
                'organic_results_checked': 20, 'relevant_results': 15,
                'relevance_band': 'high', 'decision': 'adopt',
                'local_volume_claimed': False,
            },
        ]
        return data

    def test_listing_package_requires_all_adopted_incremental_tokens(self):
        data = self.audited_listing_package()
        listing = data['listings'][0]
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('omits audited incremental tokens: wintermotiv' in error
                            for error in result['errors']))

        listing['search_terms'] += ' wintermotiv'
        self.assertTrue(validate_legacy(data)['ready_for_delivery'])

        # Additional real ASIN reverse lookups can supply distinct local terms.
        data['competitors'].append({
            'id': 'C2', 'marketplace': 'DE', 'product_group': 'P2',
            'status': 'complete', 'asin': 'B087654321',
            'title': 'Papierdekoration Wanddeko', 'bullets': ['Material Papier'],
        })
        data['search_term_audits'].append({
            'marketplace': 'DE', 'variant_id': 'V1', 'phrase': 'wanddeko',
            'source_asin': 'B087654321', 'source_tool': 'sellersprite_reverse_asin',
            'source_marketplace': 'DE', 'organic_results_checked': 20,
            'relevant_results': 16, 'relevance_band': 'high',
            'decision': 'adopt', 'local_volume_claimed': False,
        })
        listing['search_terms'] += ' wanddeko'
        result = validate_legacy(data)
        self.assertFalse(any('SellerSprite audit source_asin' in error
                             for error in result['errors']))
        self.assertFalse(any('omits audited incremental tokens: wanddeko' in error
                             for error in result['errors']))

    def test_listing_package_rejects_unified_format_violations(self):
        data = copy.deepcopy(self.listing_package())
        data['listings'][0]['bullets'][0] = 'Lieferumfang: Nur ein kurzer Satz.'
        data['listings'][0]['search_terms'] = 'dekoration, papier papier'
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('bullet 1 must use' in error for error in result['errors']))
        self.assertTrue(any('must not contain punctuation' in error for error in result['errors']))
        self.assertTrue(any('repeats tokens: papier' in error for error in result['errors']))

    def test_listing_package_enforces_incremental_search_terms(self):
        data = copy.deepcopy(self.listing_package())
        data['listings'][0]['search_terms'] = 'Dekoartikel und Papier ' + ('ä' * 120)
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('must be lowercase' in error for error in result['errors']))
        self.assertTrue(any('contains stop words: und' in error for error in result['errors']))
        self.assertTrue(any('repeats front-end tokens: papier' in error for error in result['errors']))
        self.assertTrue(any('UTF-8 bytes' in error for error in result['errors']))

    def test_item_highlights_is_required_and_included_in_front_end_deduplication(self):
        data = copy.deepcopy(self.listing_package())
        data['listings'][0]['item_highlights'] = ''
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('item_highlights is empty' in error for error in result['errors']))

        data = copy.deepcopy(self.listing_package())
        data['listings'][0]['search_terms'] = 'innenraum'
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('repeats front-end tokens: innenraum' in error
                            for error in result['errors']))

    def test_same_product_evidence_is_a_valid_confirmed_fact_source(self):
        data = copy.deepcopy(self.listing_package())
        fact = data['facts'][0]
        fact['source_type'] = 'same_product_evidence'
        fact['same_product_confirmed'] = True
        fact['source'] = {
            'asin': 'B012345678', 'field': 'bullet_2', 'user_message': '同款确认'
        }
        self.assertTrue(validate_legacy(data)['ready_for_delivery'])

        fact['same_product_confirmed'] = False
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('lacks explicit same-product confirmation' in error
                            for error in result['errors']))

    def test_same_material_competitor_function_can_be_used(self):
        data = copy.deepcopy(self.listing_package())
        data['facts'].append({
            'id': 'F2', 'source_type': 'material_match_evidence',
            'status': 'confirmed', 'field': 'function',
            'value': 'Papierdekoration lässt sich leicht umstellen',
            'variant_ids': ['V1'],
            'source': {
                'asin': 'B012345678', 'field': 'bullet_2',
                'material': 'Papier', 'own_material_fact_id': 'F1',
            },
        })
        listing = data['listings'][0]
        listing['claim_fact_ids'].append('F2')
        listing['field_fact_ids']['description'].append('F2')
        data['mappings'][0]['fact_ids'].append('F2')
        self.assertTrue(validate_legacy(data)['ready_for_delivery'])

    def test_bullet_copy_has_no_artificial_maximum(self):
        data = copy.deepcopy(self.listing_package())
        data['listings'][0]['bullets'][0] = (
            '📦【Ausführliche Produktangabe】Das Set enthält vier einzeln '
            'gestaltete Figuren mit klar erkennbaren Formen, stabilen Aufstellflächen und '
            'abgestimmten Details für Regale, Tische und saisonale Dekorationen. '
            'Jede Figur lässt sich separat platzieren, nach dem Umstellen erneut ausrichten '
            'und mit vorhandenen Dekorationen kombinieren, ohne dass eine feste Reihenfolge '
            'oder ein bestimmter Aufbau erforderlich ist.'
        )
        self.assertTrue(validate_legacy(data)['ready_for_delivery'])

    def test_listing_package_rejects_repeated_bullet_sentences(self):
        data = copy.deepcopy(self.listing_package())
        repeated = (
            'Die vier Figuren lassen sich einzeln auf Regalen, Tischen und Fensterbänken '
            'platzieren und nach dem Umstellen erneut ausrichten'
        )
        data['listings'][0]['bullets'][0] = (
            f'📦【Flexible Platzierung】{repeated}. {repeated}.'
        )
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('repeats the same sentence' in error
                            for error in result['errors']))

    def test_listing_package_rejects_thin_bullet_copy(self):
        data = copy.deepcopy(self.listing_package())
        data['listings'][0]['bullets'][0] = (
            '📦【Klarer Lieferumfang】Die Variante enthält die angegebenen Teile. '
            'Der Inhalt lässt sich vorab überblicken.'
        )
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('more than 200 visible characters' in error
                            for error in result['errors']))

    def test_listing_package_requires_three_core_keywords_and_optional_scene(self):
        data = copy.deepcopy(self.listing_package())
        self.assertTrue(validate_legacy(data)['ready_for_delivery'])
        data['listings'][0]['title_keywords'] = ['Dekoration', 'Papierdekoration']
        result = validate_legacy(data)
        self.assertTrue(any('title_keywords must contain exactly 3 distinct phrases' in error
                            for error in result['errors']))
        data['listings'][0]['title_keywords'] = ['Dekoration', 'Papierdekoration', 'Festschmuck', 'Feierdeko']
        result = validate_legacy(data)
        self.assertTrue(any('title_keywords must contain exactly 3 distinct phrases' in error
                            for error in result['errors']))
        data = copy.deepcopy(self.listing_package())
        data['keywords'][2]['is_core'] = False
        data['listings'][0]['title_scene'] = 'für Hochzeiten'
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('title keyword is not marked as core: Festschmuck' in error
                            for error in result['errors']))
        self.assertTrue(any('title does not contain title_scene: für Hochzeiten' in error
                            for error in result['errors']))

    def test_short_title_is_valid_without_size_padding(self):
        data = copy.deepcopy(self.listing_package())
        data['listings'][0]['title'] = 'Dekoration, Papierdekoration und Festschmuck für Feiern'
        self.assertTrue(validate_legacy(data)['ready_for_delivery'])

    def test_item_highlights_below_target_is_editorial_warning(self):
        data = copy.deepcopy(self.listing_package())
        data['listings'][0]['item_highlights'] = (
            'Dekoration aus Papier für Tisch und Regal mit klarer Form und '
            'flexibler Platzierung bei saisonalen Feiern'
        )
        result = validate_legacy(data)
        self.assertTrue(result['ready_for_delivery'])
        self.assertTrue(any('120-125 editorial target' in warning
                            for warning in result['warnings']))

    def test_title_size_requires_reference_reason_and_final_position(self):
        data = copy.deepcopy(self.listing_package())
        listing = data['listings'][0]
        listing['title_scene'] = ''
        listing['title'] = 'Dekoration, Papierdekoration und Festschmuck, 33 x 183 cm'
        result = validate_legacy(data)
        self.assertTrue(any('title_size_term must match' in error for error in result['errors']))
        self.assertTrue(any('title_size_reason must record' in error for error in result['errors']))
        listing['title_size_term'] = '33 x 183 cm'
        listing['title_size_reason'] = 'B012345678 title: Table runner 33 x 183 cm'
        self.assertTrue(validate_legacy(data)['ready_for_delivery'])
        listing['title'] = '33 x 183 cm Dekoration, Papierdekoration und Festschmuck'
        result = validate_legacy(data)
        self.assertTrue(any('title size must appear at the end' in error for error in result['errors']))
        listing['title'] = 'Dekoration 10 cm, Papierdekoration und Festschmuck, 33 x 183 cm'
        result = validate_legacy(data)
        self.assertTrue(any('title may contain at most one size group' in error
                            for error in result['errors']))

    def test_listing_package_requires_html_and_rejects_multicolor_title_terms(self):
        data = copy.deepcopy(self.listing_package())
        listing = data['listings'][0]
        listing['color_mode'] = 'multi'
        listing['title_color_terms'] = ['Rot']
        listing['title'] += ' Rot'
        listing['description'] = (
            'Eine Beschreibung. Eigenschaften: 1. Material: Papier. '
            '2. Form: Dekoration. 3. Anlass: Feiern. '
            'Produktdetails: Material: Papier. Lieferumfang: 1 Dekoration.'
        )
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('must not use title_color_terms when color_mode is multi' in error
                            for error in result['errors']))
        self.assertTrue(any('description must use basic HTML' in error
                            for error in result['errors']))

    def test_listing_package_rejects_bad_title_spacing_and_checks_optional_notices(self):
        data = copy.deepcopy(self.listing_package())
        listing = data['listings'][0]
        listing['title'] = 'Dekoration,Papierdekoration und Festschmuck für Feiern'
        listing['notice_fact_ids'] = ['F1']
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('title uses non-standard punctuation spacing' in error
                            for error in result['errors']))
        self.assertTrue(any('must include localized notices' in error
                            for error in result['errors']))

    def test_title_quantity_is_omitted_for_one_and_precedes_first_keyword_for_multipacks(self):
        data = copy.deepcopy(self.listing_package())
        listing = data['listings'][0]
        listing['title_quantity'] = 12
        listing['title_quantity_term'] = '12 Stück'
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('title quantity must immediately precede' in error
                            for error in result['errors']))

        listing['title'] = ('12 Stück Dekoration, Dekoration aus Papier und Festschmuck '
                            'für Feiern')
        self.assertTrue(validate_legacy(data)['ready_for_delivery'])

        listing['title_quantity'] = 1
        listing['title_quantity_term'] = '1 Stück'
        listing['title'] = '1 Dekoration, Papierdekoration und Festschmuck für Feiern'
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('title_quantity_term must be empty' in error
                            for error in result['errors']))
        self.assertTrue(any('title must omit quantity 1' in error
                            for error in result['errors']))

    def test_item_highlights_does_not_repeat_title_pack_count(self):
        data = copy.deepcopy(self.listing_package())
        listing = data['listings'][0]
        listing['title_quantity'] = 8
        listing['title_quantity_term'] = '8'
        listing['title'] = '8 Dekoration, Papierdekoration und Festschmuck für Feiern'
        listing['item_highlights'] = '8 Dekoration aus Papier für Tisch und Regal mit klarer Form'
        result = validate_legacy(data)
        self.assertTrue(any('item_highlights repeats the title pack quantity'
                            in error for error in result['errors']))
        self.assertTrue(package_validator.repeats_pack_count_in_highlights(
            '8 sacs cadeaux de Noël en tissu non tissé',
            8, '8', ['sacs cadeaux de Noël']))
        self.assertFalse(package_validator.repeats_pack_count_in_highlights(
            'Sacs cadeaux de Noël avec 8 motifs pour différentes occasions',
            8, '8', ['sacs cadeaux de Noël']))

    @classmethod
    def localized_search_package(cls):
        data = cls.audited_listing_package()
        data['search_term_policy_version'] = 4
        data['listings'][0]['search_terms'] += ' wintermotiv'
        data['competitors'][0]['title'] += ' Winter motif'
        data['title_candidate_audit'] = {'all_user_references': [{'asin': 'B012345678'}]}
        for audit, original in zip(data['search_term_audits'],
                                   ['paper ornaments hanging decoration', 'Winter motif']):
            audit.update(source_phrase=original, source_marketplace='US',
                         source_language='en', target_language='de', localization_method='ai',
                         search_marketplace='DE', search_query=audit['phrase'],
                         search_evidence='internal/de-local-query-results.json',
                         specificity='product_specific', product_match_reason='Matches the paper ornament product type.')
        data['search_term_concepts'] = []
        data['search_query_inventory'] = []
        for index, audit in enumerate(data['search_term_audits']):
            audit['concept_id'] = str(index)
            data['search_term_concepts'].append({'id': str(index), 'source_phrase': audit['source_phrase'], 'meaning': 'Product-specific source concept'})
        data['listings'][0]['search_term_expansion'] = {
            'queries': [a['phrase'] for a in data['search_term_audits']],
            'reviewed_phrases': [a['phrase'] for a in data['search_term_audits']],
            'shortfall_reason': 'All original title and reverse-ASIN concepts were localized; available incremental terms are included.',
        }
        return data

    def test_ai_localized_keywords_can_supply_incremental_terms(self):
        data = self.localized_search_package()
        self.assertTrue(package_validator.validate(data, require_current_search_policy=True)['ready_for_delivery'])
        data['listings'][0]['search_terms'] = 'dekoartikel schmuckanhänger festbedarf'
        self.assertTrue(any('omits audited incremental tokens: wintermotiv' in e
                            for e in package_validator.validate(data)['errors']))

    def test_foreign_language_is_not_an_exclusion_reason(self):
        data = self.localized_search_package()
        audit = data['search_term_audits'][0]
        audit.update(decision='exclude', exclusion_reason='English term; no European local volume')
        del audit['localization_method']
        errors = package_validator.validate(data)['errors']
        self.assertTrue(any('must AI-localize' in e for e in errors))
        self.assertTrue(any('source language and absent local volume are not exclusion reasons' in e for e in errors))
        audit.update(localization_method='ai', decision='adopt', search_query=audit['source_phrase'])
        self.assertTrue(any('must search the localized phrase' in e
                            for e in package_validator.validate(data)['errors']))
        audit.update(search_query=audit['phrase'], phrase=audit['source_phrase'])
        self.assertTrue(any('unchanged foreign phrase requires' in e
                            for e in package_validator.validate(data)['errors']))

    def test_shortfall_needs_per_candidate_and_all_reference_review(self):
        data = self.localized_search_package()
        expansion = data['listings'][0]['search_term_expansion']
        del expansion['reviewed_phrases']
        self.assertTrue(any('shortfall requires reviewed_phrases' in e
                            for e in package_validator.validate(data)['errors']))
        expansion['reviewed_phrases'] = [a['phrase'] for a in data['search_term_audits']]
        data['title_candidate_audit']['all_user_references'].append({'asin': 'B087654321'})
        self.assertTrue(any('shortfall omits reference-title localization: B087654321' in e
                            for e in package_validator.validate(data)['errors']))
        audit = data['search_term_audits'][1]
        audit.update(decision='exclude', exclusion_reason='front_end_covered', exclusion_detail='Already in the title')
        self.assertTrue(any('front_end_covered exclusion has uncovered localized tokens: wintermotiv' in e
                            for e in package_validator.validate(data)['errors']))
        audit['source_phrase'] = 'Invented original phrase'
        self.assertTrue(any('original search-term source_phrase must exist' in e
                            for e in package_validator.validate(data)['errors']))

    def test_delivery_cannot_use_legacy_search_policy(self):
        data = self.audited_listing_package()
        data['search_term_policy_version'] = 2
        errors = package_validator.validate(data)['errors']
        self.assertTrue(any('Delivery requires search_term_policy_version=4' in e for e in errors))

    def test_specific_source_supported_long_tail_does_not_need_majority_results(self):
        data = self.localized_search_package()
        audit = data['search_term_audits'][1]
        audit.update(relevance_basis='source_product_match', relevance_fact_ids=['F1'],
                     source_evidence='reference title and product material record',
                     search_intent='same_product', relevant_results=3, relevance_band='low')
        self.assertTrue(package_validator.validate(data)['ready_for_delivery'])
        audit['relevance_fact_ids'] = ['invented']
        self.assertTrue(any('applicable confirmed' in e for e in package_validator.validate(data)['errors']))
        audit['relevance_fact_ids'] = ['F1']
        audit['search_intent'] = 'different_product'
        self.assertTrue(any('different-product search intent' in e for e in package_validator.validate(data)['errors']))
        audit.update(search_intent='not_measured', relevance_band='not_measured',
                     organic_results_checked=None, relevant_results=None)
        self.assertTrue(package_validator.validate(data)['ready_for_delivery'])
        audit['relevant_results'] = 20
        self.assertTrue(any('invented counts' in e for e in package_validator.validate(data)['errors']))

    def test_v4_search_concentration_is_not_an_adoption_gate(self):
        data = self.localized_search_package()
        audit = data['search_term_audits'][1]
        audit.update(relevant_results=1, relevance_band='low')
        self.assertTrue(package_validator.validate(data)['ready_for_delivery'])
        audit.update(organic_results_checked=3, relevant_results=1)
        self.assertTrue(package_validator.validate(data)['ready_for_delivery'])
        audit['search_intent'] = 'different_product'
        self.assertTrue(any('different-product search intent' in e
                            for e in package_validator.validate(data)['errors']))

    def test_research_inventory_and_prior_terms_cannot_be_silently_dropped(self):
        data = self.localized_search_package()
        data['search_query_inventory'] = [{'marketplace': 'DE', 'query': 'gedrehte Creolen'}]
        data['listings'][0]['previous_search_terms'] = 'zuckerstangen'
        data['search_term_concepts'].append({'id': 'candy', 'source_phrase': 'Candy', 'meaning': 'Candy theme'})
        errors = package_validator.validate(data)['errors']
        self.assertTrue(any('saved search query has no decision' in e for e in errors))
        self.assertTrue(any('prior search terms removed without review' in e for e in errors))
        self.assertTrue(any('omits source concepts: candy' in e for e in errors))

    def test_search_packing_uses_utf8_and_keeps_filling_after_large_candidate(self):
        with patch.dict(sys.modules, {'validate_listing_package': package_validator}):
            packer = module('pack_search_terms')
        result = packer.pack('FR', '<p>créoles de Noël</p>',
                             ['créoles sucre orge', 'créoles multicolores', 'créoles étoile', 'créoles orge'], limit=18)
        self.assertEqual(result['search_terms'], 'sucre orge étoile')
        self.assertEqual(result['utf8_bytes'], 18)
        self.assertEqual(result['characters'], 17)
        self.assertEqual(result['decisions'][0]['front_end_tokens'], ['créoles'])
        self.assertEqual(result['decisions'][0]['added_tokens'], ['sucre', 'orge'])
        self.assertEqual(result['decisions'][1]['exclusion_reason'], 'byte_limit')
        self.assertEqual(result['decisions'][3]['decision'], 'adopt')
        exact = packer.pack('FR', '', ['é' * 124, 'x'])
        self.assertEqual(exact['utf8_bytes'], 248)
        self.assertEqual(exact['decisions'][1]['exclusion_reason'], 'byte_limit')

    def test_keyword_report_separates_markets_and_full_phrases_from_tokens(self):
        data = {'targets': ['DE', 'FR', 'IT'], 'variants': [{'id': 'V1'}],
                'search_term_concepts': [{'id': 'hat'}],
                'listings': [
                    {'marketplace': 'DE', 'variant_id': 'V1', 'title': 'Weihnachtsmütze', 'search_terms': 'bommel'},
                    {'marketplace': 'FR', 'variant_id': 'V1', 'title': 'Bonnet de Noël',
                     'description': '<p>maille &amp; lumière</p>', 'search_terms': 'pompon'}],
                'search_term_audits': [
                    {'marketplace': 'DE', 'variant_id': 'V1', 'concept_id': 'hat',
                     'phrase': 'Weihnachtsmütze Bommel', 'decision': 'adopt'},
                    {'marketplace': 'FR', 'variant_id': 'V1', 'concept_id': 'hat',
                     'phrase': 'bonnet de Noël à pompon', 'decision': 'adopt'},
                    {'marketplace': 'FR', 'variant_id': 'V1', 'concept_id': 'hat',
                     'phrase': 'bonnet lumière', 'decision': 'exclude', 'exclusion_reason': 'front_end_covered'}]}
        rows = package_validator.keyword_coverage(data)
        fr = rows[1]
        self.assertEqual(fr['candidate_count'], 2)
        self.assertEqual(fr['incremental_pool_tokens'], 1)
        self.assertEqual(fr['candidates'][0]['backend_tokens'], ['pompon'])
        self.assertEqual(fr['candidates'][0]['placement'], 'across_fields')
        self.assertEqual(fr['candidates'][0]['exact_front_end_fields'], [])
        self.assertEqual(fr['candidates'][1]['front_end_tokens_by_field']['description'], ['lumière'])
        self.assertEqual(rows[2]['next_actions'], ['create_missing_listing'])
        self.assertIn('review_target_market_expansion', fr['next_actions'])
        self.assertNotIn('reconcile_backend_with_reviewed_candidates', fr['next_actions'])

    def test_keyword_report_detects_frontend_changes_and_unreviewed_loss(self):
        data = {'listings': [{'marketplace': 'FR', 'variant_id': 'A', 'title': 'bonnet',
                             'search_terms': 'pompon', 'previous_search_terms': 'pompon tricoté cadeau',
                             'previous_term_exclusions': [{'term': 'cadeau', 'reason': 'low_relevance',
                                                          'detail': 'Broad gift intent'}]}],
                'search_term_audits': [{'marketplace': 'FR', 'variant_id': 'A', 'decision': 'adopt',
                                       'phrase': 'bonnet tricoté pompon'}]}
        before = package_validator.keyword_coverage(data)[0]
        self.assertEqual(before['unplaced_adopted_tokens'], ['tricoté'])
        self.assertEqual(before['previous_terms']['removed_without_review'], ['tricoté'])
        self.assertEqual(before['previous_terms']['removed_with_reason'], ['cadeau'])
        data['listings'][0]['description'] = '<p>tricot&#233;</p>'
        after = package_validator.keyword_coverage(data)[0]
        self.assertEqual(after['unplaced_adopted_tokens'], [])
        self.assertEqual(after['previous_terms']['removed_without_review'], [])
        self.assertEqual(after['previous_terms']['covered_in_front_end'], ['tricoté'])

    def test_long_keyword_field_without_audits_is_not_delivery_ready(self):
        data = self.localized_search_package()
        data['competitors'] = []
        data['search_term_audits'] = []
        data['listings'][0]['search_terms'] = 'x' * 240
        result = package_validator.validate(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('without an adopted relevance audit' in error for error in result['errors']))
        self.assertTrue(result['keyword_coverage'][0]['target_reached'])
        self.assertIn('complete_localized_candidate_review', result['keyword_coverage'][0]['next_actions'])

    def test_short_keyword_field_can_pass_when_evidence_review_is_complete(self):
        result = package_validator.validate(self.localized_search_package())
        self.assertTrue(result['ready_for_delivery'])
        self.assertFalse(result['keyword_coverage'][0]['target_reached'])
        self.assertTrue(result['keyword_coverage'][0]['semantic_review_required'])

    def test_keyword_report_flags_false_frontend_and_byte_exclusions(self):
        data = {'listings': [{'marketplace': 'FR', 'variant_id': 'A', 'title': 'bonnet',
                             'search_terms': 'pompon'}],
                'search_term_audits': [
                    {'marketplace': 'FR', 'variant_id': 'A', 'phrase': 'bonnet tricoté',
                     'decision': 'exclude', 'exclusion_reason': 'front_end_covered'},
                    {'marketplace': 'FR', 'variant_id': 'A', 'phrase': 'bonnet maille',
                     'decision': 'exclude', 'exclusion_reason': 'byte_limit'},
                    {'marketplace': 'FR', 'variant_id': 'A', 'phrase': 'pochette cadeau',
                     'decision': 'exclude', 'exclusion_reason': 'product_mismatch'}]}
        row = package_validator.keyword_coverage(data)[0]
        self.assertEqual(row['incorrect_exclusions'], ['bonnet tricoté', 'bonnet maille'])
        self.assertIn('reconcile_backend_with_reviewed_candidates', row['next_actions'])
        self.assertEqual(row['incremental_pool_tokens'], 2)

    def test_keyword_report_counts_utf8_separately_and_keeps_variant_scope(self):
        data = {'listings': [{'marketplace': 'FR', 'variant_id': 'A', 'search_terms': 'écharpe'},
                             {'marketplace': 'FR', 'variant_id': 'B', 'search_terms': 'pompon'}],
                'search_term_audits': [{'marketplace': 'FR', 'variant_id': 'A',
                                       'phrase': 'écharpe', 'decision': 'adopt'}]}
        first, second = package_validator.keyword_coverage(data)
        self.assertEqual((first['characters'], first['utf8_bytes']), (7, 8))
        self.assertEqual(first['unaudited_backend_tokens'], [])
        self.assertEqual(second['unaudited_backend_tokens'], ['pompon'])

    def test_keyword_report_previous_package_does_not_mutate_current(self):
        with patch.dict(sys.modules, {'validate_listing_package': package_validator}):
            reviewer = module('review_keyword_coverage')
        data = {'listings': [{'marketplace': 'FR', 'variant_id': 'A', 'search_terms': 'pompon'},
                             {'marketplace': 'IT', 'variant_id': 'A', 'search_terms': 'berretto'}]}
        previous = {'listings': [{'marketplace': 'FR', 'variant_id': 'A', 'search_terms': 'tricoté'}]}
        result = reviewer.review(data, previous)
        self.assertEqual(result['market_variants'][0]['previous_terms']['removed_without_review'], ['tricoté'])
        self.assertEqual(result['unmatched_previous_listings'], [{'marketplace': 'IT', 'variant_id': 'A'}])
        self.assertNotIn('previous_search_terms', data['listings'][0])
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'coverage.csv'
            reviewer.write_csv(result, path)
            import csv
            with path.open(encoding='utf-8-sig', newline='') as stream:
                rows = list(csv.DictReader(stream))
            self.assertEqual([row['marketplace'] for row in rows], ['FR', 'IT'])

    def test_packer_uses_related_broad_only_after_precise_shortfall(self):
        with patch.dict(sys.modules, {'validate_listing_package': package_validator}):
            packer = module('pack_search_terms')
        result = packer.pack('FR', '', ['x' * 235], broad_candidates=['fête', 'hiver'])
        self.assertEqual(result['precise_utf8_bytes'], 235)
        self.assertEqual(result['utf8_bytes'], 241)
        self.assertEqual(result['broad_added_tokens'], ['fête'])
        self.assertEqual(result['decisions'][-1]['exclusion_reason'], 'fallback_not_needed')
        full = packer.pack('FR', '', ['x' * 240], broad_candidates=['fête'])
        self.assertEqual(full['broad_added_tokens'], [])
        self.assertEqual(full['utf8_bytes'], 240)
        small = packer.pack('FR', '', ['x' * 239], broad_candidates=['accessoires', 'été'])
        self.assertEqual(small['utf8_bytes'], 245)
        self.assertEqual(small['broad_added_tokens'], ['été'])

    def test_related_broad_requires_precise_research_and_product_match(self):
        data = self.localized_search_package()
        audit = data['search_term_audits'][1]
        audit['specificity'] = 'related_broad'
        self.assertTrue(package_validator.validate(data)['ready_for_delivery'])
        audit.update(relevance_basis='source_product_match', relevance_fact_ids=['F1'],
                     source_evidence='Recorded reference title and product facts',
                     search_intent='not_measured', relevance_band='not_measured',
                     organic_results_checked=None, relevant_results=None)
        self.assertTrue(package_validator.validate(data)['ready_for_delivery'])
        del data['listings'][0]['search_term_expansion']
        errors = package_validator.validate(data)['errors']
        self.assertTrue(any('related_broad requires actual precise expansion' in e for e in errors))
        audit['search_intent'] = 'different_product'
        self.assertTrue(any('different-product' in e for e in package_validator.validate(data)['errors']))

    def test_broad_fallback_cannot_displace_precise_terms_or_hide_research(self):
        listing = {'marketplace': 'FR', 'variant_id': 'V1', 'search_terms': 'x' * 230 + ' fête',
                   'search_term_expansion': {'queries': ['actual target-market query'],
                       'reviewed_phrases': ['x' * 230, 'bonnet'], 'shortfall_reason': 'Precise terms reviewed'}}
        audits = [{'phrase': 'x' * 230, 'specificity': 'product_specific', 'decision': 'adopt'},
                  {'phrase': 'bonnet', 'specificity': 'product_specific', 'decision': 'exclude',
                   'exclusion_reason': 'byte_limit'},
                  {'phrase': 'fête', 'specificity': 'related_broad', 'decision': 'adopt'}]
        errors = []
        package_validator.validate_broad_fallback(listing, audits, errors)
        self.assertTrue(any('must not displace a fitting precise' in e for e in errors))
        audits.pop(1)
        audits[0]['phrase'] = 'x' * 240
        listing['search_terms'] = 'x' * 240 + ' fête'
        listing['search_term_expansion']['reviewed_phrases'] = ['x' * 240]
        errors = []
        package_validator.validate_broad_fallback(listing, audits, errors)
        self.assertTrue(any('precise terms already reach 240' in e for e in errors))
        audits[0]['phrase'] = 'x' * 235
        listing['search_terms'] = 'x' * 235 + ' fête'
        del listing['search_term_expansion']
        errors = []
        package_validator.validate_broad_fallback(listing, audits, errors)
        self.assertTrue(any('requires actual precise expansion' in e for e in errors))

    def test_fallback_not_needed_cannot_discard_precise_or_underfilled_terms(self):
        audit = {'phrase': 'fête', 'specificity': 'related_broad', 'decision': 'exclude',
                 'exclusion_reason': 'fallback_not_needed'}
        listing = {'marketplace': 'FR', 'variant_id': 'V1', 'search_terms': 'bonnet'}
        errors = []
        package_validator.validate_broad_fallback(listing, [audit], errors)
        self.assertTrue(errors)
        listing['search_terms'] = 'x' * 240
        errors = []
        package_validator.validate_broad_fallback(listing, [audit], errors)
        self.assertFalse(errors)
        audit['specificity'] = 'product_specific'
        package_validator.validate_broad_fallback(listing, [audit], errors)
        self.assertTrue(errors)

    def test_query_inventory_collects_queries_missing_from_final_review(self):
        collector = module('collect_search_query_inventory')
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'search-extra.json').write_text(
                '{"ES2":{"url":"https://www.amazon.es/s?k=aros","query":"aros", "results":[]},'
                '"FR":{"query":"créoles rondes","results":[]}}')
            rows = collector.collect(root)
            self.assertEqual({(r['marketplace'], r['query']) for r in rows},
                             {('ES', 'aros'), ('FR', 'créoles rondes')})

    def test_search_terms_v2_requires_expansion_before_accepting_short_field(self):
        data = self.audited_listing_package()
        data['search_term_policy_version'] = 2
        listing = data['listings'][0]
        listing['search_terms'] += ' wintermotiv'
        for audit in data['search_term_audits']:
            audit.update(specificity='product_specific',
                         product_match_reason='The full phrase matches the product type and intended use.')
        result = validate_legacy(data)
        self.assertTrue(any('below 240 bytes requires' in e for e in result['errors']))
        for invalid in (None, [], {'queries': [''], 'shortfall_reason': 'short'},
                        {'queries': ['actual query'], 'shortfall_reason': ''}):
            listing['search_term_expansion'] = invalid
            self.assertFalse(validate_legacy(data)['ready_for_delivery'])
        listing['search_term_expansion'] = {
            'queries': ['DE reference titles and additional paper ornament reverse-ASIN candidates'],
            'shortfall_reason': 'Other audited terms overlap front-end copy or describe different structures.',
        }
        self.assertTrue(validate_legacy(data)['ready_for_delivery'])
        audit = data['search_term_audits'][0]
        audit.update(relevant_results=12, relevance_band='medium')
        self.assertTrue(any('only high-relevance' in e
                            for e in validate_legacy(data)['errors']))
        audit.update(relevant_results=16, relevance_band='high', specificity='generic')
        self.assertTrue(any('not generic' in e
                            for e in validate_legacy(data)['errors']))
        audit.update(specificity='product_specific', product_match_reason='')
        self.assertTrue(any('product_match_reason' in e
                            for e in validate_legacy(data)['errors']))

    def test_search_terms_v2_byte_limits_count_utf8(self):
        data = self.listing_package()
        data['search_term_policy_version'] = 2
        listing = data['listings'][0]
        # Synthetic terms isolate the byte boundaries from semantic review.
        for byte_count in (239, 240, 249, 250):
            listing['search_terms'] = 'é' * (byte_count // 2) + ('x' if byte_count % 2 else '')
            errors = validate_legacy(data)['errors']
            self.assertEqual(any('below 240 bytes requires' in e for e in errors), byte_count < 240)
            self.assertEqual(any('maximum is 249' in e for e in errors), byte_count > 249)
        data['search_term_policy_version'] = True
        self.assertTrue(any('search_term_policy_version must be' in e
                            for e in validate_legacy(data)['errors']))

    def test_search_terms_v2_rejects_false_byte_limit_exclusion(self):
        data = self.audited_listing_package()
        data['search_term_policy_version'] = 2
        data['listings'][0]['search_term_expansion'] = {
            'queries': ['DE reference and reverse-ASIN expansion'],
            'shortfall_reason': 'Remaining phrases were checked for product match and front-end overlap.',
        }
        for audit in data['search_term_audits']:
            audit.update(specificity='product_specific', product_match_reason='Matches product structure.')
        last = data['search_term_audits'][-1]
        last.update(decision='exclude', exclusion_reason='byte_limit')
        self.assertTrue(any('byte_limit exclusion still fits' in e
                            for e in validate_legacy(data)['errors']))
        last['phrase'] = 'é' * 130
        self.assertTrue(validate_legacy(data)['ready_for_delivery'])
        last.update(phrase='dekoration', exclusion_reason='front_end_covered')
        self.assertTrue(validate_legacy(data)['ready_for_delivery'])

    def test_short_search_terms_prompts_more_research_without_inventing_words(self):
        data = copy.deepcopy(self.listing_package())
        data['listings'][0]['search_terms'] = 'dekoartikel festbedarf'
        result = validate_legacy(data)
        self.assertTrue(result['ready_for_delivery'])
        self.assertTrue(any('search_terms has only 2 distinct tokens'
                            in warning for warning in result['warnings']))

    def test_own_listing_package_rejects_missing_or_competitor_facts(self):
        data = copy.deepcopy(self.listing_package())
        data['facts'][0]['source_type'] = 'competitor'
        data['facts'][0]['status'] = 'unconfirmed'
        data['listings'][0]['bullets'] = ['Only one']
        data['listings'][0]['description'] = 'Eigenschaften:\n1. A: B\nProduktdetails:\nLieferumfang:'
        data['listings'][0]['title'] += ' B012345678'
        result = validate_legacy(data)
        self.assertFalse(result['ready_for_delivery'])
        self.assertTrue(any('source_type must be own_product, same_product_evidence, or material_match_evidence' in error
                            for error in result['errors']))
        self.assertTrue(any('five non-empty bullets' in error for error in result['errors']))
        self.assertTrue(any('uses non-confirmed fact' in error for error in result['errors']))
        self.assertTrue(any('contains an ASIN' in error for error in result['errors']))


if __name__ == '__main__':
    unittest.main()


class ClauseSourceTests(unittest.TestCase):
    def setUp(self):
        self.competitors = [
            {'asin': 'B012345678', 'title': 'Fine chain',
             'bullets': ['Three white beads', 'Fine chain'],
             'description': 'Wear with a shirt.'},
            {'asin': 'B087654321', 'bullets': ['Central bead is larger.']},
        ]
        self.facts = {'F1': {'status': 'confirmed', 'variant_ids': ['V1']}}
        self.listing = {
            'bullets': ['【Trois perles】Trois perles blanches. Centre plus grand.'] * 5,
            'description': '<p>Avec une chemise.</p><p>Longueur : 40 cm.</p>',
            'field_fact_ids': {**{f'bullet_{i}': ['F1'] for i in range(1, 6)}, 'description': ['F1']},
            'clause_sources': {}, 'bullet_references': [],
        }
        for i in range(1, 6):
            self.listing['clause_sources'][f'bullet_{i}'] = [
                {'text': 'Trois perles', 'sources': [{'asin': 'B012345678', 'source_field': 'bullet_1', 'source_phrase': 'Three white beads'}]},
                {'text': 'Trois perles blanches.', 'sources': [{'asin': 'B012345678', 'source_field': 'bullet_1', 'source_phrase': 'Three white beads'}]},
                {'text': 'Centre plus grand.', 'sources': [{'asin': 'B087654321', 'source_field': 'bullet_1', 'source_phrase': 'Central bead is larger.'}]},
            ]
            self.listing['bullet_references'].append(
                'Trois perles ← B012345678 第1点（Three white beads）；'
                'Trois perles blanches. ← B012345678 第1点（Three white beads）；'
                'Centre plus grand. ← B087654321 第1点（Central bead is larger.）')
        self.listing['clause_sources']['description'] = [
            {'text': 'Avec une chemise.', 'sources': [{'asin': 'B012345678', 'source_field': 'description', 'source_phrase': 'Wear with a shirt.'}]},
            {'text': 'Longueur : 40 cm.', 'sources': [{'fact_id': 'F1'}]},
        ]
        self.listing['description_reference'] = 'Avec une chemise. ← B012345678 详情（Wear with a shirt.）；Longueur : 40 cm. ← 自有事实 F1'

    def errors(self):
        errors = []
        package_validator.validate_clause_sources(errors, 'FR', 'V1', self.listing, self.competitors, self.facts)
        return errors

    def test_combines_sources_per_clause_and_handles_html(self):
        self.assertEqual(self.errors(), [])

    def test_title_rejects_real_bullet_source_but_highlights_accepts_it(self):
        for field in ('title', 'item_highlights'):
            listing = {
                field: 'Fine chain',
                f'{field}_reference': 'Fine chain ← B012345678 第2点（Fine chain）',
                f'{field}_phrase_sources': [
                    {'phrase': 'Fine chain', 'asin': 'B012345678',
                     'source_field': 'bullet_2', 'source_phrase': 'Fine chain'},
                ],
            }
            errors = []
            package_validator.validate_phrase_sources(
                errors, 'FR', 'V1', listing, field, self.competitors, self.facts, strict=True)
            if field == 'title':
                self.assertTrue(any('competitor title sources only' in error for error in errors))
                listing['title_phrase_sources'][0]['source_field'] = 'title'
                listing['title_reference'] = 'Fine chain ← B012345678 标题（Fine chain）'
                errors = []
                package_validator.validate_phrase_sources(
                    errors, 'FR', 'V1', listing, field, self.competitors, self.facts, strict=True)
            self.assertEqual(errors, [])

    def test_bullet_rejects_cross_field_source_even_with_valid_bullet_sources(self):
        for source_field, label in (('title', '标题'), ('description', '详情'), ('aplus_text', 'A+')):
            with self.subTest(source_field=source_field):
                self.setUp()
                self.competitors[0][source_field] = 'Three white beads'
                source = self.listing['clause_sources']['bullet_1'][0]['sources'][0]
                source['source_field'] = source_field
                self.listing['bullet_references'][0] = self.listing['bullet_references'][0].replace(
                    'Trois perles ← B012345678 第1点', f'Trois perles ← B012345678 {label}')
                self.assertTrue(any('competitor bullet sources only' in error for error in self.errors()))

    def test_omitted_detail_cannot_hide_behind_bullet_reference(self):
        self.listing['clause_sources']['bullet_1'].pop()
        self.assertTrue(any('without source coverage' in error for error in self.errors()))

    def test_wrong_point_and_invented_original_are_rejected(self):
        source = self.listing['clause_sources']['bullet_1'][0]['sources'][0]
        source['source_field'] = 'bullet_2'
        self.assertTrue(any('source_phrase not found' in error for error in self.errors()))

    def test_visible_reference_must_bind_the_actual_clause(self):
        self.listing['bullet_references'][0] = 'B012345678 第1点；B087654321 第1点'
        self.assertTrue(any('visible reference must map clause' in error for error in self.errors()))

    def test_wrong_variant_fact_is_rejected(self):
        self.facts['F1']['variant_ids'] = ['V2']
        self.assertTrue(any('applicable field fact' in error for error in self.errors()))

    def test_numbered_reference_without_actual_bullet_mapping_is_rejected(self):
        for entry in self.listing['clause_sources']['bullet_1']:
            entry['sources'] = [{'fact_id': 'F1'}]
        self.assertTrue(any('actual competitor bullet source' in error for error in self.errors()))

    def test_v2_full_package_requires_clause_sources(self):
        data = ListingTests.listing_package()
        data['source_mapping_version'] = 2
        self.assertTrue(any('clause_sources must map' in error for error in validate_legacy(data)['errors']))

    def test_v2_full_package_with_complete_sources(self):
        import re
        data = ListingTests.listing_package()
        data['source_mapping_version'] = 2
        listing = data['listings'][0]
        listing['clause_sources'] = {}
        listing['bullet_references'] = []
        # This fixture tests structural traceability; factual entailment remains a semantic review.
        for field in ('title', 'item_highlights'):
            listing[f'{field}_phrase_sources'] = [{'phrase': listing[field], 'fact_id': 'F1'}]
            listing[f'{field}_reference'] = listing[field] + ' ← 自有事实 F1'
        for field in (*(f'bullet_{i}' for i in range(1, 6)), 'description'):
            text = listing['description'] if field == 'description' else listing['bullets'][int(field[-1]) - 1]
            clauses = [x.strip() for x in re.split(r'(?<=[.!?])\s+|<[^>]+>|[【】]', text) if x.strip()]
            listing['clause_sources'][field] = [{'text': x, 'sources': [{'fact_id': 'F1'}]} for x in clauses]
            reference = '\n'.join(x + ' ← 自有事实 F1' for x in clauses)
            if field == 'description':
                listing['description_reference'] = reference
            else:
                listing['bullet_references'].append(reference)
        result = validate_legacy(data)
        self.assertTrue(result['ready_for_delivery'], result['errors'])
        listing['title'] += ' extra'
        self.assertTrue(any('without source coverage' in error for error in validate_legacy(data)['errors']))


    def test_description_cannot_use_only_own_facts_when_references_exist(self):
        for entry in self.listing['clause_sources']['description']:
            entry['sources'] = [{'fact_id': 'F1'}]
        self.listing['description_reference'] = '\n'.join(
            entry['text'] + ' ← 自有事实 F1'
            for entry in self.listing['clause_sources']['description'])
        self.assertTrue(any('needs actual competitor bullet/description/A+ content'
                            in error for error in self.errors()))

    def test_fact_checks_do_not_need_duplicate_content_source_entries(self):
        self.assertTrue(self.listing['field_fact_ids']['description'])
        self.assertEqual(self.errors(), [])
        self.assertNotIn('fact_id', self.listing['clause_sources']['description'][0]['sources'][0])

    def test_description_can_draw_from_actual_competitor_bullet(self):
        entry = self.listing['clause_sources']['description'][0]
        entry['text'] = 'Trois perles blanches.'
        entry['sources'] = [{'asin': 'B012345678', 'source_field': 'bullet_1',
                             'source_phrase': 'Three white beads'}]
        self.listing['description'] = '<p>Trois perles blanches.</p><p>Longueur : 40 cm.</p>'
        self.listing['description_reference'] = 'Trois perles blanches. ← B012345678 第1点（Three white beads）；Longueur : 40 cm. ← 自有事实 F1'
        self.assertEqual(self.errors(), [])
