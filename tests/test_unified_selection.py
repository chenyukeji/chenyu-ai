import json
from pathlib import Path
import sqlite3
import sys

import pytest

SCRIPTS=Path(__file__).resolve().parents[1]/'plugins/chenyu-kaifa/skills/chenyu-xuanpin/scripts'
sys.path.insert(0,str(SCRIPTS))
import selection_flow as flow
from selection_analysis import analyze_keywords, merge_products, summarize_products
from selection_collectors import amazon_identity, collect_database


def scope(sources=None):
    return {'marketplaces':['US'],'sources':sources or ['sellersprite_new','keyword_trends'],
            'categories':[{'name':'圣诞树装饰','keywords':{'US':['christmas tree ornaments']}}]}


def keyword(month,volume,**extra):
    return {'marketplace':'US','keyword':'tree ornaments','month':month,'search_volume':volume,
            'source_ref':'https://www.sellersprite.com/fixture',**extra}


def product(**extra):
    return {'marketplace':'US','asin':'B012345678','product_name':'Tree ornaments',
            'source_available_date':'2026-09-01','estimated_sales':500,'review_count':8,'price':15,
            'source_ref':'https://www.sellersprite.com/fixture','parent_asin':extra.get('asin','B012345678'),**extra}


def test_input_market_identity_and_unsupported_domains():
    assert amazon_identity('https://www.amazon.de/dp/B012345678')['marketplace']=='DE'
    assert amazon_identity('https://www.amazon.com/sp?seller=A123456789')['kind']=='store'
    assert amazon_identity('https://www.amazon.com/stores/ABC/page/123')['seller_id'] is None
    for bad in ('https://amazon.com.evil.test/dp/B012345678','https://www.amazon.fr/dp/B012345678','https://user@amazon.com/dp/B012345678'):
        with pytest.raises(ValueError):amazon_identity(bad)


def test_product_link_does_not_default_to_two_markets():
    plan=flow.plan_selection({'request':'分析 https://www.amazon.de/dp/B012345678'})
    assert plan['marketplaces']==['DE'] and plan['blocking_items']
    assert plan['queries'][0]['source']=='amazon_input'


def test_multi_source_plan_keeps_history_windows_and_explicit_scope():
    value=scope(list(flow.SOURCES));value['following_months']=3
    plan=flow.plan_selection({'request':'圣诞树装饰','scope':value,'as_of_date':'2026-10-11'})
    historical=[q for q in plan['queries'] if q['source']=='sellersprite_history']
    assert [q['history_month'] for q in historical]==['2025-10','2025-11','2025-12','2026-01']
    assert all(q['keyword']=='christmas tree ornaments' for q in historical)
    assert any(q.get('unavailable') for q in plan['queries'] if q['source']=='amazon_new_releases')
    assert len({q['id'] for q in plan['queries']})==len(plan['queries'])


def test_keyword_only_evidence_survives_without_asin_and_excludes_partial_month():
    report=analyze_keywords([keyword('2025-09',100),keyword('2026-08',200),keyword('2026-09',300),
                             keyword('2026-10',10000)],'2026-10-11')
    d=report['directions'][0]
    assert d['latest_month']=='2026-09' and d['mom_percent']==50 and d['yoy_percent']==200
    assert d['absolute_change']==100 and len(d['history'])==3


def test_keyword_missing_month_and_zero_base_never_invent_growth():
    d=analyze_keywords([keyword('2026-07',100),keyword('2026-09',200)],'2026-10-11')['directions'][0]
    assert d['mom_percent'] is None and d['yoy_percent'] is None
    d=analyze_keywords([keyword('2026-08',0),keyword('2026-09',200)],'2026-10-11')['directions'][0]
    assert d['mom_percent'] is None and d['absolute_change']==200


def test_conflicting_keyword_observations_require_review():
    d=analyze_keywords([keyword('2026-08',100),keyword('2026-09',200),keyword('2026-09',300)],'2026-10-11')['directions'][0]
    assert d['confidence']=='待核实' and d['mom_percent'] is None


def test_history_and_database_do_not_overwrite_current_sales():
    values=[{'query':{'source':'sellersprite_history','history_month':'2025-12'},'records':[product(estimated_sales=9999)]},
            {'query':{'source':'sellersprite_new'},'records':[product()]},
            {'query':{'source':'new_releases_db'},'records':[product(estimated_sales=0)]}]
    products=merge_products(values)
    assert len(products)==1 and products[0]['current']['estimated_sales']==500
    assert products[0]['historical'][0]['estimated_sales']==9999
    assert len(products[0]['source_types'])==3


def test_missing_data_cannot_receive_positive_priority():
    products=merge_products([{'query':{'source':'sellersprite_new'},'records':[product(price=None)]}])
    rows=summarize_products(products,[{'marketplace':'US','asin':'B012345678','relevant':True,'reason':'相关','priority':'优先研究'}])
    assert rows[0]['优先级']=='待补数据'


def test_resume_reuses_success_and_retries_blocked_source(monkeypatch,tmp_path):
    calls=[]
    def collect(q,p,known):
        calls.append(q['source'])
        if q['source']=='keyword_trends':return {'collection_status':'blocked','records':[],'block_reason':'permission'}
        return {'collection_status':'complete','records':[product()]}
    monkeypatch.setattr(flow,'collect_query',collect)
    payload={'request':'圣诞树装饰','scope':scope(),'run_dir':str(tmp_path),'enrich':False,'as_of_date':'2026-10-11'}
    result=flow.run_selection(payload)
    assert result['status']=='AWAITING_ANALYSIS' and len(result['analysis_requests'])==1
    assert not (tmp_path/'开品结果.xlsx').exists()
    result=flow.run_selection({'run_dir':str(tmp_path),'enrich':False})
    assert calls==['sellersprite_new','keyword_trends','keyword_trends']
    assert result['coverage'][1]['status']=='blocked'


def test_query_budget_resumes_without_recollecting_success(monkeypatch,tmp_path):
    calls=[]
    def collect(q,p,known):
        calls.append(q['source']);return {'collection_status':'complete','records':[product()] if q['source']=='sellersprite_new' else [keyword('2026-09',300)]}
    monkeypatch.setattr(flow,'collect_query',collect)
    p={'scope':scope(),'request':'圣诞树装饰','run_dir':str(tmp_path),'enrich':False,'max_queries':1,'as_of_date':'2026-10-11'}
    assert flow.run_selection(p)['status']=='RUNNING'
    result=flow.run_selection({'run_dir':str(tmp_path),'enrich':False,'max_queries':1})
    assert result['status']=='AWAITING_ANALYSIS' and calls==['sellersprite_new','keyword_trends']



def test_default_combined_selection_does_not_query_last_year():
    value = {'marketplaces': ['US'], 'categories': [
        {'name': '圣诞树装饰', 'keywords': {'US': ['christmas tree ornaments']}}
    ]}
    plan = flow.plan_selection({'request': '分析圣诞树装饰', 'scope': value,
                                'as_of_date': '2026-10-11'})
    assert 'sellersprite_history' not in {query['source'] for query in plan['queries']}
    assert 'sellersprite_surge' in {query['source'] for query in plan['queries']}



def test_keyword_only_evidence_stays_in_json_and_workbook_has_one_sheet(monkeypatch,tmp_path):
    monkeypatch.setattr(flow,'collect_query',lambda *args:{'collection_status':'complete','records':[keyword('2026-08',100),keyword('2026-09',300)]})
    result=flow.run_selection({'scope':scope(['keyword_trends']),'run_dir':str(tmp_path),'as_of_date':'2026-10-11'})
    from openpyxl import load_workbook
    wb=load_workbook(result['workbook']['path'])
    evidence=json.loads((tmp_path/'selection-evidence.json').read_text())
    assert result['status']=='PARTIAL' and result['counts']['shortfall']==50 and wb.sheetnames==['开品结果']
    assert wb.active.max_row==1
    assert evidence['keywords']['directions'][0]['absolute_change']==200
    assert len(evidence['keywords']['directions'][0]['history'])==2


def test_database_scope_and_readonly_window(tmp_path):
    path=tmp_path/'db.sqlite';con=sqlite3.connect(path)
    con.execute('CREATE TABLE observations(marketplace,category,snapshot_date,rank,asin,title,source_url)')
    con.executemany('INSERT INTO observations VALUES (?,?,?,?,?,?,?)',[
        ('US','party','2026-10-10',1,'B012345678','Party','https://www.amazon.com/gp/new-releases/toys'),
        ('US','other','2026-10-10',1,'B012345679','Other','https://www.amazon.com/gp/new-releases/home')])
    con.commit();con.close();before=path.read_bytes()
    result=collect_database({'db_path':str(path),'marketplace':'US','db_category':'party','as_of_date':'2026-10-11'})
    assert len(result['records'])==1 and path.read_bytes()==before
    with pytest.raises(ValueError):collect_database({'db_path':str(path),'marketplace':'US','as_of_date':'2026-10-11'})


def test_changed_scope_does_not_reuse_reviews(monkeypatch,tmp_path):
    monkeypatch.setattr(flow,'collect_query',lambda *args:{'collection_status':'complete','records':[product()]})
    p={'scope':scope(['sellersprite_new']),'run_dir':str(tmp_path),'as_of_date':'2026-10-11','enrich':False,
       'reviews':[{'marketplace':'US','asin':'B012345678','relevant':True,'priority':'继续观察','reason':'本次商品用途与圣诞装饰相关'}]}
    assert flow.run_selection(p)['status']=='PARTIAL'  # Sales/review analysis still missing.
    p.pop('reviews');p['scope']['categories'][0]['keywords']['US']=['birthday decorations']
    assert flow.run_selection(p)['status']=='AWAITING_ANALYSIS'


def test_filters_are_enforced_and_formula_text_is_not_executable(monkeypatch,tmp_path):
    monkeypatch.setattr(flow,'collect_query',lambda *args:{'collection_status':'complete','records':[product(product_name='=1+1',price=15)]})
    value=scope(['sellersprite_new']);value['filters']={'price_max':10}
    result=flow.run_selection({'scope':value,'run_dir':str(tmp_path),'enrich':False,
       'reviews':[{'marketplace':'US','asin':'B012345678','relevant':True,'priority':'优先研究','reason':'商品相关'}]})
    assert result['workbook']['row_count']==0
    value['filters']={'price_max':20}
    result=flow.run_selection({'scope':value,'run_dir':str(tmp_path),'enrich':False,
       'reviews':[{'marketplace':'US','asin':'B012345678','relevant':True,'priority':'优先研究','reason':'=1+1'}]})
    from openpyxl import load_workbook
    ws=load_workbook(result['workbook']['path'])['开品结果']
    assert ws.cell(2,11).value=='B012345678'
    assert ws.cell(2,15).value.startswith('=1+1') and ws.cell(2,15).data_type=='s'



def test_unified_workbook_restores_original_main_table(monkeypatch, tmp_path):
    from openpyxl import load_workbook
    from discovery_workbook import HEADERS
    monkeypatch.setattr(flow, 'collect_query', lambda *args: {
        'collection_status': 'complete', 'records': [product(bsr=47302, category_name='圣诞树装饰')]
    })
    payload = {'scope': scope(['sellersprite_new']), 'run_dir': str(tmp_path),
               'enrich': False, 'reviews': [{'marketplace': 'US', 'asin': 'B012345678',
               'relevant': True, 'priority': '优先研究',
               'reason': '圣诞树挂饰具有手作特点。销量口径需要核实。'}]}
    result = flow.run_selection(payload)
    wb = load_workbook(result['workbook']['path'])
    assert wb.sheetnames == ['开品结果']
    ws = wb.active
    assert [cell.value for cell in ws[1]] == HEADERS
    assert ws.cell(2, 5).value == 47302
    assert ws.cell(2, 8).value == '圣诞树挂饰具有手作特点'
    assert ws.cell(2, 9).value is None
    assert ws.cell(2, 10).value is None
    assert ws.cell(2, 14).value == '🟢 优先研究'
    assert ws.cell(2, 15).value.startswith('圣诞树挂饰')
    assert ws.cell(1, 1).fill.fgColor.rgb == '00C6E0B4'
    assert ws.cell(2, 14).fill.fgColor.rgb == '0033A36B'
    assert ws.freeze_panes == 'A2' and ws.auto_filter.ref == 'A1:O2'
    wb.close()

def test_fresh_response_matching_rejects_other_market_keyword_and_host():
    from types import SimpleNamespace
    from selection_collectors import response_matches
    def response(market='US',keyword='party',host='www.sellersprite.com'):
        url=f'https://{host}/v3/api/product-research'
        return SimpleNamespace(url=url,request=SimpleNamespace(url=url,post_data=json.dumps({'market':market,'keywords':keyword})))
    args={'path':'/v3/api/product-research','expected':{'market':'US','keywords':'party'}}
    assert response_matches(response(),**args)
    assert not response_matches(response(market='DE'),**args)
    assert not response_matches(response(keyword='tree'),**args)
    assert not response_matches(response(host='www.sellersprite.com.evil.test'),**args)


def test_keyword_series_uses_explicit_month_axis():
    from selection_collectors import parse_keyword_series, browser
    rows=parse_keyword_series([{'key':'time','data':['2025年12月','2026年01月']},
                               {'key':'search','data':[100,200]}, {'key':'purchase','data':[2,3]}],
                              'US','tree ornaments','https://www.sellersprite.com/trend')
    assert [r['month'] for r in rows]==['2025-12','2026-01']
    assert rows[1]['purchase_volume']==3
    with pytest.raises(browser.BrowserCollectionError):
        parse_keyword_series([{'key':'time','data':['unknown']},{'key':'search','data':[10]}],
                             'US','tree ornaments','https://www.sellersprite.com/trend')


def test_shortlist_prefers_lexical_fit_but_balances_sources():
    from selection_analysis import balanced_shortlist
    data=[{'query':{'source':'sellersprite_new'},'records':[
        product(asin='B012345678',product_name='Tea tree wipes',filter_keyword='christmas tree ornaments'),
        product(asin='B012345679',product_name='Christmas tree ornaments',filter_keyword='christmas tree ornaments')]},
        {'query':{'source':'amazon_new_releases'},'records':[product(asin='B012345680')]}]
    rows=balanced_shortlist(merge_products(data),2)
    assert [r['asin'] for r in rows]==['B012345679','B012345680']


def test_invalid_scope_values_are_rejected():
    for extra in ({'input_urls':'https://www.amazon.com/dp/B012345678'},
                  {'categories':['tree']},{'filters':{'price_min':float('nan')}}):
        with pytest.raises(ValueError):flow.plan_selection({'scope':{**scope(),**extra}})


def test_placeholder_dates_and_numeric_fields_remain_missing():
    rows=merge_products([{'query':{'source':'sellersprite_new'},'records':[
        product(source_available_date='--',estimated_sales='--',price='15.5')]}])
    assert rows[0]['current'].get('source_available_date') is None
    assert rows[0]['current'].get('estimated_sales') is None
    assert rows[0]['current']['price']==15.5


def test_duplicate_observation_cannot_erase_an_existing_conflict():
    d=analyze_keywords([keyword('2026-09',200),keyword('2026-09',300),keyword('2026-09',200)],'2026-10-11')['directions'][0]
    assert d['confidence']=='待核实'


def test_authentication_failure_stops_repeated_source_logins(monkeypatch,tmp_path):
    calls=[]
    def collect(q,*args):
        calls.append(q['source'])
        return {'collection_status':'blocked','records':[],'block_reason':'sellersprite_login_required'}
    monkeypatch.setattr(flow,'collect_query',collect)
    r=flow.run_selection({'scope':scope(),'run_dir':str(tmp_path)})
    assert calls==['sellersprite_new'] and r['status']=='BLOCKED'
    assert len(r['coverage'])==2 and all(x['reason']=='sellersprite_login_required' for x in r['coverage'])


def review(item, relevant=True):
    return {'marketplace': item['marketplace'], 'asin': item['asin'], 'relevant': relevant,
            'priority': '继续观察', 'reason': '按该商品的用途及历史证据审核'}


@pytest.mark.parametrize('unlimited', [False, True])
def test_default_fifty_or_explicit_all_in_resumable_batches(monkeypatch, tmp_path, unlimited):
    calls = []
    rows = [product(asin=f'B{i:09d}') for i in range(125)]
    def collect(*args):
        calls.append(1)
        return {'collection_status': 'complete', 'records': rows}
    monkeypatch.setattr(flow, 'collect_query', collect)
    value = scope(['sellersprite_new'])
    if unlimited:
        value['shortlist_limit'] = None
    r = flow.run_selection({'scope': value, 'run_dir': str(tmp_path), 'enrich': False})
    assert r['counts']['families'] == 125 and len(r['analysis_requests']) == 50
    expected_batches = [(50, 50), (100, 25), (125, 0)] if unlimited else [(50, 0)]
    for expected_reviewed, expected_next in expected_batches:
        r = flow.run_selection({'run_dir': str(tmp_path), 'evidence_id': r['evidence_id'],
                                'reviews': [review(p) for p in r['analysis_requests']]})
        assert r['counts']['reviewed'] == expected_reviewed
        assert len(r['analysis_requests']) == expected_next
    assert r['status'] == 'PARTIAL'  # Sales/review analysis still missing.
    assert r['analysis_gaps'] and r['workbook']['row_count'] == (125 if unlimited else 50)
    assert len(calls) == 1


def test_explicit_target_replenishes_rejected_candidates(monkeypatch, tmp_path):
    monkeypatch.setattr(flow, 'collect_query', lambda *args: {
        'collection_status': 'complete', 'records': [product(asin=f'B{i:09d}') for i in range(8)]})
    value = {**scope(['sellersprite_new']), 'shortlist_limit': 3, 'review_batch_size': 2}
    r = flow.run_selection({'scope': value, 'run_dir': str(tmp_path), 'enrich': False})
    r = flow.run_selection({'run_dir': str(tmp_path), 'reviews': [review(p, False) for p in r['analysis_requests']]})
    assert r['status'] == 'AWAITING_ANALYSIS' and len(r['analysis_requests']) == 2
    r = flow.run_selection({'run_dir': str(tmp_path), 'reviews': [review(p) for p in r['analysis_requests']]})
    assert len(r['analysis_requests']) == 1
    r = flow.run_selection({'run_dir': str(tmp_path), 'reviews': [review(p) for p in r['analysis_requests']]})
    assert r['workbook']['row_count'] == 3 and r['counts']['reviewed'] == 5
    assert r['counts']['remaining'] == 3


def test_family_grouping_preserves_month_and_variant_evidence_without_summing_sales():
    from selection_analysis import group_product_families
    rows = merge_products([
        {'query': {'source': 'sellersprite_history', 'history_month': '2025-11'},
         'records': [product(asin='B000000001', parent_asin='B000000010', estimated_sales=800)]},
        {'query': {'source': 'sellersprite_history', 'history_month': '2025-12'},
         'records': [product(asin='B000000002', parent_asin='B000000010', estimated_sales=900),
                     product(asin='B000000003', parent_asin='B000000010', estimated_sales=900),
                     product(asin='B000000004', parent_asin=None), product(asin='B000000005', parent_asin=None),
                     product(asin='B000000002', parent_asin='B000000010', marketplace='DE')]}])
    families = group_product_families(rows)
    assert len(families) == 4  # US parent, DE parent, two unknown-parent ASINs.
    family = next(p for p in families if p['marketplace'] == 'US' and p['family_key'] == 'B000000010')
    assert family['asin'] == 'B000000002'
    assert family['family_asins'] == ['B000000001', 'B000000002', 'B000000003']
    assert {r['history_month'] for r in family['family_historical']} == {'2025-11', '2025-12'}
    assert family['historical'][0]['estimated_sales'] == 900
    assert family['current'] == {}  # Historical numbers never become current sales.
    assert all('family_asins' not in p for p in rows)  # Raw ASIN evidence stays intact.


def test_historical_shortlist_balances_keywords_before_unrelated_current_hits():
    from selection_analysis import group_product_families, balanced_shortlist
    sources = [{'query': {'source': 'amazon_new_releases'}, 'records': [product(asin='B000000999')]}]
    for index, term in enumerate(['heater', 'blanket', 'shovel', 'gloves']):
        sources.append({'query': {'source': 'sellersprite_history', 'keyword': term,
                                   'history_month': '2025-12'}, 'records': [
            product(asin=f'B{index * 20 + n:09d}', product_name=term, filter_keyword=term,
                    parent_asin=f'B{index * 20 + n + 500:09d}') for n in range(20)]})
    ordered = balanced_shortlist(group_product_families(merge_products(sources)),
                                 historical_months=['2025-12'])
    assert [p['historical'][0]['query_keyword'] for p in ordered[:4]] == ['heater', 'blanket', 'shovel', 'gloves']
    assert ordered[-1]['asin'] == 'B000000999'
    assert len(ordered) == 81


def test_family_deduplication_happens_before_output_target(monkeypatch, tmp_path):
    monkeypatch.setattr(flow, 'collect_query', lambda *args: {'collection_status': 'complete', 'records': [
        product(asin='B000000001', parent_asin='B000000010'),
        product(asin='B000000002', parent_asin='B000000010'),
        product(asin='B000000003', parent_asin='B000000020')]})
    r = flow.run_selection({'scope': {**scope(['sellersprite_new']), 'shortlist_limit': 2},
                            'run_dir': str(tmp_path), 'enrich': False})
    assert r['counts']['products'] == 3 and r['counts']['families'] == 2
    r = flow.run_selection({'run_dir': str(tmp_path), 'reviews': [review(p) for p in r['analysis_requests']]})
    assert r['workbook']['row_count'] == 2


def test_enrichment_advances_with_review_batches_without_requerying_previous_batch(monkeypatch, tmp_path):
    rows = [product(asin=f'B{i:09d}', price=None) for i in range(5)]
    monkeypatch.setattr(flow, 'collect_query', lambda *args: {'collection_status': 'complete', 'records': rows})
    calls = []
    def enrich(options):
        calls.append(options['asins'])
        return {'collection_status': 'complete', 'records': [product(asin=asin, price=20) for asin in options['asins']]}
    monkeypatch.setattr(flow.collectors.browser, 'collect_sellersprite_by_asin', enrich)
    r = flow.run_selection({'scope': {**scope(['sellersprite_new']), 'review_batch_size': 2}, 'run_dir': str(tmp_path)})
    for _ in range(3):
        assert all(p['current']['price'] == 20 for p in r['analysis_requests'])
        r = flow.run_selection({'run_dir': str(tmp_path), 'evidence_id': r['evidence_id'],
                                'reviews': [review(p) for p in r['analysis_requests']]})
    assert r['status'] == 'PARTIAL'  # Sales/review analysis still missing.
    assert r['analysis_gaps'] and r['workbook']['row_count'] == 5
    assert list(map(len, calls)) == [2, 2, 1]


def test_refresh_invalidates_reviews_when_product_evidence_changes(monkeypatch, tmp_path):
    rows = [product()]
    monkeypatch.setattr(flow, 'collect_query', lambda *args: {'collection_status': 'complete', 'records': rows})
    r = flow.run_selection({'scope': scope(['sellersprite_new']), 'run_dir': str(tmp_path), 'enrich': False})
    r = flow.run_selection({'run_dir': str(tmp_path), 'reviews': [review(p) for p in r['analysis_requests']]})
    assert r['status'] == 'PARTIAL'  # Sales/review analysis still missing.
    assert r['analysis_gaps']
    rows[0]['estimated_sales'] = 0
    r = flow.run_selection({'run_dir': str(tmp_path), 'refresh': True})
    assert r['status'] == 'AWAITING_ANALYSIS' and r['counts']['reviewed'] == 0


def test_product_query_rejects_all_variants_response():
    from types import SimpleNamespace
    from selection_collectors import response_matches
    url = 'https://www.sellersprite.com/v3/api/product-research'
    for flag, accepted in [(True, True), ('true', True), (False, False), ('false', False), (None, False)]:
        response = SimpleNamespace(url=url, request=SimpleNamespace(url=url, post_data=json.dumps({'symbolFlag': flag})))
        assert response_matches(response, path='/v3/api/product-research', expected={'symbolFlag': True}) is accepted


def test_large_requested_result_count_is_not_capped_at_one_hundred():
    plan = flow.plan_selection({'scope': {**scope(), 'shortlist_limit': 500}})
    assert plan['shortlist_limit'] == 500
    assert flow.plan_selection({'scope': scope()})['shortlist_limit'] == 50


@pytest.mark.parametrize('same_parent', [False, True])
def test_collector_disables_variants_and_checks_each_page(monkeypatch, same_parent):
    from contextlib import contextmanager
    from unittest.mock import MagicMock
    import selection_collectors as collectors
    page = MagicMock()
    page.get_by_text.return_value.locator.return_value.get_by_role.return_value.get_attribute.return_value = 'active'
    page.locator.return_value.filter.return_value.first.get_attribute.return_value = 'active is-checked'
    page.locator.return_value.filter.return_value.first.locator.return_value.input_value.return_value = 'heated blanket'
    disabled = []
    queries = []
    @contextmanager
    def session(*args, **kwargs):
        yield page
    def verified(_page, trigger, *, path, expected):
        assert disabled  # The real flow invokes the control before submitting.
        assert expected['symbolFlag'] is True
        queries.append(expected)
        asin = f'B{len(queries):09d}'
        return {'items': [{'asin': asin, 'parentAsin': f'B{11 if same_parent else len(queries)+10:09d}'}],
                'total': 2, 'hasNextPage': len(queries) == 1}, expected
    monkeypatch.setattr(collectors, 'session', session)
    monkeypatch.setattr(collectors, 'ensure_product_list_view', lambda p: None)
    monkeypatch.setattr(collectors.browser, '_ensure_sellersprite_variants_off', lambda p: disabled.append(p))
    monkeypatch.setattr(collectors, 'verified_query', verified)
    monkeypatch.setattr(collectors.browser, 'extract_sellersprite_table', lambda *args: [product(asin=f'B{len(queries):09d}')])
    result = collectors.collect_sellersprite_products({'marketplace': 'US', 'source': 'sellersprite_history',
                  'keyword': 'heated blanket', 'history_month': '2025-12'})
    assert [q['page'] for q in queries] == [1, 2]  # Default collection no longer stops at page one.
    assert len(disabled) >= 2 and result['source_metadata']['variants_verified'] is True
    assert result['collection_status'] == 'complete' and len(result['records']) == (1 if same_parent else 2)
    assert len(result['raw_records']) == 2
    assert len(merge_products([{**result, 'query': {'source': 'sellersprite_history'}}])) == 2


def test_stale_review_token_rejected_after_source_refresh(monkeypatch, tmp_path):
    rows = [product()]
    monkeypatch.setattr(flow, 'collect_query', lambda *args: {'collection_status': 'complete', 'records': rows})
    r = flow.run_selection({'scope': scope(['sellersprite_new']), 'run_dir': str(tmp_path), 'enrich': False})
    rows[0]['estimated_sales'] = 1
    with pytest.raises(ValueError, match='审阅证据已变更'):
        flow.run_selection({'run_dir': str(tmp_path), 'refresh': True, 'evidence_id': r['evidence_id'],
                            'reviews': [review(p) for p in r['analysis_requests']]})


def test_submitting_reviews_keeps_failed_sources_and_evidence_stable(monkeypatch, tmp_path):
    calls = []
    def collect(query, *args):
        calls.append(query['source'])
        if query['source'] == 'keyword_trends':
            return {'collection_status': 'blocked', 'records': [], 'block_reason': 'TimeoutError'}
        return {'collection_status': 'complete', 'records': [product()]}
    monkeypatch.setattr(flow, 'collect_query', collect)
    r = flow.run_selection({'scope': scope(), 'run_dir': str(tmp_path), 'enrich': False})
    before = list(calls)
    r = flow.run_selection({'run_dir': str(tmp_path), 'evidence_id': r['evidence_id'],
                            'reviews': [review(p) for p in r['analysis_requests']]})
    assert calls == before
    assert r['status'] == 'PARTIAL' and r['workbook']['row_count'] == 1
    flow.run_selection({'run_dir': str(tmp_path)})
    assert calls == before + ['keyword_trends']


@pytest.mark.parametrize('input_urls', [[], ['https://www.amazon.com/dp/B012345678'],
                                      ['https://www.amazon.com/stores/example/page/123']])
def test_default_fifty_counts_only_verified_families_and_replenishes(monkeypatch, tmp_path, input_urls):
    rows = [product(asin='B000000001', parent_asin=None), product(asin='B000000002', parent_asin=None)]
    rows += [product(asin=f'B{i:09d}', parent_asin=f'B{1000 + i // 2:09d}') for i in range(100, 202)]
    monkeypatch.setattr(flow, 'collect_query', lambda *args: {'collection_status': 'complete', 'records': rows})
    value = {**scope(['sellersprite_new']), 'input_urls': input_urls}
    r = flow.run_selection({'scope': value, 'run_dir': str(tmp_path), 'enrich': False})
    assert len(r['analysis_requests']) == 50
    r = flow.run_selection({'run_dir': str(tmp_path), 'reviews': [review(p) for p in r['analysis_requests']]})
    assert r['counts']['accepted'] == 48 and len(r['analysis_requests']) == 2
    r = flow.run_selection({'run_dir': str(tmp_path), 'reviews': [review(p) for p in r['analysis_requests']]})
    assert r['workbook']['row_count'] == r['counts']['accepted'] == 50
    assert r['counts']['pending_parent'] == 2 and r['status'] == 'PARTIAL'
    from openpyxl import load_workbook
    asins = [row[10].value for row in list(load_workbook(r['workbook']['path']).active.rows)[1:]]
    parents = {x['asin']: x.get('parent_asin') for x in rows}
    assert len({parents[a] for a in asins}) == 50 and all(parents[a] for a in asins)


def test_parent_only_missing_triggers_enrichment_before_review(monkeypatch, tmp_path):
    rows = [product(asin=f'B{i:09d}', parent_asin=None) for i in range(1, 4)]
    monkeypatch.setattr(flow, 'collect_query', lambda *args: {'collection_status': 'complete', 'records': rows})
    calls = []
    def enrich(payload):
        calls.append(payload['asins'])
        return {'collection_status': 'complete', 'records': [
            {**row, 'parent_asin': 'B000000010' if i < 2 else 'B000000020'} for i, row in enumerate(rows)]}
    monkeypatch.setattr(flow.collectors.browser, 'collect_sellersprite_by_asin', enrich)
    r = flow.run_selection({'scope': scope(['sellersprite_new']), 'run_dir': str(tmp_path)})
    assert calls == [[x['asin'] for x in rows]]
    assert len(r['analysis_requests']) == 2 and r['counts']['pending_parent'] == 0
    assert {x['verified_parent_asin'] for x in r['analysis_requests']} == {'B000000010', 'B000000020'}


def test_unknown_parents_never_become_delivered_products_and_can_retry(monkeypatch, tmp_path):
    monkeypatch.setattr(flow, 'collect_query', lambda *args: {
        'collection_status': 'complete', 'records': [product(parent_asin=None)]})
    r = flow.run_selection({'scope': scope(['sellersprite_new']), 'run_dir': str(tmp_path), 'enrich': False})
    r = flow.run_selection({'run_dir': str(tmp_path), 'reviews': [review(p) for p in r['analysis_requests']]})
    assert r['status'] == 'AWAITING_ENRICHMENT' and r['counts']['accepted'] == 0
    assert 'workbook' not in r and not (tmp_path / '开品结果.xlsx').exists()
    monkeypatch.setattr(flow.collectors.browser, 'collect_sellersprite_by_asin', lambda p: {
        'collection_status': 'complete', 'records': [product()]})
    r = flow.run_selection({'run_dir': str(tmp_path), 'enrich': True})
    assert r['status'] == 'AWAITING_ANALYSIS' and r['counts']['pending_parent'] == 0


def test_complete_lookup_without_parent_does_not_prevent_parent_retry(monkeypatch, tmp_path):
    row = product(parent_asin=None)
    monkeypatch.setattr(flow, 'collect_query', lambda *args: {'collection_status': 'complete', 'records': [row]})
    calls = []
    def enrich(payload):
        calls.append(payload['asins'])
        return {'collection_status': 'complete', 'records': [row if len(calls) == 1 else product()]}
    monkeypatch.setattr(flow.collectors.browser, 'collect_sellersprite_by_asin', enrich)
    r = flow.run_selection({'scope': scope(['sellersprite_new']), 'run_dir': str(tmp_path)})
    assert r['counts']['pending_parent'] == 1
    r = flow.run_selection({'run_dir': str(tmp_path)})
    assert len(calls) == 2 and r['counts']['pending_parent'] == 0


def test_resume_adds_evidence_to_workbook_and_clears_partial_status(monkeypatch, tmp_path):
    from test_product_evidence import analysis
    from openpyxl import load_workbook
    monkeypatch.setattr(flow, 'collect_query', lambda *args: {
        'collection_status': 'complete', 'records': [product()]})
    response = flow.run_selection({'scope': {**scope(['sellersprite_new']), 'shortlist_limit': 1}, 'run_dir': str(tmp_path),
                                   'as_of_date': '2026-10-13', 'enrich': False})
    plain = review(response['analysis_requests'][0])
    response = flow.run_selection({'run_dir': str(tmp_path), 'reviews': [plain]})
    assert response['status'] == 'PARTIAL' and response['analysis_gaps']
    completed = {**plain, **analysis()}
    response = flow.run_selection({'run_dir': str(tmp_path), 'reviews': [completed]})
    assert response['status'] == 'COMPLETE' and response['analysis_gaps'] == []
    wb = load_workbook(response['workbook']['path'])
    assert '所读样本 1/2 条' in wb.active.cell(2, 9).value
    assert wb.active.cell(2, 10).value == completed['lifecycle_analysis']['summary']
    wb.close()
    evidence = json.loads(Path(response['evidence_path']).read_text())
    assert evidence['reviews'][0]['review_analysis']['samples'][0]['text'] == 'The fan is noisy.'
    assert flow.run_selection({'run_dir': str(tmp_path)})['status'] == 'COMPLETE'


@pytest.mark.parametrize('target, expected', [(50, 46), (10, 6), (None, None)])
def test_fully_analyzed_four_products_do_not_complete_larger_target(monkeypatch, tmp_path, target, expected):
    from test_product_evidence import analysis
    from copy import deepcopy
    rows = [product(asin=f'B{i:09d}') for i in range(4)]
    monkeypatch.setattr(flow, 'collect_query', lambda *args: {'collection_status': 'complete', 'records': rows})
    values = scope(['sellersprite_new'])
    if target != 50:
        values['shortlist_limit'] = target
    response = flow.run_selection({'scope': values, 'run_dir': str(tmp_path),
                                    'as_of_date': '2026-10-13', 'enrich': False})
    reviews = []
    for item in response['analysis_requests']:
        data = deepcopy(analysis())
        for point in data['lifecycle_analysis']['history'] + data['review_analysis']['samples']:
            point['asin'] = item['asin']
        reviews.append({**data, **review(item)})
    response = flow.run_selection({'run_dir': str(tmp_path), 'reviews': reviews})
    assert response['counts']['target'] == target
    assert response['counts']['accepted'] == 4
    assert response['counts']['shortfall'] == expected
    assert response['status'] == ('COMPLETE' if target is None else 'PARTIAL')


def test_missing_history_requests_real_collection_and_reuses_saved_points(monkeypatch, tmp_path):
    monkeypatch.setattr(flow, 'collect_query', lambda *args: {'collection_status': 'complete', 'records': [product()]})
    response = flow.run_selection({'scope': scope(['sellersprite_new']), 'run_dir': str(tmp_path), 'enrich': False})
    response = flow.run_selection({'run_dir': str(tmp_path), 'reviews': [review(response['analysis_requests'][0])]})
    assert response['product_evidence_requests'][0]['skill_action'] == 'collect_product_evidence'
    assert response['product_evidence_requests'][0]['asins'] == ['B012345678']
    cached = {'marketplace': 'US', 'asin': 'B012345678', 'history': [{'month': '2026-01', 'estimated_sales': 58}]}
    (tmp_path / 'product-history-US-B012345678.json').write_text(json.dumps(cached))
    response = flow.run_selection({'run_dir': str(tmp_path)})
    assert response['product_history'] == [cached]
    assert not response['product_evidence_requests']
    assert response['analysis_gaps']  # Collection never pretends to be semantic analysis.
