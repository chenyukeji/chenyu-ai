import importlib.util
from pathlib import Path
import pytest
from openpyxl import load_workbook
from PIL import Image

SCRIPT = Path(__file__).resolve().parents[1] / 'plugins/chenyu-wuliu/skills/chenyu-xiangmai-fapiao/scripts/generate_shipping_table.py'
spec = importlib.util.spec_from_file_location('shipping_table', SCRIPT)
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)

def data():
    return {'cartons': [{'sourcePhotos': ['photo.jpg'], 'weightKg': 16.05, 'items': [
        {'nameZh': '装备A', 'rawText': '装备A=12', 'quantity': 12},
        {'nameZh': '装备B', 'rawText': '装备B=6', 'quantity': 6},
        {'nameZh': '装饰混配100', 'rawText': '混配100=100', 'quantity': 100}]}]}

def test_variants_shared_carton_and_missing_fields(tmp_path):
    out = module.generate(data(), tmp_path / '物流表.xlsx')
    book = load_workbook(out)
    sheet = book['物流表']
    assert [sheet.cell(row, 9).value for row in (2, 3, 4)] == [12, 6, 100]
    assert sheet['B2'].value == 16.05 and sheet['B3'].value is None
    assert 'B2:B4' in sheet.merged_cells
    assert sheet['H2'].value is None and sheet['K2'].value is None
    assert book['待补信息'].max_row == 4 and book['识别依据'].max_row == 4
    assert '申报单价' in book['待补信息']['D2'].value
    assert [c.value for c in sheet[1]] == module.HEADERS
    book.close()

def test_product_image_hs_text_and_formula_injection(tmp_path):
    image = tmp_path / 'product.png'
    Image.new('RGB', (30, 30), 'red').save(image)
    value = data()
    value['cartons'][0]['items'][0].update(imagePath=str(image), hsCode='00123456', nameZh='=HYPERLINK("bad")')
    out = module.generate(value, tmp_path / '物流表.xlsx')
    book = load_workbook(out)
    assert book['物流表']['G2'].data_type == 's'
    assert book['物流表']['K2'].value == '00123456'
    assert len(book['物流表']._images) == 1
    book.close()

@pytest.mark.parametrize('value', [True, 0, -1, 1.5, '12', float('inf')])
def test_invalid_quantities_rejected(tmp_path, value):
    payload = data()
    payload['cartons'][0]['items'][0]['quantity'] = value
    with pytest.raises(ValueError):
        module.generate(payload, tmp_path / 'out.xlsx')

def test_empty_or_duplicate_cartons_rejected(tmp_path):
    with pytest.raises(ValueError): module.generate({'cartons': []}, tmp_path / 'out.xlsx')
    payload = data()
    payload['cartons'][0]['cartonId'] = 'same'
    payload['cartons'] *= 2
    with pytest.raises(ValueError): module.generate(payload, tmp_path / 'out.xlsx')
