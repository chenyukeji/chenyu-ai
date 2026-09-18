import importlib.util
import json
import math
import subprocess
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import pytest

ROOT = Path(__file__).resolve().parents[1]
PLUGIN = ROOT / "plugins" / "chenyu-kaifa"


def load_module(name, relpath):
    path = PLUGIN / relpath
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


kaifa = load_module("chenyu_kaifa_run", "skills/chenyu-kaifa/scripts/run.py")
jihui = load_module("chenyu_jihui_run", "skills/chenyu-jihui/scripts/run.py")
pingshen = load_module("chenyu_kaifa_pingshen_run", "skills/chenyu-kaifa-pingshen/scripts/run.py")


def full_evidence(candidate_id="cand-1"):
    fields = {
        "cn_seller_evidence": 1,
        "target_market_demand": {"signal": "observed"},
        "comparable_competition": {"groups": 12},
        "comparable_price": {"min": 11.99, "max": 16.99},
        "product_specs": {"material": "silicone", "size": "10cm", "pack_count": 1},
        "review_pain_points": ["hard to clean"],
        "fulfillment_and_delivery": {"fba_presence": "observed"},
    }
    return [
        {
            "evidence_id": f"ev-{i}",
            "candidate_id": candidate_id,
            "field": field,
            "value": value,
            "evidence_status": "fact" if field != "target_market_demand" else "estimate",
            "source_type": "fixture",
            "source_ref": f"fixture:{field}",
            "source_market": "DE",
            "observed_at": "2026-09-18T12:00:00+02:00",
        }
        for i, (field, value) in enumerate(fields.items(), start=1)
    ]


def example_economics():
    return pingshen.calculate_economics(
        {
            "marketplace": "DE",
            "fulfillment_path": "FBA",
            "price_gross_eur": 15.0,
            "vat_rate": 0.19,
            "cny_per_eur": 8.0,
            "purchase_cny": 14.0,
            "packaging_cny": 2.0,
            "referral_fee_eur": 2.25,
            "fulfillment_fee_eur": 2.70,
            "inbound_eur": 0.60,
            "storage_eur": 0.30,
            "returns_reserve_eur": 0.30,
            "other_variable_eur": 0.0,
            "cpc_eur": 0.30,
            "ad_cvr": 0.10,
            "ad_order_share": 1.0,
            "ad_sales_per_order_eur": 15.0,
            "trial_units": 20,
        }
    )


def test_runtime_rules_are_identical_and_executable():
    paths = [
        PLUGIN / "skills/chenyu-kaifa/references/runtime-rules.json",
        PLUGIN / "skills/chenyu-jihui/references/runtime-rules.json",
        PLUGIN / "skills/chenyu-kaifa-pingshen/references/runtime-rules.json",
    ]
    rules = [json.loads(path.read_text(encoding="utf-8")) for path in paths]
    assert rules[0] == rules[1] == rules[2]
    assert rules[0]["plugin_name"] == "chenyu-kaifa"
    assert rules[0]["status"] == "runtime"
    assert rules[0]["sales_marketplaces"] == ["DE", "FR", "IT", "ES"]
    assert rules[0]["target_price_eur"] == {"min": 5.0, "max": 20.0, "inclusive": True}
    assert rules[0]["purchase_packaging_cap_cny"] == 20.0


def test_task_entry_applies_defaults_and_auto_routes():
    task = kaifa.create_task(
        {
            "task_id": "task-fixture",
            "category_or_need": "kitchen organization",
            "keywords": ["drawer organizer"],
            "target_season": "Q4",
        }
    )
    assert task["sales_marketplaces"] == ["DE", "FR", "IT", "ES"]
    assert task["target_price_eur"] == {"min": 5.0, "max": 20.0}
    assert task["purchase_packaging_cap_cny"] == 20.0
    assert set(["A", "E", "J", "C", "H"]).issubset(task["strategy_ids"])
    plan = kaifa.build_plan(task)
    assert [p["owner_skill"] for p in plan["phases"]] == [
        "chenyu-jihui",
        "chenyu-jihui",
        "chenyu-kaifa-pingshen",
        "chenyu-kaifa-pingshen",
    ]


def test_task_entry_rejects_relaxing_business_cap():
    with pytest.raises(kaifa.ContractError):
        kaifa.create_task({"purchase_packaging_cap_cny": 21})


def test_multi_strategy_parent_candidate_merges_without_summing_sales():
    rows = [
        {
            "marketplace": "DE",
            "asin": "CHILD-A",
            "parent_asin": "PARENT-1",
            "product_name": "Example Product",
            "source_strategy": "E",
            "source_ref": "new-release:1",
            "estimated_sales": 50,
            "sales_level": "parent",
            "observed_at": "2026-09-18T10:00:00+02:00",
        },
        {
            "marketplace": "DE",
            "asin": "CHILD-B",
            "parent_asin": "PARENT-1",
            "product_name": "Example Product Variant",
            "source_strategy": "J",
            "source_ref": "fbm:2",
            "estimated_sales": 50,
            "sales_level": "parent",
            "observed_at": "2026-09-18T10:00:00+02:00",
        },
        {
            "marketplace": "DE",
            "asin": "CHILD-A",
            "parent_asin": "PARENT-1",
            "product_name": "Example Product",
            "source_strategy": "A",
            "source_ref": "cn-seller:3",
            "estimated_sales": 50,
            "sales_level": "parent",
            "observed_at": "2026-09-18T10:00:00+02:00",
        },
    ]
    merged = jihui.merge_candidates(rows)
    assert len(merged) == 1
    candidate = merged[0]
    assert set(candidate["source_strategies"]) == {"A", "E", "J"}
    assert candidate["source_record_count"] == 3
    assert len(candidate["marketplace_listings"]) == 2
    assert candidate["sales_aggregation"]["status"] == "not_aggregated"
    assert "sales_total" not in candidate


def test_opportunity_card_preserves_pending_instead_of_inventing_values():
    candidate = {
        "candidate_id": "cand-1",
        "product_name_normalized": "Example Product",
        "source_strategies": ["E"],
        "source_refs": ["new-release:1"],
    }
    card = jihui.build_opportunity_card(candidate, full_evidence()[:2], ["DE"])
    assert card["decision_readiness"] == "partial"
    assert "product_specs" in card["pending_verification"]
    pending = card["evidence"]["product_specs"][0]
    assert pending["evidence_status"] == "pending_verification"
    assert pending["value"] is None


def test_supply_match_and_bundle_apply_20_cny_hard_gate():
    match = pingshen.match_supply(
        {
            "candidate_id": "cand-1",
            "material": "silicone",
            "size": "10cm",
            "pack_count": 1,
        },
        {
            "supplier": "fixture",
            "supplier_sku": "SKU-1",
            "material": "Silicone",
            "size": "10cm",
            "pack_count": 1,
            "unit_quote_cny": 14,
            "packaging_cny": 2,
            "sample_status": "passed",
        },
    )
    assert match["match_type"] == "exact"
    assert match["cost_gate"]["status"] == "pass"
    assert match["cost_gate"]["total_cny"] == 16

    bundle = pingshen.build_bundle(
        [
            {"supplier_sku": "A", "unit_quote_cny": 11, "quantity": 1},
            {"supplier_sku": "B", "unit_quote_cny": 8, "quantity": 1},
        ],
        packaging_cny=2,
        shared_purchase_reason="used together in the same task",
    )
    assert bundle["cost_gate"]["status"] == "fail"
    assert bundle["cost_gate"]["total_cny"] == 21


def test_economics_matches_documented_formula_and_is_reproducible():
    econ = example_economics()
    assert econ["calculation_status"] == "calculated"
    assert econ["price_gate"]["status"] == "pass"
    assert econ["purchase_packaging_gate"]["status"] == "pass"
    assert econ["net_revenue_eur"] == pytest.approx(15 / 1.19, abs=1e-6)
    assert econ["cm_before_ads_eur"] == pytest.approx(4.455042, abs=1e-6)
    assert econ["cm_after_ads_eur"] == pytest.approx(1.455042, abs=1e-6)
    assert econ["break_even_cpc_eur"] == pytest.approx(0.445504, abs=1e-6)
    assert econ["trial_procurement_packaging_cny"] == 320


def test_economics_does_not_treat_missing_platform_fee_as_zero():
    econ = pingshen.calculate_economics(
        {
            "price_gross_eur": 15,
            "vat_rate": 0.19,
            "cny_per_eur": 8,
            "purchase_cny": 14,
            "packaging_cny": 2,
        }
    )
    assert econ["calculation_status"] == "pending"
    assert "referral_fee_eur" in econ["missing_inputs"]
    assert "fulfillment_fee_eur" in econ["missing_inputs"]
    assert "cm_before_ads_eur" not in econ


def test_review_recommends_trial_only_after_market_supply_profit_and_sample():
    candidate = {
        "candidate_id": "cand-1",
        "product_name_normalized": "Example Product",
        "source_strategies": ["A", "E", "J"],
        "source_refs": ["fixture"],
    }
    card = jihui.build_opportunity_card(candidate, full_evidence(), ["DE"])
    supply = pingshen.match_supply(
        {"candidate_id": "cand-1", "material": "silicone", "size": "10cm", "pack_count": 1},
        {
            "supplier": "fixture",
            "supplier_sku": "SKU-1",
            "material": "silicone",
            "size": "10cm",
            "pack_count": 1,
            "unit_quote_cny": 14,
            "packaging_cny": 2,
            "sample_status": "passed",
        },
    )
    decision = pingshen.review_development(card, supply, example_economics())
    assert decision["decision_code"] == "TRIAL_RECOMMENDED"
    assert decision["decision_label"] == "建议试销"


def test_review_rejects_hard_cost_failure_with_reason_code():
    card = jihui.build_opportunity_card(
        {"candidate_id": "cand-1", "product_name_normalized": "Example", "source_strategies": ["A"]},
        full_evidence(),
        ["DE"],
    )
    supply = {
        "match_type": "exact",
        "sample_status": "passed",
        "cost_gate": {"status": "fail", "total_cny": 22, "cap_cny": 20},
    }
    decision = pingshen.review_development(card, supply, example_economics())
    assert decision["decision_code"] == "REJECTED"
    assert "PURCHASE_PACKAGING_CAP_FAIL" in decision["reason_codes"]



def test_cross_market_title_similarity_does_not_merge_without_verified_identity():
    rows = [
        {
            "marketplace": "DE",
            "product_name": "Same Looking Product",
            "product_signature": "sig-123",
            "source_strategy": "E",
            "source_ref": "de:1",
        },
        {
            "marketplace": "FR",
            "product_name": "Same Looking Product",
            "product_signature": "sig-123",
            "source_strategy": "E",
            "source_ref": "fr:1",
        },
    ]
    assert len(jihui.merge_candidates(rows)) == 2

    for row in rows:
        row["cross_market_identity_verified"] = True
    merged = jihui.merge_candidates(rows)
    assert len(merged) == 1
    assert merged[0]["dedupe_confidence"] == "medium"


def test_cn_seller_gate_zero_fact_is_a_real_admission_failure():
    candidate = {
        "candidate_id": "cand-cn-zero",
        "product_name_normalized": "Example",
        "source_strategies": ["A"],
    }
    evidence = full_evidence("cand-cn-zero")
    evidence[0]["value"] = 0
    evidence[0]["evidence_status"] = "fact"
    card = jihui.build_opportunity_card(candidate, evidence, ["DE"])
    assert card["cn_seller_gate"]["status"] == "fail"

    supply = {
        "match_type": "exact",
        "sample_status": "passed",
        "cost_gate": {"status": "pass", "total_cny": 16, "cap_cny": 20},
    }
    decision = pingshen.review_development(card, supply, example_economics())
    assert decision["decision_code"] == "REJECTED"
    assert "CN_SELLER_GATE_FAIL" in decision["reason_codes"]

def test_fixed_workbook_has_all_ten_delivery_sheets(tmp_path):
    target = tmp_path / "development.xlsx"
    result = pingshen.export_workbook(
        target,
        {
            "task": {"task_id": "task-fixture"},
            "candidates": [{"candidate_id": "cand-1"}],
            "evidence": full_evidence(),
            "opportunity_cards": [{"candidate_id": "cand-1", "decision_readiness": "ready_for_supply_match"}],
            "supply_match": [{"candidate_id": "cand-1", "supplier_sku": "SKU-1"}],
            "economics": [example_economics()],
            "decision": [{"decision_code": "TRIAL_RECOMMENDED"}],
            "product_master": [{"candidate_id": "cand-1", "fact_status": "confirmed"}],
            "handoff": [{"department": "采购", "task": "confirm quote validity"}],
            "trial_card": [{"candidate_id": "cand-1", "quantity_status": "pending_human_approval"}],
        },
    )
    assert result["sheet_count"] == 10
    assert target.exists()
    with zipfile.ZipFile(target) as z:
        assert all(f"xl/worksheets/sheet{i}.xml" in z.namelist() for i in range(1, 11))
        root = ET.fromstring(z.read("xl/workbook.xml"))
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    names = [sheet.attrib["name"] for sheet in root.findall("s:sheets/s:sheet", ns)]
    assert names == pingshen.load_rules()["workbook_sheets"]


def test_strategy_j_file_import_enforces_60d_fbm_and_filter_state(tmp_path):
    source = tmp_path / "j.json"
    source.write_text(
        json.dumps(
            {
                "requested_filters": {"fulfillment": "FBM", "age_days_max": 60},
                "applied_filters": {"fulfillment": "FBM", "age_days_max": 60},
                "rows": [
                    {
                        "marketplace": "DE",
                        "asin": "J-VALID",
                        "product_name": "Valid",
                        "fulfillment": "FBM",
                        "source_available_date": "2026-08-01",
                    },
                    {
                        "marketplace": "DE",
                        "asin": "J-OLD",
                        "product_name": "Old",
                        "fulfillment": "FBM",
                        "source_available_date": "2026-07-01",
                    },
                    {
                        "marketplace": "DE",
                        "asin": "J-FBA",
                        "product_name": "FBA",
                        "fulfillment": "FBA",
                        "source_available_date": "2026-09-01",
                    },
                    {
                        "marketplace": "DE",
                        "asin": "J-PENDING",
                        "product_name": "Pending",
                        "fulfillment": "FBM",
                    },
                ],
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    result = jihui.import_discovery_file(str(source), "J", as_of_date="2026-09-18")
    assert result["filter_verification"] == "verified"
    assert result["counts"] == {"rows": 4, "accepted": 1, "pending": 1, "rejected": 2}
    assert result["accepted"][0]["asin"] == "J-VALID"
    assert result["pending"][0]["asin"] == "J-PENDING"
    assert "source_available_date" in result["pending"][0]["strategy_pending_fields"]
    rejection_codes = {code for row in result["rejected"] for code in row["reason_codes"]}
    assert {"AGE_OVER_60_DAYS", "FULFILLMENT_NOT_FBM"}.issubset(rejection_codes)


def test_filter_mismatch_never_becomes_verified_strategy_match(tmp_path):
    source = tmp_path / "a.csv"
    source.write_text(
        "asin,marketplace,seller_location,product_name\n"
        "A1,DE,CN,Example\n",
        encoding="utf-8",
    )
    result = jihui.import_discovery_file(
        str(source),
        "A",
        requested_filters={"seller_location": "CN"},
        applied_filters={"seller_location": "ALL"},
    )
    assert result["filter_verification"] == "mismatch"
    assert result["counts"]["accepted"] == 0
    assert result["counts"]["pending"] == 1
    assert result["pending"][0]["strategy_match_status"] == "unknown"
    assert "filter_mismatch" in result["pending"][0]["strategy_pending_fields"]


def test_strategy_e_only_accepts_us_or_de_discovery_sources(tmp_path):
    source = tmp_path / "e.json"
    source.write_text(
        json.dumps(
            {
                "filter_verification": "verified",
                "rows": [
                    {"marketplace": "US", "asin": "US1", "product_name": "US"},
                    {"marketplace": "DE", "asin": "DE1", "product_name": "DE"},
                    {"marketplace": "FR", "asin": "FR1", "product_name": "FR"},
                ],
            }
        ),
        encoding="utf-8",
    )
    result = jihui.import_discovery_file(str(source), "E")
    assert {row["asin"] for row in result["accepted"]} == {"US1", "DE1"}
    assert result["rejected"][0]["normalized"]["asin"] == "FR1"
    assert "NEW_RELEASE_SOURCE_MARKET_NOT_US_DE" in result["rejected"][0]["reason_codes"]


def test_market_evidence_fba_state_respects_sample_coverage():
    base_hits = [
        {"asin": "A", "parent_asin": "P1", "relevant": True, "fulfillment": "FBM", "price": 12.0, "seller_location": "CN"},
        {"asin": "B", "parent_asin": "P2", "relevant": True, "fulfillment": "FBM", "price": 14.0, "seller_location": "DE"},
    ]
    partial = jihui.build_market_evidence(
        "cand-1",
        "DE",
        "2026-09-18T12:00:00+02:00",
        search_hits=base_hits,
        coverage_complete=False,
    )
    fulfillment = next(e for e in partial["evidence"] if e["field"] == "fulfillment_and_delivery")
    assert fulfillment["value"]["fba_presence"] == "unknown"

    complete = jihui.build_market_evidence(
        "cand-1",
        "DE",
        "2026-09-18T12:00:00+02:00",
        search_hits=base_hits,
        coverage_complete=True,
    )
    fulfillment = next(e for e in complete["evidence"] if e["field"] == "fulfillment_and_delivery")
    assert fulfillment["value"]["fba_presence"] == "not_observed_in_sample"

    with_fba = jihui.build_market_evidence(
        "cand-1",
        "DE",
        "2026-09-18T12:00:00+02:00",
        search_hits=base_hits + [{"asin": "C", "parent_asin": "P3", "relevant": True, "fulfillment": "FBA", "price": 13.0}],
        coverage_complete=False,
    )
    fulfillment = next(e for e in with_fba["evidence"] if e["field"] == "fulfillment_and_delivery")
    assert fulfillment["value"]["fba_presence"] == "observed"


def test_market_snapshot_can_feed_a_ready_opportunity_card():
    candidate = {
        "candidate_id": "cand-market",
        "product_name_normalized": "Example",
        "source_strategies": ["A", "E", "J"],
        "source_refs": ["fixture"],
    }
    market = jihui.build_market_evidence(
        "cand-market",
        "DE",
        "2026-09-18T12:00:00+02:00",
        search_hits=[
            {
                "asin": "A",
                "parent_asin": "P1",
                "relevant": True,
                "fulfillment": "FBA",
                "landed_price": 12.99,
                "seller_location": "CN",
                "purchase_cue": "100+ bought",
            },
            {
                "asin": "B",
                "parent_asin": "P2",
                "relevant": True,
                "fulfillment": "FBM",
                "landed_price": 15.99,
                "seller_location": "DE",
            },
        ],
        product_details=[{"asin": "A", "material": "silicone", "size": "10cm", "pack_count": 1}],
        review_annotations=[{"pain_point": "hard to clean", "count": 4, "review_refs": ["r1", "r2"]}],
        query="example keyword",
        source_ref="snapshot:de:1",
        coverage_complete=True,
    )
    card = jihui.build_opportunity_card(candidate, market["evidence"], ["DE"])
    assert card["decision_readiness"] == "ready_for_supply_match"
    assert card["cn_seller_gate"]["status"] == "pass"
    comparable_price = card["evidence"]["comparable_price"][0]["value"]
    assert comparable_price["median"] == pytest.approx(14.49)


def test_supplier_quote_import_checks_quantity_expiry_and_packaging(tmp_path):
    source = tmp_path / "quotes.json"
    source.write_text(
        json.dumps(
            [
                {
                    "supplier": "S1",
                    "supplier_sku": "SKU-VALID",
                    "currency": "CNY",
                    "unit_quote_cny": 14,
                    "packaging_cny": 2,
                    "packaging_included": False,
                    "min_qty": 10,
                    "max_qty": 50,
                    "quote_valid_until": "2026-10-01",
                },
                {
                    "supplier": "S1",
                    "supplier_sku": "SKU-EXPIRED",
                    "currency": "CNY",
                    "unit_quote_cny": 12,
                    "packaging_included": True,
                    "min_qty": 10,
                    "max_qty": 50,
                    "quote_valid_until": "2026-09-01",
                },
                {
                    "supplier": "S1",
                    "supplier_sku": "SKU-WRONG-TIER",
                    "currency": "CNY",
                    "unit_quote_cny": 5,
                    "packaging_included": True,
                    "min_qty": 1000,
                    "max_qty": 2000,
                    "quote_valid_until": "2026-10-01",
                },
            ]
        ),
        encoding="utf-8",
    )
    result = pingshen.import_supplier_quotes(
        str(source),
        planned_quantity=20,
        as_of_date="2026-09-18",
    )
    assert result["counts"] == {"rows": 3, "valid": 1, "pending": 0, "invalid": 2}
    assert result["valid"][0]["supplier_sku"] == "SKU-VALID"
    codes = {code for row in result["invalid"] for code in row["assessment"]["reason_codes"]}
    assert "QUOTE_EXPIRED" in codes
    assert "PLANNED_QTY_BELOW_QUOTE_RANGE" in codes


def test_packaging_included_quote_does_not_double_count_packaging(tmp_path):
    source = tmp_path / "quotes.csv"
    source.write_text(
        "supplier,supplier_sku,currency,unit_quote_cny,packaging_included,min_qty,max_qty,quote_valid_until\n"
        "S1,SKU-1,CNY,16,true,1,100,2026-12-31\n",
        encoding="utf-8",
    )
    result = pingshen.import_supplier_quotes(str(source), planned_quantity=20, as_of_date="2026-09-18")
    quote = result["valid"][0]
    assert quote["assessment"]["effective_packaging_cny"] == 0


def test_scenario_matrix_has_three_prices_and_three_scenarios():
    matrix = pingshen.calculate_scenario_matrix(
        pingshen.calculate_economics,
        {
            "marketplace": "DE",
            "fulfillment_path": "FBA",
            "price_gross_eur": 15.0,
            "vat_rate": 0.19,
            "cny_per_eur": 8.0,
            "purchase_cny": 14.0,
            "packaging_cny": 2.0,
            "referral_fee_eur": 2.25,
            "fulfillment_fee_eur": 2.70,
            "inbound_eur": 0.60,
            "storage_eur": 0.30,
            "returns_reserve_eur": 0.30,
            "other_variable_eur": 0.0,
            "cpc_eur": 0.30,
            "ad_cvr": 0.10,
            "ad_order_share": 1.0,
            "ad_sales_per_order_eur": 15.0,
        },
    )
    assert matrix["prices"] == [5.0, 15.0, 20.0]
    assert len(matrix["rows"]) == 9
    assert matrix["profile"]["status"] == "proposed_not_calibrated"
    conservative_5 = next(
        row for row in matrix["rows"]
        if row["price_reference_eur"] == 5.0 and row["scenario"] == "conservative"
    )
    assert conservative_5["economics"]["price_gate"]["status"] == "fail"


def test_delivery_sections_keep_unconfirmed_facts_out_of_product_master():
    delivery = pingshen.build_delivery_sections(
        {"rules_snapshot": {"config_revision": "x"}},
        {"candidate_id": "cand-1", "product_name_normalized": "Example"},
        {"decision_readiness": "ready_for_supply_match"},
        {"supplier_sku": "SKU-1", "sample_status": "passed"},
        {"marketplace": "DE", "inputs": {"price_gross_eur": 15}},
        {"decision_code": "TRIAL_RECOMMENDED"},
        {"candidate_id": "cand-1"},
        [
            {"field": "material", "value": "silicone", "status": "confirmed"},
            {"field": "size", "value": "10cm", "status": "unconfirmed"},
        ],
    )
    master = delivery["product_master"]
    assert master["facts"] == {"material": "silicone"}
    assert master["pending_facts"][0]["field"] == "size"
    assert {row["department"] for row in delivery["handoff"]} == {"采购", "运营", "美工"}


def test_trial_review_uses_explicit_thresholds_and_never_executes_actions():
    review = pingshen.review_trial(
        {
            "trial_id": "T1",
            "review_thresholds": {
                "max_acos": 0.35,
                "max_return_rate": 0.10,
                "min_orders": 10,
                "min_contribution_total": 20,
            },
        },
        [
            {
                "sessions": 200,
                "orders": 20,
                "revenue": 300,
                "ad_spend": 80,
                "ad_sales": 200,
                "returns": 4,
                "refunds": 50,
                "inventory": 10,
                "contribution_total": 15,
                "stockout_days": 0,
            }
        ],
    )
    assert review["metrics"]["acos"] == pytest.approx(0.4)
    assert review["metrics"]["return_rate"] == pytest.approx(0.2)
    assert review["recommended_action"] == "exit_review"
    assert "RETURN_RATE_ABOVE_LIMIT" in review["failure_reasons"]
    assert "no_ads_no_purchase" in review["action_boundary"]


def test_trial_review_can_recommend_replenishment_without_ordering():
    review = pingshen.review_trial(
        {
            "trial_id": "T2",
            "review_thresholds": {
                "max_acos": 0.5,
                "max_return_rate": 0.2,
                "min_orders": 5,
                "min_contribution_total": 5,
            },
        },
        [
            {
                "sessions": 100,
                "orders": 10,
                "revenue": 150,
                "ad_spend": 20,
                "ad_sales": 80,
                "returns": 0,
                "refunds": 0,
                "inventory": 0,
                "contribution_total": 20,
                "stockout_days": 0,
            }
        ],
    )
    assert review["recommended_action"] == "replenishment_review"


def test_quote_import_outputs_direct_match_supplier_item(tmp_path):
    source = tmp_path / "quote.json"
    source.write_text(
        json.dumps(
            [{
                "supplier": "S1",
                "supplier_sku": "SKU-1",
                "material": "silicone",
                "size": "10cm",
                "pack_count": 1,
                "currency": "CNY",
                "unit_quote_cny": 16,
                "packaging_included": True,
                "min_qty": 1,
                "max_qty": 100,
                "quote_valid_until": "2026-12-31",
            }]
        ),
        encoding="utf-8",
    )
    quote = pingshen.import_supplier_quotes(
        str(source),
        planned_quantity=20,
        as_of_date="2026-09-18",
    )["valid"][0]
    item = quote["match_supplier_item"]
    assert item["packaging_cny"] == 0
    match = pingshen.match_supply(
        {"candidate_id": "cand-1", "material": "silicone", "size": "10cm", "pack_count": 1},
        item,
    )
    assert match["match_type"] == "exact"
    assert match["cost_gate"]["total_cny"] == 16


def test_trial_csv_import_can_feed_review(tmp_path):
    source = tmp_path / "trial.csv"
    source.write_text(
        "period,sessions,orders,revenue,ad_spend,ad_sales,returns,refunds,inventory,stockout_days,contribution_total\n"
        "2026-09-01/2026-09-07,100,10,150,20,80,0,0,5,0,12\n"
        "2026-09-08/2026-09-14,120,12,180,25,100,1,10,0,0,15\n",
        encoding="utf-8",
    )
    imported = pingshen.import_trial_observations(str(source))
    assert imported["count"] == 2
    review = pingshen.review_trial(
        {
            "trial_id": "T3",
            "review_thresholds": {
                "max_acos": 0.4,
                "max_return_rate": 0.2,
                "min_orders": 10,
                "min_contribution_total": 20,
            },
        },
        imported["observations"],
    )
    assert review["metrics"]["orders"] == 22
    assert review["metrics"]["contribution_total"] == 27
    assert review["recommended_action"] == "replenishment_review"


def test_sample_checklist_for_bundle_requires_physical_verification():
    checklist = pingshen.build_sample_checklist(
        {
            "candidate_id": "cand-bundle",
            "material": "silicone",
            "size": "10cm",
            "pack_count": 2,
        },
        {"supplier_sku": "BUNDLE-SKU"},
        output_mode="complementary_bundle",
        bundle={"shared_purchase_reason": "used together"},
        extra_checks=["confirm closure strength"],
    )
    ids = {row["check_id"] for row in checklist["checks"]}
    assert {"identity", "material", "size", "pack_count", "bundle_reason", "bundle_compatibility", "bundle_packaging"}.issubset(ids)
    assert all(row["status"] == "pending_physical_verification" for row in checklist["checks"])
    assert "not a sample result" in checklist["completion_rule"]


def test_trial_funding_calculates_initial_cash_without_approving_quantity():
    econ = example_economics()
    funding = pingshen.calculate_trial_funding(
        econ,
        planned_units=20,
        ad_budget_eur=60,
        fixed_cost_cny=100,
        other_initial_cash_cny=50,
    )
    assert funding["status"] == "calculated"
    assert funding["components_cny"]["procurement_packaging"] == 320
    assert funding["components_cny"]["inbound"] == 96
    assert funding["components_cny"]["ad_budget"] == 480
    assert funding["initial_cash_requirement_cny"] == 1046
    assert "not a purchase quantity approval" in funding["boundary"]


def test_bundle_passes_full_review_when_cost_profit_and_sample_are_ready():
    candidate = {
        "candidate_id": "cand-bundle-pass",
        "product_name_normalized": "Bundle Example",
        "source_strategies": ["A", "E"],
        "source_refs": ["fixture"],
    }
    card = jihui.build_opportunity_card(candidate, full_evidence("cand-bundle-pass"), ["DE"])
    bundle = pingshen.build_bundle(
        [
            {"supplier_sku": "A", "unit_quote_cny": 7, "quantity": 1},
            {"supplier_sku": "B", "unit_quote_cny": 8, "quantity": 1},
        ],
        packaging_cny=2,
        shared_purchase_reason="components are used together in one customer task",
    )
    assert bundle["cost_gate"]["status"] == "pass"

    econ = pingshen.calculate_economics(
        {
            "marketplace": "DE",
            "fulfillment_path": "FBA",
            "price_gross_eur": 15.0,
            "vat_rate": 0.19,
            "cny_per_eur": 8.0,
            "purchase_cny": 15.0,
            "packaging_cny": 2.0,
            "referral_fee_eur": 2.25,
            "fulfillment_fee_eur": 2.70,
            "inbound_eur": 0.60,
            "storage_eur": 0.30,
            "returns_reserve_eur": 0.30,
            "other_variable_eur": 0.0,
            "cpc_eur": 0.30,
            "ad_cvr": 0.10,
            "ad_order_share": 1.0,
            "ad_sales_per_order_eur": 15.0,
        }
    )
    supply = {
        "supplier_sku": "BUNDLE-A+B",
        "match_type": "exact",
        "sample_status": "passed",
        "cost_gate": bundle["cost_gate"],
    }
    decision = pingshen.review_development(card, supply, econ)
    assert decision["decision_code"] == "TRIAL_RECOMMENDED"

    checklist = pingshen.build_sample_checklist(
        {
            "candidate_id": "cand-bundle-pass",
            "pack_count": 2,
        },
        supply,
        output_mode="complementary_bundle",
        bundle=bundle,
    )
    assert any(row["check_id"] == "bundle_compatibility" for row in checklist["checks"])


def test_multipack_costs_whole_sales_unit_against_20_cny_cap():
    pack = pingshen.build_multipack(
        supplier_sku="SKU-MULTI",
        unit_quote_cny=4,
        pack_count=4,
        packaging_cny=2,
    )
    assert pack["purchase_total_cny"] == 16
    assert pack["cost_gate"]["total_cny"] == 18
    assert pack["cost_gate"]["status"] == "pass"

    fail_pack = pingshen.build_multipack(
        supplier_sku="SKU-MULTI",
        unit_quote_cny=5,
        pack_count=4,
        packaging_cny=2,
    )
    assert fail_pack["cost_gate"]["total_cny"] == 22
    assert fail_pack["cost_gate"]["status"] == "fail"


def test_economics_rejects_market_outside_europe_sales_scope():
    inputs = {
        "marketplace": "UK",
        "price_gross_eur": 15,
        "vat_rate": 0.2,
        "cny_per_eur": 8,
        "purchase_cny": 14,
        "packaging_cny": 2,
        "referral_fee_eur": 2,
        "fulfillment_fee_eur": 3,
        "inbound_eur": 0.6,
    }
    with pytest.raises(pingshen.ContractError):
        pingshen.calculate_economics(inputs)


def test_workbook_export_normalizes_surrogate_pairs_and_reopens(tmp_path):
    target = tmp_path / "unicode.xlsx"
    surrogate_pair = "\ud83d\ude00"
    pingshen.export_workbook(
        target,
        {
            "task": {"note": "中文 " + surrogate_pair},
            "candidates": [],
            "evidence": [],
            "opportunity_cards": [],
            "supply_match": [],
            "economics": [],
            "decision": [],
            "product_master": [],
            "handoff": [],
            "trial_card": [],
        },
    )
    from openpyxl import load_workbook

    wb = load_workbook(target, read_only=False, data_only=False)
    try:
        assert wb["00_Task"]["A2"].value == "中文 😀"
        assert wb.sheetnames == pingshen.load_rules()["workbook_sheets"]
    finally:
        wb.close()
    assert not target.with_name(target.name + ".tmp").exists()


def test_all_three_cli_entrypoints_read_utf8_json_without_windows_mojibake(tmp_path):
    def run_cli(relpath, payload):
        proc = subprocess.run(
            [sys.executable, str(PLUGIN / relpath)],
            input=json.dumps(payload, ensure_ascii=False).encode("utf-8"),
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=False,
        )
        assert proc.returncode == 0, proc.stderr.decode("utf-8", "replace")
        return json.loads(proc.stdout.decode("utf-8"))

    main_result = run_cli(
        "skills/chenyu-kaifa/scripts/run.py",
        {
            "skill_action": "create_task",
            "task": {
                "task_id": "utf8-main",
                "category_or_need": "厨房收纳",
                "strategy_ids": ["A"],
            },
        },
    )
    assert main_result["task"]["category_or_need"] == "厨房收纳"

    research_result = run_cli(
        "skills/chenyu-jihui/scripts/run.py",
        {
            "skill_action": "merge_candidates",
            "candidates": [
                {
                    "marketplace": "DE",
                    "asin": "UTF8-1",
                    "product_name": "硅胶收纳盒",
                    "source_strategy": "A",
                }
            ],
        },
    )
    assert research_result["candidates"][0]["product_name_normalized"] == "硅胶收纳盒"

    target = tmp_path / "utf8-cli.xlsx"
    review_result = run_cli(
        "skills/chenyu-kaifa-pingshen/scripts/run.py",
        {
            "skill_action": "export_workbook",
            "output_path": str(target),
            "sections": {
                "task": {"task_id": "utf8-review", "note": "中文 😀"},
                "candidates": [],
                "evidence": [],
                "opportunity_cards": [],
                "supply_match": [],
                "economics": [],
                "decision": [{"decision_label": "建议试销 😀"}],
                "product_master": [],
                "handoff": [{"department": "采购"}],
                "trial_card": [],
            },
        },
    )
    assert review_result["workbook"]["sheet_count"] == 10

    from openpyxl import load_workbook

    wb = load_workbook(target, read_only=False)
    try:
        assert wb["00_Task"]["B2"].value == "中文 😀"
        assert wb["06_Decision"]["A2"].value == "建议试销 😀"
        assert wb["08_Handoff"]["A2"].value == "采购"
    finally:
        wb.close()
