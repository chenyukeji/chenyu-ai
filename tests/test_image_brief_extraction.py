"""Behavioral regression checks for numbered and unnumbered image briefs."""
import importlib.util
import io
import shutil
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from xml.sax.saxutils import escape

ROOT = Path(__file__).resolve().parents[1]
SCRIPTS = ROOT / 'plugins/chenyu-meigong/skills/chenyu-zuotu/scripts'
sys.path.insert(0, str(SCRIPTS))
from brief_candidates import add_candidates
from extract_brief import extract


def sheet(cells, name='拍图需求', merged=()):
    value = {'name': name, 'cells': [{'cell': address, 'value': value, 'formula': None}
                                    for address, value in cells.items()],
             'merged_ranges': list(merged)}
    add_candidates(value)
    return value


def mixed_cells():
    labels = ['主图', '尺寸图', '功能图1', '功能图2', '场景图1', '场景图2', '场景图3']
    cells = {}
    for row, label in enumerate(labels, 2):
        cells[f'A{row}'] = label
        cells[f'G{row}'] = '制作一张当前产品的图片'
    cells.update({'A9': 'A+', 'G9': '注意：整套A+图的色调要一致。',
                  'G10': '做一张尺寸为970*600的A+图',
                  'G11': '做一张由三张场景拼成一张的A+图',
                  'G12': '参考例图，做一张细节图',
                  'G13': '做一张展示侧面、正面和背面的拼版'})
    return cells


def write_xlsx(path, cells, merged=()):
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    relns = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
    rows = {}
    import re
    for address, value in cells.items():
        row = int(re.search(r'\d+', address)[0])
        rows.setdefault(row, []).append(f'<c r="{address}" t="inlineStr"><is><t>{escape(value)}</t></is></c>')
    contents = ''.join(f'<row r="{row}">{"".join(values)}</row>' for row, values in sorted(rows.items()))
    merges = '<mergeCells>' + ''.join(f'<mergeCell ref="{ref}"/>' for ref in merged) + '</mergeCells>'
    with zipfile.ZipFile(path, 'w') as archive:
        archive.writestr('xl/workbook.xml', f'<workbook xmlns="{ns}" xmlns:r="{relns}"><sheets><sheet name="拍图需求" sheetId="1" r:id="rId1"/></sheets></workbook>')
        archive.writestr('xl/_rels/workbook.xml.rels', f'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="{relns}/worksheet" Target="worksheets/sheet1.xml"/></Relationships>')
        archive.writestr('xl/worksheets/sheet1.xml', f'<worksheet xmlns="{ns}"><sheetData>{contents}</sheetData>{merges}</worksheet>')


class BriefCandidatesTest(unittest.TestCase):
    def test_mixed_named_gallery_and_unnumbered_aplus(self):
        result = sheet(mixed_cells())
        rows = result['task_row_candidates']
        self.assertEqual([item['row'] for item in rows], [2,3,4,5,6,7,8,10,11,12,13])
        self.assertEqual([item['content_type'] for item in rows], ['gallery'] * 7 + ['aplus'] * 4)
        self.assertEqual([item['inferred_order'] for item in rows], list(range(1, 12)))
        self.assertEqual(len(result['task_marker_candidates']), 7)
        self.assertEqual(rows[-1]['evidence_cells'], ['G13'])
        self.assertEqual(result['section_candidates'][0]['cell'], 'A9')

    def test_legacy_ordinals_fullwidth_names_and_numeric_header(self):
        result = sheet({'B1': '图片序号', 'B2': '第一张', 'B3': '图 ２',
                        'B4': ' ３ ', 'B5': '4.0', 'B6': '功能图１',
                        'D2': '1254', 'D3': '970', 'F3': '参考图1'})
        self.assertEqual([c['cell'] for c in result['task_marker_candidates']], ['B2','B3','B4','B5','B6'])
        self.assertEqual(len(result['task_row_candidates']), 5)

    def test_empty_and_product_numbers_are_not_images(self):
        self.assertEqual(sheet({}, '变体')['task_row_candidates'], [])
        result = sheet({'A1': '数量', 'A2': '4', 'B1': '宽度', 'B2': '20',
                        'C2': '产品详情见图1'}, '产品内容')
        self.assertEqual(result['task_row_candidates'], [])

    def test_missing_labels_and_multiple_panels_keep_one_row_candidate(self):
        result = sheet({'D2': '做一张由四张图拼成的四宫格',
                        'D3': '做一张尺寸图，产品有4个尺寸'}, '作图要求')
        self.assertEqual([c['row'] for c in result['task_row_candidates']], [2, 3])

    def test_requirement_column_and_aplus_sheet_without_ordinals(self):
        result = sheet({'C1': '作图要求', 'C2': '纯白背景展示产品',
                        'C3': '注意：整体色调统一', 'C4': '展示产品细节'}, 'FR A+要求')
        self.assertEqual([c['row'] for c in result['task_row_candidates']], [2,4])
        self.assertTrue(all(c['content_type'] == 'aplus' for c in result['task_row_candidates']))

    def test_xlsx_disguised_as_xls_preserves_merges_and_wps_warning(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            source = root / 'brief.xls'
            cells = mixed_cells()
            cells['C2'] = '=DISPIMG("example",1)'
            write_xlsx(source, cells, ['A2:A3'])
            original = source.read_bytes()
            result = extract(source, root / 'out')
            self.assertEqual(result['source_format'], 'xlsx')
            self.assertEqual(result['sheets'][0]['merged_ranges'], ['A2:A3'])
            self.assertEqual(len(result['sheets'][0]['task_row_candidates']), 11)
            self.assertEqual(result['warnings'][0]['cell'], 'C2')
            self.assertEqual(source.read_bytes(), original)

    @unittest.skipUnless(importlib.util.find_spec('xlrd') and importlib.util.find_spec('openpyxl') and shutil.which('libreoffice'),
                         'XLS integration requires xlrd, openpyxl and LibreOffice')
    def test_legacy_xls_with_zip_attachment_and_merged_cells(self):
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            book = Workbook()
            page = book.active
            page.title = '拍图需求'
            for address, value in mixed_cells().items():
                page[address] = value
            page.merge_cells('G13:G14')
            book.create_sheet('变体')
            source = root / 'brief.xlsx'
            book.save(source)
            subprocess.run(['libreoffice', '-env:UserInstallation=' + (root / 'lo-profile').as_uri(),
                            '--headless', '--convert-to', 'xls:MS Excel 97', '--outdir', str(root), str(source)],
                           check=True, capture_output=True, timeout=60)
            legacy = root / 'brief.xls'
            # OLE containers with embedded/appended ZIPs can pass is_zipfile.
            attachment = io.BytesIO()
            with zipfile.ZipFile(attachment, 'w') as archive:
                archive.writestr('example.txt', 'attachment, not workbook content')
            with legacy.open('ab') as stream:
                stream.write(attachment.getvalue())
            self.assertTrue(zipfile.is_zipfile(legacy))
            original = legacy.read_bytes()
            result = extract(legacy, root / 'out')
            self.assertEqual(result['source_format'], 'xls')
            self.assertEqual(len(result['sheets'][0]['task_row_candidates']), 11)
            self.assertEqual(result['sheets'][0]['merged_ranges'], ['G13:G14'])
            self.assertEqual(result['sheets'][1]['task_row_candidates'], [])
            self.assertIn('legacy_xls_media_not_extracted', [w['kind'] for w in result['warnings']])
            self.assertEqual(legacy.read_bytes(), original)


if __name__ == '__main__':
    unittest.main()
