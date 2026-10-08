"""The server warehouse path preserves Lingxing templates and converts status correctly."""
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parents[1]
SKILL = ROOT / 'plugins/chenyu-cangku/skills/chenyu-lingxing-luru'
SCRIPT = SKILL / 'scripts/generate_lingxing_imports.py'
ASSETS = SKILL / 'assets'
HEADERS = ['日期', '账号', '运营', '开发', '产品', '产品名称', 'Msku', '进货成本', '单个重量(净重）', '尺寸', '毛重（取最大者）', '采购价', '状态']


class WarehouseImportsTests(unittest.TestCase):
    def source(self, path: Path, *, duplicate=False):
        book = Workbook()
        sheet = book.active
        sheet.append(HEADERS)
        sheet.append(['2026-10-08', '久阅科技', '运营甲', '开发乙', 'SPU-1', '示例产品', 'MSKU-1｜', 12, 100, '10*20*30', 120, 15, None])
        sheet.append(['2026-10-09', '久阅科技', '运营甲', '开发乙', 'SPU-2', '另一产品', 'MSKU-1' if duplicate else 'MSKU-2', None, None, None, None, None, '待售'])
        book.save(path)

    def generate(self, root: Path):
        source = root / '新品补录.xlsx'
        product = root / '领星产品录用-V392.xlsx'
        pairing = root / '领星产品配对-按MSKU.xlsx'
        result = subprocess.run([sys.executable, str(SCRIPT), '--input', str(source), '--product-output', str(product), '--pairing-output', str(pairing)], capture_output=True, text=True)
        return result, product, pairing

    def test_two_outputs_preserve_templates_and_status(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            self.source(root / '新品补录.xlsx')
            result, product_path, pairing_path = self.generate(root)
            self.assertEqual(result.returncode, 0, result.stderr)
            original_product = load_workbook(ASSETS / 'Product-V392.xlsx')
            product = load_workbook(product_path)
            original_pairing = load_workbook(ASSETS / '导入配对商品模板 (按MSKU).xlsx')
            pairing = load_workbook(pairing_path)
            self.assertEqual(product.sheetnames, original_product.sheetnames)
            self.assertEqual(pairing.sheetnames, original_pairing.sheetnames)
            self.assertEqual(pairing['Sheet2'].sheet_state, 'hidden')
            self.assertEqual(len(product['产品'].data_validations.dataValidation), len(original_product['产品'].data_validations.dataValidation))
            self.assertEqual(len(pairing['Sheet1'].data_validations.dataValidation), len(original_pairing['Sheet1'].data_validations.dataValidation))
            headers = {cell.value: cell.column for cell in product['产品'][1] if cell.value}
            sheet = product['产品']
            self.assertEqual(sheet.cell(2, headers['*SKU']).value, original_product['产品'].cell(2, headers['*SKU']).value)
            self.assertEqual(sheet.cell(3, headers['*SKU']).value, 'MSKU-1')
            self.assertEqual(sheet.cell(3, headers['状态']).value, '在售')
            self.assertEqual(sheet.cell(4, headers['状态']).value, '待售')
            self.assertEqual(sheet.cell(3, headers['全部国家头程费用(含税)']).value, 36)
            self.assertEqual(pairing['Sheet1']['A2'].value, 'MSKU-1')
            self.assertEqual(pairing['Sheet1']['B2'].value, 'MSKU-1')
            self.assertEqual(pairing['Sheet1']['C2'].value, 'jiuyuekeji-FR')
            self.assertEqual(pairing['Sheet1']['E2'].value, '是')
            second, _, _ = self.generate(root)
            self.assertNotEqual(second.returncode, 0)
            self.assertIn('不能覆盖', second.stderr)

    def test_duplicate_msku_stops_before_writing_outputs(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            self.source(root / '新品补录.xlsx', duplicate=True)
            result, product, pairing = self.generate(root)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn('Msku 重复', result.stderr)
            self.assertFalse(product.exists())
            self.assertFalse(pairing.exists())


if __name__ == '__main__':
    unittest.main()
