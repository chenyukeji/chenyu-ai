#!/usr/bin/env python3
"""Generate a photo-derived logistics worksheet; unknown facts stay blank."""
from __future__ import annotations
import argparse
import json
import math
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = ["箱号/箱唛", "箱重（kg）", "长（cm）", "宽（cm）", "高（cm）",
           "产品英文品名*", "产品中文品名*", "产品申报单价*", "产品申报数量*",
           "产品材质*", "产品海关编码*", "产品用途*", "产品品牌*", "产品型号*",
           "产品销售链接", "产品销售价格", "产品图片"]
ITEM_FIELDS = ["nameEn", "nameZh", "declaredUnitPrice", "quantity", "material", "hsCode",
               "purpose", "brand", "model", "salesUrl", "salePrice", "imagePath"]

def number(value, label, integer=False):
    if value is None or value == "":
        return None
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value) or value <= 0 or (integer and int(value) != value):
        raise ValueError(f"{label}须为正{'整数' if integer else '数'}；不确定请留空")
    return value

def text_cell(ws, row, column, value):
    cell = ws.cell(row, column, value)
    if isinstance(value, str):
        cell.data_type = 's'  # OCR and user strings must never become formulas.
    return cell

def generate(data, output):
    cartons = data.get('cartons')
    if not isinstance(cartons, list) or not cartons:
        raise ValueError('至少需要一个已识别的箱及商品，不能交付空表')
    wb = Workbook()
    ws = wb.active
    ws.title = '物流表'
    pending = wb.create_sheet('待补信息')
    evidence = wb.create_sheet('识别依据')
    ws.append(HEADERS)
    pending.append(['物流表行号', '箱号/箱唛', '商品', '待补字段', '说明'])
    evidence.append(['物流表行号', '箱号/箱唛', '来源照片', '商品原文', 'SKU/MSKU', '补充资料/核对说明'])
    seen = set()
    for index, carton in enumerate(cartons, 1):
        carton_id = str(carton.get('cartonId') or f'照片箱{index:03d}')
        if carton_id in seen:
            raise ValueError(f'箱标识重复：{carton_id}，同箱多张照片应合并为一个箱')
        seen.add(carton_id)
        sources = carton.get('sourcePhotos')
        if not isinstance(sources, list) or not sources or not all(isinstance(s, str) and s.strip() for s in sources):
            raise ValueError(f'{carton_id}缺少来源照片')
        items = carton.get('items')
        if not isinstance(items, list) or not items:
            raise ValueError(f'{carton_id}没有可识别商品')
        dimensions = carton.get('dimensionsCm') or [None, None, None]
        if not isinstance(dimensions, list) or len(dimensions) != 3:
            raise ValueError('dimensionsCm 须为长宽高三个值；缺失位置填 null')
        box_values = [carton_id, number(carton.get('weightKg'), '箱重'), *[number(v, '尺寸') for v in dimensions]]
        start = ws.max_row + 1
        for item in items:
            if not isinstance(item, dict) or not str(item.get('nameZh') or item.get('rawText') or '').strip():
                raise ValueError('商品须有可追溯的品名或原文')
            row = ws.max_row + 1
            values = [item.get(key) for key in ITEM_FIELDS]
            values[1] = item.get("nameZh") or item.get("rawText")
            for key in ('quantity', 'declaredUnitPrice', 'salePrice'):
                values[ITEM_FIELDS.index(key)] = number(item.get(key), key, key == 'quantity')
            if item.get('hsCode') is not None and not isinstance(item['hsCode'], str):
                raise ValueError('海关编码须为字符串，保留前导零')
            # Product pictures are local, verified assets; carton photos are not product pictures.
            image_path = item.get('imagePath')
            values[-1] = None
            for column, value in enumerate(box_values + values, 1):
                text_cell(ws, row, column, value)
            if image_path:
                path = Path(image_path)
                if not path.is_file():
                    raise ValueError(f'产品图片不存在：{path}')
                picture = Image(str(path))
                ratio = min(85 / picture.width, 65 / picture.height, 1)
                picture.width *= ratio
                picture.height *= ratio
                ws.add_image(picture, f'Q{row}')
            ws.row_dimensions[row].height = 58
            missing = [HEADERS[i] for i, value in enumerate(box_values + values)
                       if (value is None or value == '') and not (i == 16 and image_path)]
            notes = [str(v) for v in [*(carton.get('notes') or []), *(item.get('notes') or [])] if v]
            if not carton.get('cartonId'):
                notes.append('箱号为照片临时编号，非 FBA 箱号')
            if missing or notes:
                for c, v in enumerate([row, carton_id, item.get('nameZh') or item.get('rawText'), '、'.join(missing), '；'.join(notes) or '照片及补充资料未提供，保留空白'], 1):
                    text_cell(pending, pending.max_row + 1 if c == 1 else pending.max_row, c, v)
            for c, v in enumerate([row, carton_id, '、'.join(sources), item.get('rawText') or item.get('nameZh'), item.get('msku') or item.get('sku'), item.get('evidence')], 1):
                text_cell(evidence, evidence.max_row + 1 if c == 1 else evidence.max_row, c, v)
        end = ws.max_row
        if end > start:
            for column in range(1, 6):
                ws.merge_cells(start_row=start, start_column=column, end_row=end, end_column=column)
    if pending.max_row == 1:
        pending.append([None, None, None, None, '未发现缺项；请核对申报资料后使用。'])
    thin = Side(style='thin', color='AAB4BF')
    for sheet in wb:
        sheet.freeze_panes = 'F2' if sheet == ws else 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        sheet.sheet_view.showGridLines = False
        sheet.row_dimensions[1].height = 46
        for cells in sheet:
            for cell in cells:
                cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
                if cell.row == 1:
                    cell.font = Font(name='Microsoft YaHei', bold=True, color='C00000' if '*' in str(cell.value) else '1F2937', size=11)
                    cell.fill = PatternFill('solid', fgColor='F3F4F6')
                else:
                    cell.fill = PatternFill('solid', fgColor='FFFF99' if sheet == ws else 'FFFFFF')
        for col in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(col)].width = 20 if sheet == ws else 30
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = '1:1'
    for col in ('B', 'C', 'D', 'E', 'H', 'I', 'P'):
        ws.column_dimensions[col].width = 13
    ws.column_dimensions['O'].width = 35
    for row in ws.iter_rows(min_row=2):
        row[10].number_format = '@'
        for index in (1, 2, 3, 4, 7, 15):
            row[index].number_format = '0.00'
        row[8].number_format = '0'
    output = Path(output)
    if output.suffix.lower() != '.xlsx':
        raise ValueError('照片物流表须输出 .xlsx')
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    wb.close()
    with output.open('rb') as stream:
        check = load_workbook(stream)
        assert check['物流表'].max_row > 1
        check.close()
    return output

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', required=True, type=Path)
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    generate(json.loads(args.input.read_text(encoding='utf-8')), args.output)
    print(args.output)

if __name__ == '__main__':
    main()
