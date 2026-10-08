#!/usr/bin/env python3
"""Fill the two official Lingxing workbooks from one 新品补录 workbook."""

from __future__ import annotations

import argparse
import json
import re
import unicodedata
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


ASSETS = Path(__file__).resolve().parents[1] / "assets"
STORES = {"久阅科技": ("jiuyuekeji-FR", "法国")}
ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}


def clean(value: object) -> str:
    return re.sub(r"\s+", " ", str(value).replace("\r", " ").replace("\n", " ")).strip() if value is not None else ""


def canonical(value: object) -> str:
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", clean(value))).casefold()


def column(headers: tuple, *names: str, required: bool = True) -> int | None:
    lookup = {canonical(value): index for index, value in enumerate(headers) if clean(value)}
    for name in names:
        if canonical(name) in lookup:
            return lookup[canonical(name)]
    if required:
        raise ValueError(f"缺少必需列：{names[0]}")
    return None


def number(value: object, label: str, row: int) -> Decimal | None:
    if not clean(value):
        return None
    try:
        result = Decimal(clean(value).replace(",", ""))
    except InvalidOperation as exc:
        raise ValueError(f"第 {row} 行“{label}”不是有效的非负数字：{value}") from exc
    if not result.is_finite() or result < 0:
        raise ValueError(f"第 {row} 行“{label}”不是有效的非负数字：{value}")
    return result


def normalized_date(value: object, row: int) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        return (date(1899, 12, 30) + timedelta(days=round(value))).isoformat()
    match = re.fullmatch(r"(\d{4})[.\-/年](\d{1,2})[.\-/月](\d{1,2})日?", clean(value))
    if not match:
        raise ValueError(f"第 {row} 行“日期”格式无法识别：{value}")
    try:
        return date(*(int(part) for part in match.groups())).isoformat()
    except ValueError as exc:
        raise ValueError(f"第 {row} 行“日期”不是有效日期：{value}") from exc


def dimensions(value: object, row: int) -> tuple[Decimal, Decimal, Decimal] | None:
    if not clean(value):
        return None
    parts = re.split(r"\s*[*×xX]\s*", clean(value))
    if len(parts) != 3:
        raise ValueError(f"第 {row} 行“尺寸”必须是 长*宽*高：{value}")
    try:
        result = tuple(Decimal(part) for part in parts)
    except InvalidOperation as exc:
        raise ValueError(f"第 {row} 行“尺寸”必须是 长*宽*高：{value}") from exc
    if any(not part.is_finite() or part <= 0 for part in result):
        raise ValueError(f"第 {row} 行“尺寸”必须是 长*宽*高：{value}")
    return result


def source_records(workbook) -> list[dict]:
    sheet = next((item for item in workbook if any(canonical(cell.value) == "msku" for cell in item[1])), None)
    if sheet is None:
        raise ValueError("输入文件中找不到包含 Msku 列的新品补录工作表。")
    headers = tuple(cell.value for cell in sheet[1])
    indexes = {
        "date": column(headers, "日期"),
        "account": column(headers, "账号"),
        "operator": column(headers, "运营"),
        "developer": column(headers, "开发"),
        "product": column(headers, "产品"),
        "product_name": column(headers, "产品名称"),
        "msku": column(headers, "Msku", "MSKU"),
        "cost": column(headers, "进货成本"),
        "net": column(headers, "单个重量(净重）", "单个重量(净重)"),
        "size": column(headers, "尺寸"),
        "gross": column(headers, "毛重（取最大者）", "毛重(取最大者)"),
        "status": column(headers, "状态", required=False),
        "price": column(headers, "采购价", required=False),
    }
    records = []
    seen = set()
    for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
        if not any(clean(value) for value in values):
            continue
        def get(key: str):
            index = indexes[key]
            return values[index] if index is not None else None

        msku = clean(get("msku")).rstrip("|｜").strip()
        if not msku:
            raise ValueError(f"第 {row_number} 行有数据但 Msku 为空。")
        if msku in seen:
            raise ValueError(f"Msku 重复：{msku}（第 {row_number} 行）")
        seen.add(msku)
        record = {"msku": msku, "date": normalized_date(get("date"), row_number)}
        for key, label in (("account", "账号"), ("operator", "运营"), ("developer", "开发"),
                           ("product", "产品"), ("product_name", "产品名称")):
            record[key] = clean(get(key))
            if not record[key]:
                raise ValueError(f"第 {row_number} 行“{label}”为空。")
        if record["account"] not in STORES:
            raise ValueError(f"第 {row_number} 行账号“{record['account']}”没有店铺映射，请先在 Skill 规则中确认。")
        record.update(
            store=STORES[record["account"]],
            status=clean(get("status")) or "在售",
            cost=number(get("cost"), "进货成本", row_number),
            price=number(get("price"), "采购价", row_number),
            net=number(get("net"), "单个重量(净重)", row_number),
            gross=number(get("gross"), "毛重(取最大者)", row_number),
            size=dimensions(get("size"), row_number),
        )
        records.append(record)
    if not records:
        raise ValueError("输入文件没有可转换的产品记录。")
    return records


def first_blank_row(sheet, columns: int) -> int:
    for row in range(2, sheet.max_row + 2):
        if not any(clean(sheet.cell(row, col).value) for col in range(1, columns + 1)):
            return row
    return sheet.max_row + 1


def fill_product(sheet, records: list[dict]) -> None:
    headers = tuple(cell.value for cell in sheet[1])
    if sum(bool(clean(header)) for header in headers) < 90:
        raise ValueError("领星产品模板的“产品”表头不完整。")
    lookup = {canonical(header): index for index, header in enumerate(headers, 1) if clean(header)}
    start = first_blank_row(sheet, len(headers))
    if start + len(records) - 1 > sheet.max_row:
        raise ValueError("领星产品模板预留行不足，不能超出原母版格式。")
    for row, record in enumerate(records, start):
        size = record["size"]
        freight = max(Decimal("2"), size[0] * size[1] * size[2] * Decimal("0.006")).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP) if size else None
        fields = {
            "*SKU": record["msku"], "品名": record["product_name"], "产品类型": "普通产品",
            "状态": record["status"], "单位": "件", "开发人": record["developer"],
            "产品负责人": record["operator"], "SPU": record["product"],
            "采购成本(CNY)": record["cost"],
            "采购备注": f"原始采购价：{record['price']:.2f} CNY" if record["price"] is not None else None,
            "单品净重": record["net"], "单品净重单位": "g" if record["net"] is not None else None,
            "单品毛重": record["gross"], "单品毛重单位": "g" if record["gross"] is not None else None,
            "包装规格长": size[0] if size else None, "包装规格宽": size[1] if size else None,
            "包装规格高": size[2] if size else None, "包装规格单位": "cm" if size else None,
            "全部国家头程费用(含税)": freight,
            "全部国家头程费用币种": "CNY" if freight is not None else None,
        }
        if canonical("开发日期") in lookup:
            fields["开发日期"] = record["date"]
        for name, value in fields.items():
            index = lookup.get(canonical(name))
            if index is None:
                raise ValueError(f"领星产品模板缺少字段：{name}")
            if value is not None:
                sheet.cell(row, index).value = float(value) if isinstance(value, Decimal) else value


def fill_pairing(sheet, records: list[dict]) -> None:
    headers = [clean(sheet.cell(1, col).value) for col in range(1, 6)]
    expected = ["*SKU", "*MSKU", "店铺名称", "国家", "是否同步listing图"]
    if headers != expected:
        raise ValueError(f"配对模板表头不匹配：{headers}")
    start = first_blank_row(sheet, 5)
    if start + len(records) - 1 > sheet.max_row:
        raise ValueError("领星配对模板预留行不足，不能超出原母版格式。")
    for row, record in enumerate(records, start):
        for column_index, value in enumerate((record["msku"], record["msku"], *record["store"], "是"), 1):
            sheet.cell(row, column_index).value = value


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True)
    parser.add_argument("--product-template", type=Path, default=ASSETS / "Product-V392.xlsx")
    parser.add_argument("--pairing-template", type=Path, default=ASSETS / "导入配对商品模板 (按MSKU).xlsx")
    parser.add_argument("--product-output", type=Path, required=True)
    parser.add_argument("--pairing-output", type=Path, required=True)
    args = parser.parse_args()
    outputs = (args.product_output.resolve(), args.pairing_output.resolve())
    if outputs[0] == outputs[1]:
        raise ValueError("两份输出文件必须使用不同路径。")
    if any(path.exists() for path in outputs):
        raise ValueError("输出文件已存在，不能覆盖。")
    records = source_records(load_workbook(args.input, read_only=True, data_only=True))
    product = load_workbook(args.product_template)
    pairing = load_workbook(args.pairing_template)
    fill_product(product["产品"], records)
    fill_pairing(pairing["Sheet1"], records)
    for label, workbook in (("领星产品录用表", product), ("领星产品配对表", pairing)):
        for sheet in workbook:
            for row in sheet:
                for cell in row:
                    if cell.data_type == "e" or (isinstance(cell.value, str) and cell.value in ERRORS):
                        raise ValueError(f"{label} 存在公式错误：{sheet.title}!{cell.coordinate}")
    for path in outputs:
        path.parent.mkdir(parents=True, exist_ok=True)
    try:
        product.save(outputs[0])
        pairing.save(outputs[1])
    except Exception:
        for path in outputs:
            path.unlink(missing_ok=True)
        raise
    print(json.dumps({"records": len(records), "productOutput": str(outputs[0]), "pairingOutput": str(outputs[1])}, ensure_ascii=False))


if __name__ == "__main__":
    main()
