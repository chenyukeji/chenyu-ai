"""Read tabular evidence without guessing the historical period from download time."""
from __future__ import annotations

import csv
import json
from datetime import date, datetime
from pathlib import Path

ALIASES = {
    "站点": "marketplace", "ASIN": "asin", "产品名称": "product_name", "产品标题": "product_name",
    "月份": "history_month", "历史月份": "history_month", "月销量": "estimated_sales",
    "销量(父)": "estimated_sales", "月销量(父)": "estimated_sales",
    "预估月销量": "estimated_sales", "月销量增长率": "sales_growth_percent",
    "月销量环比增长率": "sales_growth_percent", "增长率": "sales_growth_percent",
    "搜索词": "keyword", "关键词": "keyword", "搜索量": "search_volume", "月搜索量": "search_volume",
    "上期搜索量": "previous_search_volume", "数据周期": "search_period", "上期周期": "previous_search_period",
    "卖家ID": "seller_id", "店铺ID": "seller_id", "快照日期": "snapshot_date", "快照完整": "snapshot_complete",
    "当前排名": "current_rank", "此前排名": "previous_rank", "榜单类目": "rank_category",
    "统计窗口小时数": "window_hours", "观察时间": "observed_at", "来源": "source_ref",
    "来源链接": "source_ref", "图片": "image_url", "类目": "category_name", "售价": "price",
    "评论数": "review_count", "BSR": "bsr",
}
CONTEXT_FIELDS = {"marketplace", "history_month", "seller_id", "snapshot_date", "snapshot_complete", "search_period", "previous_search_period", "window_hours"}


def _row(values, headers, source, context):
    result = dict(context)
    for key, value in zip(headers, values):
        if not key or value is None or value == "":
            continue
        if isinstance(value, (date, datetime)):
            value = value.isoformat()
        key = ALIASES.get(str(key).strip(), str(key).strip())
        if key == "snapshot_complete" and isinstance(value, str):
            value = value.strip().lower() in {"true", "yes", "1", "是", "完整"}
        if key == "marketplace":
            value = {"美国": "US", "美国站": "US", "德国": "DE", "德国站": "DE"}.get(str(value).strip(), str(value).strip().upper())
        result[key] = value
    result.setdefault("source_ref", source)
    return result


def read_input(spec):
    context = {}
    if isinstance(spec, dict):
        path = Path(spec["path"]).expanduser().resolve()
        context = {key: value for key, value in spec.items() if key != "path"}
        if set(context) - CONTEXT_FIELDS:
            raise ValueError("unsupported discovery path context")
    else:
        path = Path(spec).expanduser().resolve()
    if path.suffix.lower() == ".json":
        content = json.loads(path.read_text(encoding="utf-8-sig"))
        rows = content if isinstance(content, list) else content.get("records")
        if isinstance(rows, list):
            normalized = []
            for index, row in enumerate(rows, 1):
                if not isinstance(row, dict):
                    raise ValueError("each input record must be an object")
                normalized.append(_row(row.values(), row.keys(), f"{path}#row={index}", context))
            content = normalized if isinstance(content, list) else {**content, "records": normalized}
        return content, str(path)
    rows = []
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as stream:
            sample = stream.read(4096)
            stream.seek(0)
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
            except csv.Error:
                dialect = csv.excel
            reader = csv.reader(stream, dialect)
            headers = next(reader, [])
            for index, values in enumerate(reader, 2):
                if any(value not in (None, "") for value in values):
                    rows.append(_row(values, headers, f"{path}#row={index}", context))
    elif path.suffix.lower() == ".xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("读取 Excel 需要安装插件 requirements-browser.txt 中的 openpyxl") from exc
        workbook = load_workbook(path, read_only=True, data_only=True)
        try:
            for sheet in workbook:
                iterator = sheet.iter_rows()
                headers = [cell.value for cell in next(iterator, [])]
                for index, cells in enumerate(iterator, 2):
                    values = []
                    for header, cell in zip(headers, cells):
                        value = cell.value
                        field = ALIASES.get(str(header).strip(), str(header).strip())
                        if field == "sales_growth_percent" and isinstance(value, (int, float)) and "%" in cell.number_format:
                            value *= 100
                        values.append(value)
                    if any(value is not None for value in values):
                        rows.append(_row(values, headers, f"{path}#sheet={sheet.title}&row={index}", context))
        finally:
            workbook.close()
    else:
        raise ValueError("discovery_paths 支持 JSON、CSV、XLSX；其他格式请先规范化")
    return {"records": rows}, str(path)
