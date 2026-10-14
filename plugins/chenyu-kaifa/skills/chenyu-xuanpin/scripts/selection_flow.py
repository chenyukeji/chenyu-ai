"""Unified, resumable selection orchestration for one skill and multiple sources."""
from __future__ import annotations

from datetime import date
import hashlib
from io import BytesIO
import json
import math
import os
from pathlib import Path
import re

from category_sources import normalize_url
from selection_analysis import analyze_keywords, merge_products, summarize_products, balanced_shortlist, group_product_families, product_identity
import selection_collectors as collectors
from product_evidence import history_is_sufficient, ANALYSIS_REQUIREMENTS, analysis_columns
from strategy_h import target_months

SOURCES = ('amazon_new_releases', 'sellersprite_new', 'sellersprite_surge',
           'sellersprite_history', 'keyword_trends', 'new_releases_db')
DEFAULT_SOURCES = tuple(source for source in SOURCES if source != 'sellersprite_history')
VERSION = 'selection-v1'


def query_id(spec):
    return hashlib.sha256(json.dumps(spec, ensure_ascii=False, sort_keys=True).encode()).hexdigest()[:16]


def plan_selection(payload):
    request = str(payload.get('request') or '').strip()
    scope = dict(payload.get('scope') or {})
    if set(scope)-{'input_urls','marketplaces','sources','categories','following_months','shortlist_limit','review_batch_size','filters','exclusions'}:
        raise ValueError('scope 包含不支持的字段')
    filters=scope.get('filters') or {}
    if set(filters)-{'price_min','price_max','review_max','sales_min','fulfillment'}:
        raise ValueError('filters 包含不支持的筛选条件')
    for key,value in filters.items():
        if key=='fulfillment':
            if value not in ('FBM','FBA','AMZ'): raise ValueError('配送方式仅支持 FBM/FBA/AMZ')
        elif isinstance(value,bool) or not isinstance(value,(int,float)) or not math.isfinite(value) or value<0:
            raise ValueError('数值筛选条件须为非负数字')
    links = scope.get('input_urls') or re.findall(r'https://[^\s<>"，。]+', request)
    if not isinstance(links,list) or any(not isinstance(x,str) for x in links):
        raise ValueError('input_urls 必须为链接数组')
    identities = [collectors.amazon_identity(url.rstrip(').）')) for url in links]
    markets = scope.get('marketplaces') or list(dict.fromkeys(x['marketplace'] for x in identities))
    if not markets:
        mentioned = [m for m, words in {'US': ('美国', '美站'), 'DE': ('德国', '德站')}.items()
                     if any(word in request for word in words)]
        markets = mentioned or ['US', 'DE']
    if not isinstance(markets, list) or any(m not in ('US', 'DE') for m in markets):
        raise ValueError('scope.marketplaces 仅支持 US/DE')
    markets = list(dict.fromkeys(markets))
    if any(x['marketplace'] not in markets for x in identities):
        raise ValueError('链接站点不在目标站点中；请先明确跨站目录映射')
    sources = scope.get('sources', list(DEFAULT_SOURCES))
    if not isinstance(sources, list) or not sources or set(sources) - set(SOURCES):
        raise ValueError('scope.sources 包含不支持的数据来源')
    as_of = date.fromisoformat(payload.get('as_of_date') or str(date.today())).isoformat()
    months = target_months(as_of, scope.get('following_months', 2))
    categories = scope.get('categories') or []
    if not isinstance(categories, list):
        raise ValueError('scope.categories 必须为数组')
    queries, missing = [], []
    for identity in identities:
        queries.append({'source': 'amazon_input', **identity})
    if not categories:
        missing.append('请模型根据用户类目或链接事实，填入 scope.categories 和各站点关键词/已核实榜单入口')
    for category in categories:
        if not isinstance(category,dict): raise ValueError('类目必须为对象')
        name = category.get('name')
        if not isinstance(name, str) or not name.strip():
            raise ValueError('类目必须包含非空 name')
        for market in markets:
            terms = (category.get('keywords') or {}).get(market, [])
            if isinstance(terms, str): terms = [terms]
            if not isinstance(terms, list) or any(not isinstance(x, str) or not x.strip() or len(x)>100 for x in terms):
                raise ValueError('每站点 keywords 必须是非空短词数组')
            terms = list(dict.fromkeys(x.strip() for x in terms))
            if any(x in sources for x in ('sellersprite_new','sellersprite_surge','sellersprite_history','keyword_trends')) and not terms:
                missing.append(f'{name}/{market}：缺少已核对的本地语言关键词')
            for source in sources:
                base = {'source': source, 'marketplace': market, 'category': name, 'as_of_date': as_of}
                if source == 'amazon_new_releases':
                    urls = (category.get('amazon_new_releases') or {}).get(market) or []
                    if isinstance(urls, str): urls = [urls]
                    if not urls:
                        queries.append({**base, 'unavailable': '尚未找到并核实对应目录的新品榜入口'})
                    for url in urls:
                        queries.append({**base, 'url': normalize_url(url, market)})
                elif source == 'new_releases_db':
                    mapped = (category.get('db_categories') or {}).get(market) or []
                    if isinstance(mapped, str): mapped = [mapped]
                    for db_category in mapped or [None]:
                        queries.append({**base, 'db_category': db_category})
                else:
                    for term in terms:
                        for month in months if source == 'sellersprite_history' else [None]:
                            spec = {**base, 'keyword': term}
                            if month: spec['history_month'] = month
                            queries.append(spec)
    limit = scope.get('shortlist_limit', 50)
    if limit is not None and (type(limit) is not int or limit < 1):
        raise ValueError('shortlist_limit 必须为正整数或 null（审阅全部已采集父体）')
    batch_size = scope.get('review_batch_size', 50)
    if type(batch_size) is not int or not 1 <= batch_size <= 200:
        raise ValueError('review_batch_size 必须为 1–200；仅控制单批审阅量，不限制总量')
    unique = {query_id(q): q for q in queries}
    if len(unique) > 120:
        raise ValueError('查询任务超过 120 项，请按类目分批调查')
    return {'version': VERSION, 'request': request, 'scope': scope, 'as_of_date': as_of,
            'marketplaces': markets, 'historical_months': months, 'shortlist_limit': limit, 'review_batch_size': batch_size,
            'queries': [{'id': key, **value} for key, value in unique.items()], 'blocking_items': missing}


def write_json(path, value):
    temporary = path.with_suffix(path.suffix + '.tmp')
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')
    temporary.replace(path)


def collect_query(query, payload, known_products):
    if query.get('unavailable'):
        return {'collection_status': 'unavailable', 'records': [], 'block_reason': query['unavailable']}
    options = {**(payload.get('browser') or {}), **query}
    source = query['source']
    if source == 'amazon_input': return collectors.collect_amazon_input(options)
    if source == 'amazon_new_releases': return collectors.collect_amazon_rankings(options)
    if source in ('sellersprite_new','sellersprite_surge','sellersprite_history'):
        return collectors.collect_sellersprite_products(options)
    if source == 'keyword_trends': return collectors.collect_keyword_trends(options)
    if source == 'new_releases_db':
        options['db_path'] = payload.get('db_path') or os.environ.get('CHENYU_NEW_RELEASES_DB')
        if not options.get('db_category'):
            options['asins'] = [x['asin'] for x in known_products if x['marketplace'] == query['marketplace']]
        return collectors.collect_database(options)
    raise ValueError('unknown selection source')



def original_style_rows(rows, report):
    """Put combined evidence in the familiar 15-column selection table."""
    products = {(p['marketplace'], p['asin']): p for p in report.get('families', report['products'])}
    result = []
    for row in rows:
        current = (products.get((row['站点'], row['ASIN'])) or {}).get('current') or {}
        analysis = str(row.get('分析') or '').strip()
        summary = re.split(r'[。；;]', analysis, maxsplit=1)[0].strip()
        if not re.search(r'[\u3400-\u9fff]', summary):
            summary = ''
        item = products.get((row['站点'], row['ASIN'])) or {'marketplace': row['站点'], 'asin': row['ASIN']}
        review = next((r for r in report.get('reviews', [])
                       if (r.get('marketplace'), r.get('asin')) == (row['站点'], row['ASIN'])), {})
        evidence = analysis_columns(item, review, report.get('plan', {}).get('as_of_date'))
        priority = row.get('优先级') or ''
        conclusion = {'优先研究': '🟢 优先研究', '继续观察': '🟡 继续观察',
                      '低优先级': '🔴 低优先级', '待补数据': '🟡 待补数据',
                      '待核对相关性': '🟡 待核对相关性'}.get(priority, priority)
        reasons = [analysis, *evidence['analysis_notes']]
        if evidence['analysis_missing']:
            reasons.append("分析缺项：" + "；".join(evidence['analysis_missing']))
        for label, key in (('命中来源', '命中来源'), ('历史命中月份', '历史命中月份'),
                           ('关联关键词', '关联关键词'), ('证据缺项', '证据缺项')):
            if row.get(key):
                reasons.append(f'{label}：{row[key]}')
        result.append({'站点': row['站点'], '上架日期': row.get('上架日期'),
                       'Review数量': row.get('评论量'), '售价（当地币种）': row.get('售价'),
                       '大品类排名': current.get('bsr'), '所在品类': row.get('所在品类'),
                       '预估月销量': row.get('当前预估月销量'), '产品优点&特征': summary,
                       '缺点': evidence['缺点'], '生命周期': evidence['生命周期'], 'ASIN': row['ASIN'],
                       '亚马逊产品链接': row.get('亚马逊产品链接'), '图片': row.get('图片'),
                       '结论': conclusion, '理由': '\n'.join(x for x in reasons if x)})
    return result


def export_workbook(path, rows, report, image_loader=None):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.drawing.image import Image
    from openpyxl.utils import get_column_letter
    from discovery_workbook import HEADERS, _download_image
    loader = image_loader or _download_image
    wb = Workbook(); wb.remove(wb.active)
    main_rows = original_style_rows(rows, report)
    image_cache = {}; errors = []
    # Reuse images from the prior workbook by marketplace/ASIN on a resumed task.
    previous_images = {}
    if path.exists():
        previous = load_workbook(path)
        previous_title = next((name for name in ('开品结果', '选品结果') if name in previous.sheetnames), None)
        if previous_title:
            sheet = previous[previous_title]; headers = [c.value for c in sheet[1]]
            if '站点' in headers and 'ASIN' in headers:
                for image in sheet._images:
                    n = image.anchor._from.row + 1
                    previous_images[(sheet.cell(n, headers.index('站点')+1).value,
                                     sheet.cell(n, headers.index('ASIN')+1).value)] = image._data()
        previous.close()
    sheets = [('开品结果', main_rows)]
    border = Border(*(Side(style='thin', color='666666') for _ in range(4)))
    main_widths = dict(zip(HEADERS, [8, 13, 11, 15, 12, 16, 13, 28, 27, 24, 14, 31, 22, 15, 65]))
    for title, values in sheets:
        ws = wb.create_sheet(title)
        headers = HEADERS if title == '开品结果' else list(dict.fromkeys(k for row in values for k in row)) or ['说明']
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(name='Microsoft YaHei', size=11, bold=True)
            cell.fill = PatternFill('solid', fgColor='C6E0B4')
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border
        ws.row_dimensions[1].height = 32
        for row in values:
            ws.append([None if key=='图片' else row.get(key) for key in headers])
            for cell in ws[ws.max_row]:
                if isinstance(cell.value, str): cell.data_type = 's'
                cell.font = Font(name='Microsoft YaHei', size=10)
                cell.alignment = Alignment(horizontal='left' if title == '开品结果' and cell.column == 15 else 'center', vertical='center', wrap_text=True)
                cell.border = border
            if title == '开品结果':
                ws.row_dimensions[ws.max_row].height = 110
                conclusion_cell = ws.cell(ws.max_row, headers.index('结论') + 1)
                conclusion = str(conclusion_cell.value or '')
                color = '33A36B' if '优先研究' in conclusion else 'F4CCCC' if '低优先级' in conclusion else 'FFD966'
                conclusion_cell.fill = PatternFill('solid', fgColor=color)
                conclusion_cell.font = Font(name='Microsoft YaHei', size=10, bold=True,
                                            color='FFFFFF' if color == '33A36B' else '000000')
            if '图片' in headers and row.get('图片'):
                try:
                    url = row['图片']; identity = (row.get('站点'), row.get('ASIN'))
                    if identity in previous_images: image_cache[url] = previous_images[identity]
                    if url not in image_cache: image_cache[url] = loader(url)[0]
                    img = Image(BytesIO(image_cache[url])); ratio = min(110/img.width,85/img.height)
                    img.width *= ratio; img.height *= ratio
                    ws.add_image(img, f'{get_column_letter(headers.index("图片")+1)}{ws.max_row}')
                    if title != '开品结果': ws.row_dimensions[ws.max_row].height = 68
                except Exception:
                    errors.append({'marketplace': row.get('站点'), 'asin': row.get('ASIN'), 'reason': 'image_embedding_failed'})
                    ws.cell(ws.max_row,headers.index('图片')+1,'图片未取得')
            if '亚马逊产品链接' in headers and row.get('亚马逊产品链接'):
                ws.cell(ws.max_row,headers.index('亚马逊产品链接')+1).hyperlink = row['亚马逊产品链接']
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for i, header in enumerate(headers,1):
            ws.column_dimensions[get_column_letter(i)].width = (main_widths[header] if title == '开品结果'
                else 42 if header in ('分析','说明','统计上下文','产品名称') else 22)
    temporary = path.with_suffix('.tmp.xlsx'); wb.save(temporary); temporary.replace(path)
    return {'path': str(path), 'row_count': len(rows), 'image_failures': errors}


def run_selection(payload):
    raw_dir = payload.get('run_dir')
    if not raw_dir: raise ValueError('run_selection 需要按输出目录规范传入 run_dir')
    folder = Path(raw_dir).expanduser().resolve(); folder.mkdir(parents=True,exist_ok=True)
    plan_path = folder/'selection-plan.json'
    if plan_path.exists() and not payload.get('request') and not payload.get('scope'):
        plan=json.loads(plan_path.read_text())
        if plan.get('version')!=VERSION: raise ValueError('不能续跑不兼容的选品目录')
    else:
        plan=plan_selection(payload); write_json(plan_path,plan)
    settings_path=folder/'selection-settings.json'
    settings=json.loads(settings_path.read_text()) if settings_path.exists() else {}
    for key in ('browser','db_path','enrich','associations'):
        if key in payload: settings[key]=payload[key]
    browser_options=settings.get('browser') or {}
    if set(browser_options)-{'profile_dir','headless','max_pages','limit','navigation_timeout_ms','query_timeout_ms','query_delay_ms'}:
        raise ValueError('browser 包含不支持的字段；账号密码仅通过环境变量传入')
    write_json(settings_path,settings)
    payload={**settings,**payload}
    options = payload.get('browser') or {}
    budget = payload.get('max_queries', 20)
    if type(budget) is not int or not 1 <= budget <= 120: raise ValueError('max_queries 必须为 1–120')
    results=[]; pending=[]; executed=0; auth_block=None
    for query in plan['queries']:
        digest=query_id({'collector_version': 2, 'query':query,'browser':{k:v for k,v in options.items() if k not in ('profile_dir','headless')},
                         'db_path':payload.get('db_path') or os.environ.get('CHENYU_NEW_RELEASES_DB')})
        cache=folder/f'source-{digest}.json'
        result=json.loads(cache.read_text()) if cache.exists() and not payload.get('refresh') else None
        # Review the exact evidence shown to the model; retry failures on a
        # query-only resume or explicit refresh, not while submitting reviews.
        if result and result.get('collection_status') in ('blocked','unavailable','pending') and not payload.get('reviews'):
            result=None
        if result is None:
            if executed>=budget:
                pending.append(query['id']); continue
            try:
                if auth_block and query['source'] in ('sellersprite_new','sellersprite_surge','sellersprite_history','keyword_trends'):
                    result={'collection_status':'blocked','block_reason':auth_block,'records':[]}
                else:
                    result=collect_query(query,payload,merge_products(results)); executed+=1
            except Exception as exc:
                # Library exceptions can include URLs or account-page text. Keep only safe class names.
                reason=str(exc) if isinstance(exc,(ValueError,collectors.browser.BrowserCollectionError)) else type(exc).__name__
                result={'collection_status':'blocked','block_reason':reason[:240],'records':[]}; executed+=1
            reason=result.get('block_reason') or ''
            if query['source'] in ('sellersprite_new','sellersprite_surge','sellersprite_history','keyword_trends') and any(
                term in reason for term in ('login','session_lost','verification_required','authentication','permission_required')):
                auth_block=reason
            result['query']=query; write_json(cache,result)
        results.append(result)
        if result.get('product_records'):
            results.append({'query':{'id':query['id']+'-products','source':'keyword_products','marketplace':query.get('marketplace')},
                            'records':result['product_records'],'collection_status':'complete',
                            'coverage_note':'关键词页面实际展示的关联商品样本，相关性待模型核对'})
    associations=[a for result in results for a in result.get('associations',[])]+(payload.get('associations') or [])
    # Keep all verified families available. Batch size is a context limit, not an output cap.
    base_results = list(results)
    source_ids = [query_id(result) for result in base_results]
    enrichment_context = query_id(source_ids)
    for cache in sorted(folder.glob('enrichment-*.json')):
        cached = json.loads(cache.read_text())
        if cached.get('selection_context') == enrichment_context and not payload.get('refresh'):
            results.append(cached)

    context_plan = {**plan, 'scope': {k: v for k, v in plan['scope'].items()
                                    if k not in ('shortlist_limit', 'review_batch_size')}}
    context_plan.pop('shortlist_limit', None); context_plan.pop('review_batch_size', None)
    review_context = query_id(context_plan)
    review_path = folder/'selection-reviews.json'
    stored = json.loads(review_path.read_text()) if review_path.exists() else {}
    stored = stored if isinstance(stored, dict) and stored.get('context_id') == review_context else {}

    def prepare():
        raw = merge_products(results, associations=associations)
        families = group_product_families(raw)
        ordered = balanced_shortlist(families, historical_months=plan['historical_months']
                                    if any(q['source'] == 'sellersprite_history' for q in plan['queries']) else None)
        for item in ordered:
            checks = []
            for key, value in (plan['scope'].get('filters') or {}).items():
                field = {'price_min':'price','price_max':'price','review_max':'review_count',
                         'sales_min':'estimated_sales','fulfillment':'fulfillment'}[key]
                actual = item['current'].get(field)
                if actual is None:
                    checks.append({'condition':key,'status':'missing'})
                else:
                    passed = actual == value if key == 'fulfillment' else actual >= value if key.endswith('_min') else actual <= value
                    checks.append({'condition':key,'status':'pass' if passed else 'fail','actual':actual,'expected':value})
            item['filter_checks'] = checks
        return raw, ordered

    def saved_reviews(items):
        fingerprints = stored.get('fingerprints', {})
        known = {product_identity(item): item for item in items}
        return [review for review in stored.get('reviews', [])
                if (item := known.get((review.get('marketplace'), review.get('asin')))) is not None
                and fingerprints.get('|'.join(product_identity(item))) == query_id(item)
                and (item.get('verified_parent_asin') or payload.get('reviews'))]

    def pending_batch(items, reviews):
        covered = {(r['marketplace'], r['asin']) for r in reviews}
        eligible = [item for item in items if not any(c['status'] == 'fail' for c in item['filter_checks'])]
        eligible_ids = {product_identity(p) for p in eligible if p.get('verified_parent_asin')}
        accepted = sum(r['relevant'] for r in reviews if (r['marketplace'], r['asin']) in eligible_ids)
        target = plan.get('shortlist_limit')
        remaining = None if target is None else max(0, target - accepted)
        waiting = [item for item in eligible if product_identity(item) not in covered]
        return waiting[:min(plan.get('review_batch_size', 50), remaining) if remaining is not None else plan.get('review_batch_size', 50)]

    products, shortlisted = prepare()
    reviews = saved_reviews(shortlisted)
    submitted = payload.get('reviews') or []
    known = {product_identity(p) for p in shortlisted}
    if any((r.get('marketplace'),r.get('asin')) not in known or type(r.get('relevant')) is not bool
           or not r.get('reason') for r in submitted):
        raise ValueError('reviews 必须对应本次父体代表商品，包含 relevant 布尔值和证据理由')
    for review in submitted:
        if review['relevant']:
            item = next(p for p in shortlisted if product_identity(p) == (review['marketplace'], review['asin']))
            analysis_columns(item, review, plan['as_of_date'])
    if submitted and payload.get('evidence_id') is not None:
        current_keywords = analyze_keywords([row for result in results
            if result['query']['source'] == 'keyword_trends' for row in result.get('records', [])], plan['as_of_date'])
        current_evidence = query_id({'context': review_context, 'products': shortlisted, 'keywords': current_keywords})
        if current_evidence != payload['evidence_id']:
            raise ValueError('审阅证据已变更，请读取新的 analysis_requests')
    review_index = {(r['marketplace'], r['asin']): r for r in reviews}
    review_index.update({(r['marketplace'], r['asin']): r for r in submitted})
    reviews = list(review_index.values())
    before_enrichment = {product_identity(item): query_id(item) for item in shortlisted}
    # Enrich only the next review batch. Its cached facts survive all later batches.
    batch = pending_batch(shortlisted, reviews)
    if payload.get('enrich', True) and batch and not pending and not plan['blocking_items'] and not auth_block:
        for market in plan['marketplaces']:
            asins = [p['asin'] for p in batch if p['marketplace'] == market and
                     (not p.get('verified_parent_asin') or
                      any(p['current'].get(f) is None for f in ('price','estimated_sales','source_available_date','review_count')))]
            if not asins:
                continue
            spec = {'source':'sellersprite_enrichment','marketplace':market,'asins':asins}
            cache = folder/f'enrichment-{query_id({"context": enrichment_context, **spec})}.json'
            result = json.loads(cache.read_text()) if cache.exists() and not payload.get('refresh') else None
            parent_missing = any(not p.get('verified_parent_asin') for p in batch if p['marketplace'] == market)
            if not result or result.get('collection_status') != 'complete' or parent_missing:
                try:
                    result = collectors.browser.collect_sellersprite_by_asin({**options,'marketplace':market,'asins':asins,'headless':True})
                except Exception as exc:
                    result = {'collection_status':'blocked','block_reason':type(exc).__name__,'records':[]}
                result['query'] = {'id':query_id(spec),**spec}
                result['selection_context'] = enrichment_context
                write_json(cache, result)
            if result not in results:
                results.append(result)
        products, shortlisted = prepare()
        unchanged = {product_identity(item) for item in shortlisted
                     if before_enrichment.get(product_identity(item)) == query_id(item)}
        reviews = [r for r in reviews if (r['marketplace'], r['asin']) in unchanged]
    keyword_records=[r for result in results if result['query']['source']=='keyword_trends' for r in result.get('records',[])]
    report={'version':VERSION,'plan':plan,'results':results,'products':products,'families':shortlisted,
            'keywords':analyze_keywords(keyword_records,plan['as_of_date']), 'reviews': reviews, 'pending_queries':pending,
            'candidate_count':len(products),'family_count':len(shortlisted),'shortlist_count':len(shortlisted)}
    write_json(folder/'selection-evidence.json',report)
    evidence_id=query_id({'context':review_context,'products':shortlisted,'keywords':report['keywords']})
    known = {product_identity(p) for p in shortlisted}
    review_index = {(r['marketplace'], r['asin']): r for r in reviews}
    waiting = pending_batch(shortlisted, reviews)
    fingerprints = {'|'.join(product_identity(item)): query_id(item) for item in shortlisted
                    if product_identity(item) in review_index}
    if reviews:
        write_json(review_path,{'evidence_id':evidence_id,'context_id':review_context,
                                'fingerprints':fingerprints,'reviews':reviews})
    partial=any(r.get('collection_status')!='complete' for r in results)
    status='AWAITING_SCOPE' if plan['blocking_items'] else 'RUNNING' if pending else 'AWAITING_ANALYSIS' if waiting else 'PARTIAL' if partial else 'COMPLETE'
    if not products and not report['keywords']['directions'] and not pending and not plan['blocking_items']:
        status='BLOCKED' if partial else 'NO_CANDIDATES'
    eligible_ids = {product_identity(p) for p in shortlisted
                    if p.get('verified_parent_asin') and not any(c['status'] == 'fail' for c in p.get('filter_checks', []))}
    output={'ok':True,'status':status,'run_dir':str(folder),'evidence_path':str(folder/'selection-evidence.json'),
            'blocking_items':plan['blocking_items'],'pending_queries':pending,'analysis_requests':waiting,
            'analysis_requirements': ANALYSIS_REQUIREMENTS,
            'evidence_id':evidence_id,'counts':{'products':len(products),'shortlist':len(shortlisted),'families':len(shortlisted),
            'reviewed':len(reviews),'remaining':len(eligible_ids - set(review_index)),
            'accepted':sum(r['relevant'] for r in reviews if (r['marketplace'], r['asin']) in eligible_ids),
            'pending_parent':sum(not p.get('verified_parent_asin') for p in shortlisted),'keywords':len(report['keywords']['directions'])},
            'coverage':[{'source':r['query']['source'],'marketplace':r['query'].get('marketplace'),
                         'status':r.get('collection_status'),'reason':r.get('block_reason'),'rows':len(r.get('records',[]))} for r in results]}
    target = plan.get('shortlist_limit', 50)
    shortfall = max(0, target - output['counts']['accepted']) if target is not None else None
    output['counts'].update(target=target, shortfall=shortfall)
    output['product_history'] = []
    output['product_evidence_requests'] = []
    for review in reviews:
        if not review['relevant']:
            continue
        key = (review['marketplace'], review['asin'])
        if key not in eligible_ids:
            continue
        history_path = folder / f'product-history-{key[0]}-{key[1]}.json'
        history = json.loads(history_path.read_text()) if history_path.exists() else {}
        if (history.get('marketplace'), history.get('asin')) == key and history.get('history'):
            output['product_history'].append(history)
        life_history = (review.get('lifecycle_analysis') or {}).get('history') or []
        if not history_is_sufficient(life_history) and not (
                (history.get("marketplace"), history.get("asin")) == key
                and history_is_sufficient(history.get("history") or [])):
            output['product_evidence_requests'].append({
                'skill_action': 'collect_product_evidence', 'marketplace': key[0],
                'asins': [key[1]], 'run_dir': str(folder), 'as_of_date': plan['as_of_date'],
                'reason': 'history_coverage_incomplete', 'required_complete_months': 12,
                'next_steps': ['检查全部时间范围', '核对父体及同族子体，保留独立口径', '实际访问受阻时保留证据，不交付有限观察结论']})
    if shortfall:
        output['next_required_action'] = (
            f'目标{target}个已核实父体，当前{output["counts"]["accepted"]}个，还缺{shortfall}个。'
            '继续审阅、补数或按原范围扩大查询；不得用固定ASIN名单批量淘汰。'
            '已穷尽来源或遇到实际访问限制时，保留PARTIAL并说明尝试与缺口。')
        if output['status'] == 'COMPLETE':
            output['status'] = 'PARTIAL'
    if output['product_evidence_requests']:
        output['evidence_next_action'] = '先调用collect_product_evidence读取销量趋势表，再提交有实际月度点的lifecycle_analysis；不能只检查查竞品首页文字就认定无历史。'
    if not waiting and not pending and not plan['blocking_items'] and (products or report['keywords']['directions']):
        eligible=[p for p in shortlisted if p.get('verified_parent_asin') and product_identity(p) in review_index
                  and not any(c['status']=='fail' for c in p.get('filter_checks',[]))]
        if plan.get('shortlist_limit') is not None:
            eligible = [p for p in eligible if review_index[product_identity(p)]['relevant']][:plan['shortlist_limit']]
        rows=summarize_products(eligible,reviews)
        output['analysis_gaps'] = [
            {'marketplace': p['marketplace'], 'asin': p['asin'], 'missing': evidence['analysis_missing']}
            for p in eligible if review_index[product_identity(p)]['relevant']
            if (evidence := analysis_columns(p, review_index[product_identity(p)], plan['as_of_date']))['analysis_missing']
        ]
        if output['analysis_gaps']:
            output['status'] = 'PARTIAL'
        missing={(p['marketplace'],p['asin']) for p in eligible if any(c['status']=='missing' for c in p.get('filter_checks',[]))}
        for row in rows:
            if (row['站点'],row['ASIN']) in missing: row['优先级']='待补数据';row['证据缺项']+='；筛选条件无法完整核实'
        if missing or any(row['优先级']=='待补数据' for row in rows):output['status']='PARTIAL'
        if output['counts']['pending_parent']:
            output['status'] = 'PARTIAL' if rows else 'AWAITING_ENRICHMENT'
        if output['status'] != 'AWAITING_ENRICHMENT':
            output['workbook']=export_workbook(folder/'开品结果.xlsx',rows,report)
            if output['workbook']['image_failures']: output['status']='PARTIAL'
    write_json(folder/'selection-status.json',output)
    return output
