import json
from pathlib import Path
import sys
import tempfile
import unittest

from PIL import Image
from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))
from asin_keyword_breakout import normalize_item, render
from discovery_workbook import HEADERS


class ReverseKeywordReportTests(unittest.TestCase):
    def test_only_result_workbook_is_delivered_and_raw_evidence_is_preserved(self):
        item = normalize_item({
            "keywords": "brotbeutel bienenwachs", "keywordCn": "面包袋蜂蜡",
            "searches": 16740, "searchesGrowth": 16.23, "exactPpc": 0.97,
            "purchases": 71, "latest7daysAds": 189, "products": 412,
            "searchesTrend": [{"month": month, "searches": searches} for month, searches in
                              [("202606", 0), ("202607", 120), ("202608", 8000), ("202609", 16740)]],
        })
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            Image.new("RGB", (80, 80), "#e0b36d").save(root / "product.png")
            evidence = root / "evidence.json"
            evidence.write_text(json.dumps({
                "marketplace": "DE", "asin": "B0HJ66GXZK",
                "product": {"name": "Bienenwachs Brotbeutel", "url": "https://www.amazon.de/dp/B0HJ66GXZK",
                            "image_file": "product.png", "Review数量": 27, "预估月销量": 125,
                            "结论": "优先调研", "理由": "已核实的商品判断"},
                "source_url": "https://www.sellersprite.com/v3/keyword-reverse/?q=B0HJ66GXZK",
                "observed_at": "2026-10-19T00:00:00+00:00", "collected_count": 174,
                "reported_total": 174, "keywords": [item, {**item, "keyword": "bienenwachs brotbeutel"}],
            }), encoding="utf-8")
            selected = root / "selected.json"
            selected.write_text(json.dumps([{"marketplace": "DE", "asin": "B0HJ66GXZK",
                                             "keyword": "brotbeutel bienenwachs",
                                             "direction": "蜂蜡面包袋"},
                                            {"marketplace": "DE", "asin": "B0HJ66GXZK",
                                             "keyword": "bienenwachs brotbeutel"}]), encoding="utf-8")
            evidence_before = evidence.read_bytes()
            result = render([evidence], selected, root / "report")
            workbook = load_workbook(result["workbook"])
            self.assertEqual(workbook.sheetnames, ["开品结果"])
            self.assertEqual(Path(result["workbook"]).name, "开品结果.xlsx")
            self.assertEqual(list((root / "report").iterdir()), [Path(result["workbook"])])
            self.assertNotIn("preview", result)
            self.assertEqual(result["rows"], 1)
            sheet = workbook["开品结果"]
            self.assertEqual([cell.value for cell in sheet[1]], HEADERS)
            self.assertEqual(sheet.max_row, 2)
            self.assertEqual(sheet["C2"].value, 27)
            self.assertEqual(sheet["G2"].value, 125)  # Product sales, not keyword searches.
            self.assertEqual(sheet["K2"].value, "B0HJ66GXZK")
            self.assertEqual(sheet["L2"].hyperlink.target, "https://www.amazon.de/dp/B0HJ66GXZK")
            self.assertEqual(sheet["N2"].value, "优先调研")
            self.assertIn("已核实的商品判断", sheet["O2"].value)
            self.assertIn("brotbeutel bienenwachs", sheet["O2"].value)
            self.assertIn("bienenwachs brotbeutel", sheet["O2"].value)
            self.assertIn("16740", sheet["O2"].value)
            self.assertEqual(len(sheet._images), 1)  # Product image only, no trend charts.
            workbook.close()
            self.assertEqual(evidence.read_bytes(), evidence_before)
            # Rebuilding reuses embedded product images and still creates only one file.
            rebuilt = render([evidence], selected, root / "report")
            with_image = load_workbook(rebuilt["workbook"])
            self.assertEqual(len(with_image.active._images), 1)
            with_image.close()

    def test_markets_stay_separate_and_missing_product_facts_stay_blank(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            sources, choices = [], []
            for market in ("US", "DE"):
                source = root / f"{market}.json"
                source.write_text(json.dumps({
                    "marketplace": market, "asin": "B0HJ66GXZK",
                    "product": {"name": "Bread Bag", "url": f"https://example.com/{market}"},
                    "source_url": "source", "observed_at": "today",
                    "collected_count": 1, "reported_total": 1,
                    "keywords": [normalize_item({"keywords": "bread bag", "searches": 1000})],
                }))
                sources.append(source)
                choices.append({"marketplace": market, "asin": "B0HJ66GXZK", "keyword": "bread bag"})
            selection = root / "selected.json"
            selection.write_text(json.dumps(choices))
            result = render(sources, selection, root / "out")
            self.assertEqual(result["rows"], 2)
            workbook = load_workbook(result["workbook"])
            self.assertEqual(workbook.sheetnames, ["开品结果"])
            sheet = workbook.active
            self.assertEqual([sheet.cell(i, 1).value for i in (2, 3)], ["US", "DE"])
            for i in (2, 3):
                for column in (2, 3, 4, 5, 6, 7, 9, 10, 14):
                    self.assertFalse(sheet.cell(i, column).value)
            self.assertEqual(len(sheet._images), 0)
            workbook.close()

    def test_unverified_keyword_is_rejected(self):
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            evidence = root / "evidence.json"
            evidence.write_text(json.dumps({
                "marketplace": "DE", "asin": "B0HJ66GXZK", "product": {"name": "Bread Bag"},
                "source_url": "source", "observed_at": "today", "collected_count": 0,
                "reported_total": 0, "keywords": [],
            }))
            selection = root / "selection.json"
            selection.write_text(json.dumps([{"marketplace": "DE", "asin": "B0HJ66GXZK",
                                              "keyword": "invented term"}]))
            with self.assertRaisesRegex(ValueError, "反查证据没有关键词"):
                render([evidence], selection, root / "out")


class ReversePaginationTests(unittest.TestCase):
    def run_collection(self, folder, fail_second=False):
        from contextlib import contextmanager
        from types import SimpleNamespace
        from unittest.mock import MagicMock, patch
        import asin_keyword_breakout as collector
        import selection_collectors
        asin = 'B000000001'
        page = MagicMock()
        page.url = 'https://www.sellersprite.com/v3/keyword-reverse/?q=' + asin
        responses = []
        for skip in (0, 2):
            response = MagicMock()
            response.url = 'https://www.sellersprite.com' + collector.API_PATH
            response.status = 200
            response.request.method = 'POST'
            response.request.post_data = json.dumps({'asin': asin, 'skip': skip, 'limit': 2})
            response.json.return_value = {'code': 'OK', 'data': {
                'asin': asin, 'station': 'DE', 'total': 4,
                'items': [{'keywords': f'word{n}'} for n in range(skip, skip + 2)]}}
            responses.append(response)
        calls = []
        @contextmanager
        def response_context(predicate, **kw):
            response = responses[len(calls)]
            self.assertTrue(predicate(response))
            calls.append(response)
            yield SimpleNamespace(value=response)
        page.expect_response.side_effect = response_context
        page.locator.return_value.first.count.return_value = 1
        page.locator.return_value.first.is_enabled.return_value = True
        output = Path(folder) / 'evidence.json'
        def click():
            self.assertEqual(json.loads(output.read_text())['collected_count'], 2)
            if fail_second:
                raise RuntimeError('page timed out')
        page.locator.return_value.first.click.side_effect = click
        @contextmanager
        def session(payload, url, **kw):
            self.assertNotIn('q=', url)  # no duplicate first-page query
            yield page
        with patch.object(selection_collectors, 'session', session), \
             patch.object(collector, 'assert_access'), \
             patch.object(collector, '_product', return_value={'image_url': ''}):
            if fail_second:
                with self.assertRaisesRegex(RuntimeError, 'page timed out'):
                    collector.collect_reverse('DE', asin, output)
                result = json.loads(output.read_text())
            else:
                result = collector.collect_reverse('DE', asin, output)
        page.request.post.assert_not_called()
        page.locator.return_value.first.click.assert_called_once()
        return result

    def test_pagination_clicks_browser_and_preserves_matching_rows(self):
        with tempfile.TemporaryDirectory() as folder:
            result = self.run_collection(folder)
            self.assertEqual(result['collected_count'], 4)
            self.assertEqual(result['status'], 'complete')

    def test_later_page_error_preserves_earlier_evidence(self):
        with tempfile.TemporaryDirectory() as folder:
            result = self.run_collection(folder, fail_second=True)
            self.assertEqual(result['collected_count'], 2)
            self.assertEqual(result['status'], 'partial')


if __name__ == "__main__":
    unittest.main()
